from __future__ import annotations

import json
import re
from pathlib import Path
from typing import Any

from .schema import KnowledgeUnit


TERM_PATTERN = re.compile(r"\b[A-Z][A-Z0-9-]{1,}\b")
FORMULA_PATTERN = re.compile(r"[^。；;\n]{0,80}(?:=|→|⇒|∑|∫|∂|×|÷)[^。；;\n]{0,80}")
PPT_TERM_OVERRIDES = {
    3: ["GOFAI", "符号主义", "知识", "规则"],
    5: ["物理符号系统", "PSSH", "符号处理"],
    6: ["中文屋", "符号操作", "理解", "强人工智能"],
    7: ["IF-THEN", "知识表示", "产生式规则"],
    8: ["IF-THEN", "规则表达", "专家系统"],
    9: ["正向推理", "前向链", "数据驱动推理"],
    10: ["反向推理", "后向链", "目标驱动推理"],
    11: ["ELIZA", "自然语言处理", "模式匹配"],
}


def _unique(values: list[str]) -> list[str]:
    return list(dict.fromkeys(value.strip() for value in values if value and value.strip()))


class PptxCourseIngestor:
    def ingest(
        self,
        source: str | Path,
        *,
        course_id: str,
        course_version: str,
        source_id: str,
        language: str = "en-US",
        teacher_confirmed: bool = True,
    ) -> list[KnowledgeUnit]:
        from pptx import Presentation

        path = Path(source).resolve()
        presentation = Presentation(path)
        slide_width = max(1, int(presentation.slide_width))
        slide_height = max(1, int(presentation.slide_height))
        units: list[KnowledgeUnit] = []
        page_ids: list[str] = []

        for page, slide in enumerate(presentation.slides, start=1):
            elements: list[dict[str, Any]] = []
            for shape_index, shape in enumerate(slide.shapes, start=1):
                values: list[str] = []
                if getattr(shape, "has_text_frame", False):
                    values.extend(line.strip() for line in shape.text.splitlines() if line.strip())
                if getattr(shape, "has_table", False):
                    for row in shape.table.rows:
                        values.append(" | ".join(cell.text.strip() for cell in row.cells if cell.text.strip()))
                text = "\n".join(_unique(values)).strip()
                if not text:
                    continue
                elements.append(
                    {
                        "index": shape_index,
                        "text": text,
                        "bbox": [
                            round(float(shape.left) / slide_width, 6),
                            round(float(shape.top) / slide_height, 6),
                            round(float(shape.width) / slide_width, 6),
                            round(float(shape.height) / slide_height, 6),
                        ],
                    }
                )
            try:
                notes = "\n".join(
                    line.strip()
                    for line in slide.notes_slide.notes_text_frame.text.splitlines()
                    if line.strip()
                )
            except Exception:
                notes = ""
            title = elements[0]["text"].splitlines()[0] if elements else f"第{page}页"
            content_parts = [item["text"] for item in elements]
            if notes:
                content_parts.append("教师讲稿：" + notes)
            content = "\n".join(content_parts)
            page_id = f"{course_id}:ppt:p{page:03d}"
            page_ids.append(page_id)
            terms = _unique([*TERM_PATTERN.findall(content), *PPT_TERM_OVERRIDES.get(page, [])])
            formulas = _unique(match.group(0).strip() for match in FORMULA_PATTERN.finditer(content))
            units.append(
                KnowledgeUnit(
                    knowledge_id=page_id,
                    course_id=course_id,
                    course_version=course_version,
                    title=title,
                    content=content,
                    source_id=source_id,
                    source_name=path.name,
                    source_page=page,
                    chapter="早期人工智能技术",
                    section=title,
                    unit_type="slide",
                    language=language,
                    authority=1,
                    teacher_confirmed=teacher_confirmed,
                    terms=terms,
                    formulas=formulas,
                    metadata={"elementCount": len(elements), "hasNotes": bool(notes)},
                )
            )
            for item in elements:
                if len(item["text"]) < 24:
                    continue
                element_id = f"{page_id}:e{int(item['index']):02d}"
                element_terms = _unique([*TERM_PATTERN.findall(item["text"]), *PPT_TERM_OVERRIDES.get(page, [])])
                element_formulas = _unique(
                    match.group(0).strip() for match in FORMULA_PATTERN.finditer(item["text"])
                )
                units.append(
                    KnowledgeUnit(
                        knowledge_id=element_id,
                        course_id=course_id,
                        course_version=course_version,
                        title=title,
                        content=item["text"],
                        source_id=source_id,
                        source_name=path.name,
                        source_page=page,
                        chapter="早期人工智能技术",
                        section=title,
                        unit_type="slide_element",
                        language=language,
                        authority=1,
                        teacher_confirmed=teacher_confirmed,
                        terms=element_terms,
                        formulas=element_formulas,
                        bbox=item["bbox"],
                        parent_id=page_id,
                    )
                )

        page_lookup = {value: index for index, value in enumerate(page_ids)}
        for unit in units:
            if unit.knowledge_id not in page_lookup:
                continue
            index = page_lookup[unit.knowledge_id]
            unit.adjacent_ids = page_ids[max(0, index - 1):index] + page_ids[index + 1:index + 2]
        return units


class CuratedWebIngestor:
    def ingest(self, registry: str | Path, *, course_id: str, course_version: str) -> list[KnowledgeUnit]:
        values = json.loads(Path(registry).read_text(encoding="utf-8"))
        units: list[KnowledgeUnit] = []
        for source in values.get("sources", []):
            source_id = str(source["source_id"])
            common = {
                "course_id": course_id,
                "course_version": course_version,
                "source_id": source_id,
                "source_name": str(source["name"]),
                "source_url": str(source.get("url", "")),
                "chapter": str(source.get("chapter", "")),
                "language": str(source.get("language", "zh-CN")),
                "authority": int(source.get("authority", 2)),
                "teacher_confirmed": bool(source.get("teacher_confirmed", False)),
            }
            summary = str(source.get("summary", "")).strip()
            if summary:
                units.append(
                    KnowledgeUnit(
                        knowledge_id=f"{course_id}:web:{source_id}:summary",
                        title=str(source.get("title", source["name"])),
                        content=summary,
                        section=str(source.get("title", source["name"])),
                        unit_type="external_summary",
                        terms=_unique(list(source.get("terms", []))),
                        metadata={"ingestMode": source.get("ingest_mode", "curated")},
                        **common,
                    )
                )
            for index, point in enumerate(source.get("key_points", []), start=1):
                text = str(point).strip()
                if not text:
                    continue
                units.append(
                    KnowledgeUnit(
                        knowledge_id=f"{course_id}:web:{source_id}:p{index:02d}",
                        title=str(source.get("title", source["name"])),
                        content=text,
                        section=str(source.get("title", source["name"])),
                        unit_type="external_key_point",
                        terms=_unique([*source.get("terms", []), *TERM_PATTERN.findall(text)]),
                        parent_id=f"{course_id}:web:{source_id}:summary",
                        metadata={"ingestMode": source.get("ingest_mode", "curated")},
                        **common,
                    )
                )
        return units


class WebSnapshotIngestor:
    def ingest(
        self,
        snapshot_dir: str | Path,
        registry: str | Path,
        *,
        course_id: str,
        course_version: str,
    ) -> list[KnowledgeUnit]:
        root = Path(snapshot_dir)
        registry_value = json.loads(Path(registry).read_text(encoding="utf-8"))
        sources = {str(item["source_id"]): item for item in registry_value.get("sources", [])}
        units: list[KnowledgeUnit] = []
        for source_id, source in sources.items():
            snapshot_path = root / f"{source_id}.json"
            if not snapshot_path.is_file():
                continue
            snapshot = json.loads(snapshot_path.read_text(encoding="utf-8"))
            if snapshot.get("status") != "ok":
                continue
            if source_id == "google-mlcc":
                continue
            heading = str(source.get("title", source.get("name", source_id)))
            buffer: list[str] = []
            chunk_index = 0

            def append_chunk() -> None:
                nonlocal buffer, chunk_index
                text = "\n".join(buffer).strip()
                if len(text) < 40:
                    buffer = []
                    return
                chunk_index += 1
                units.append(
                    KnowledgeUnit(
                        knowledge_id=f"{course_id}:snapshot:{source_id}:c{chunk_index:03d}",
                        course_id=course_id,
                        course_version=course_version,
                        title=heading,
                        content=text,
                        source_id=source_id,
                        source_name=str(source.get("name", source_id)),
                        source_url=str(source.get("url", "")),
                        chapter=str(source.get("chapter", "")),
                        section=heading,
                        unit_type="official_web_chunk",
                        language=str(source.get("language", "zh-CN")),
                        authority=int(source.get("authority", 2)),
                        teacher_confirmed=False,
                        terms=_unique([*source.get("terms", []), *TERM_PATTERN.findall(text)]),
                        metadata={"snapshot": snapshot_path.name, "collectedAt": snapshot.get("collectedAt")},
                    )
                )
                buffer = []

            for block in snapshot.get("blocks", []):
                block_type = str(block.get("type", ""))
                text = str(block.get("text", "")).strip()
                if not text:
                    continue
                if block_type in {"h1", "h2", "h3", "h4"}:
                    append_chunk()
                    heading = text
                    buffer = [text]
                    continue
                if sum(len(value) for value in buffer) + len(text) > 1100:
                    append_chunk()
                buffer.append(text)
            append_chunk()
        return units

import io,struct,tempfile,unittest,zipfile
from pathlib import Path
from lxml import etree as E
from fontTools.ttLib import TTFont
from app.pptx_font_embedding import eot_font,embed_fonts,P,R

class EmbeddedFontTests(unittest.TestCase):
    def font(self):
        path=Path(__file__).resolve().parents[3]/'.local-tools/SmartCourseJP-Regular.ttf'
        if not path.is_file():self.skipTest('font fixture not installed')
        return path
    def test_eot_header_and_payload(self):
        source=self.font().read_bytes();data=eot_font(source)
        size,font_size,version,flags=struct.unpack_from('<IIII',data)
        self.assertEqual(size,len(data));self.assertEqual(font_size,len(source));self.assertEqual(version,0x00020001)
        self.assertEqual(data[-font_size:],source)
        self.assertEqual(struct.unpack_from('<H',data,34)[0],0x504C)
        self.assertEqual(TTFont(io.BytesIO(data[-font_size:]))['name'].getDebugName(1),'SmartCourse JP')
    def test_package_relationships_and_two_faces(self):
        from pptx import Presentation
        with tempfile.TemporaryDirectory() as folder:
            path=Path(folder)/'deck.pptx';p=Presentation();p.slides.add_slide(p.slide_layouts[6]);p.save(path)
            embed_fonts(path,self.font(),self.font(),'SmartCourse JP')
            with zipfile.ZipFile(path) as z:
                root=E.fromstring(z.read('ppt/presentation.xml'));fonts=root.find('{'+P+'}embeddedFontLst')
                self.assertEqual(root.get('embedTrueTypeFonts'),'1');self.assertEqual(len(fonts),1)
                self.assertEqual(len(fonts[0]),3)
                for kind in ['regular','bold']:
                    self.assertIsNotNone(fonts[0].find('{'+P+'}'+kind).get('{'+R+'}id'))
                    self.assertTrue(z.read('ppt/fonts/course-'+kind+'.fntdata'))

if __name__=='__main__':unittest.main()

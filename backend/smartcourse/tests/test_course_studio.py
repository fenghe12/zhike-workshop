import copy
import json
import tempfile
import threading
import time
import unittest
import wave
import zipfile
from http.server import ThreadingHTTPServer
from pathlib import Path
from urllib import request, error
from unittest.mock import Mock, patch
from xml.etree import ElementTree as ET

from pptx import Presentation
from app.account_policy import POLICY_VERSION
from app.config import Settings
from app.course_revisions import create_revision, preserve_revision_narration
from app.feedback_repair import _work, submit_repair
from app.server import SmartCourseApplication, make_handler
from app.task_model import create_task
from app.teaching_memory import generation_memory, memory_prompt
from app.teacher_features import create_question, review_question, generate_question_batch


class CourseStudioTests(unittest.TestCase):
    def test_memory_candidate_can_select_an_exact_source_snippet(self):
        self.app.pipeline_runner.gateway.respond_json.return_value={'entries':[{
            'kind':'definition','name':'神经元','value':'输入与权重共同决定输出。','pages':[1],'sourceIndex':2}]}
        status,value=self.call('extract-memory',{})
        self.assertEqual(status,200,value)
        self.assertEqual(value['entries'][0]['status'],'candidate')
        self.assertEqual(value['entries'][0]['pages'],[1])
        prompt=self.app.pipeline_runner.gateway.respond_json.call_args.args[0]
        data=json.loads(prompt.rsplit('\n',1)[1])
        self.assertEqual(data['sourceSnippets']['1'][2],self.pages[0]['script_text'])

    def test_content_review_uses_selected_translation_and_original_facts(self):
        self.app.pipeline_runner.gateway.respond_json.return_value={'page':1,'issues':[
            {'quote':self.translated[0]['script_text'],'message':'Check the selected narration.','suggestion':'Keep the original facts.'}],
            'summary':'Selected translation reviewed.'}
        status,value=self.call('review',{'page':1,'language':'en-US','withImages':False})
        self.assertEqual(status,200,value);self.assertEqual(value['language'],'en-US')
        supplied=json.loads(self.app.pipeline_runner.gateway.respond_json.call_args.args[0].rsplit('\n',1)[1])
        self.assertEqual(supplied['reviewedPage']['script_text'],self.translated[0]['script_text'])
        self.assertEqual(supplied['originalPage']['script_text'],self.pages[0]['script_text'])
        self.assertEqual(self.call('review',{'page':1,'language':'ru-RU','withImages':False})[0],400)

    def test_narration_feedback_retry_preserves_slide_text(self):
        feedback=self.feedback()
        with patch('app.feedback_repair.start_worker'):
            record=submit_repair(self.app,self.task,self.uid('owner'),
                {'feedbackId':feedback['id'],'kind':'narration'})
            self.app.studio.save('repairs',record['id'],{**record,'status':'failed','attempts':4,
                'message':'Previous model response used the wrong slide text.'})
            retried=submit_repair(self.app,self.task,self.uid('owner'),
                {'feedbackId':feedback['id'],'kind':'narration'})
        self.assertEqual(retried['attempts'],0);self.assertEqual(retried['status'],'queued')
        self.assertEqual(len(retried['previousFailures']),1)
        self.app.pipeline_runner.gateway.respond_json.return_value={
            'script_text':'Training improves predictions by adjusting model weights according to the prediction errors.',
            'replacements':[{'oldText':'Text not in this slide','text':'Unrequested slide edit'}]}
        original=(self.course_root/'s2/en-US/translated.pptx').read_bytes()
        _work(self.app,record['id'])
        saved=self.app.studio.state()['repairs'][record['id']]
        revised=self.app.get_task(saved['releaseTaskId'])
        self.assertEqual({x['kind'] for x in revised['revisionChanges']},{'script'})
        self.assertEqual((self.app.studio.root(revised['id'])/'s2/en-US/translated.pptx').read_bytes(),original)

    def test_template_batch_inherits_saved_pages_minutes_and_requirements(self):
        self.app.mutate_task(self.task['id'],lambda task,_:task['inputs'].update(
            totalPages=2,durationMinutes=1,requirements='Preserve binary examples and use short sentences.'))
        status,template=self.call('/v1/studio/templates',{'courseId':self.task['id'],'name':'Saved course settings'})
        self.assertEqual(status,201,template)
        status,batch=self.call('/v1/studio/templates/'+template['id']+'/launch',
            {'topics':['First short lesson','Second short lesson'],'execute':False})
        self.assertEqual(status,200,batch);self.assertEqual(len(batch['taskIds']),2)
        for identifier in batch['taskIds']:
            task=self.app.get_task(identifier)
            self.assertEqual(task['inputs']['totalPages'],2)
            self.assertEqual(task['inputs']['durationMinutes'],1)
            self.assertIn('Preserve binary examples',task['inputs']['requirements'])
            self.assertEqual(task['targetLanguages'],['英语'])

    def test_question_generation_retries_source_language_for_target_question(self):
        good={'questions':[{'page_id':1,'statement':'Inputs and weights determine the neuron output.',
            'correct':True,'explanation':'The course states that inputs and weights jointly determine the output.',
            'source_excerpt':self.pages[0]['script_text']}]}
        bad=copy.deepcopy(good)
        bad['questions'][0]['statement']='输入与权重共同决定神经元的输出。'
        model=self.app.pipeline_runner.gateway.respond_json;model.side_effect=[bad,good]
        result=generate_question_batch(self.app.settings,self.task['id'],self.uid('owner'),self.app.pipeline_runner.gateway,1,'en-US')
        self.assertEqual(model.call_count,2)
        self.assertEqual(result['questions'][0]['prompt'],good['questions'][0]['statement'])

    def test_quiz_media_is_readable_only_by_its_own_course_participant(self):
        self.member('learner','learner');self.member('other','learner');self.call('publish',{'status':'published'})
        folder=self.course_root/'interruptions/answer-fixture';folder.mkdir(parents=True)
        video=folder/'answer.mp4';video.write_bytes(b'private quiz media')
        audio=folder/'answer.wav';audio.write_bytes(b'private quiz audio')
        poster=folder/'answer.png';poster.write_bytes(b'private quiz poster')
        private=folder/'request.json';private.write_text('{}')
        record={'id':'answer-fixture','courseId':self.task['id'],'learnerId':self.uid('learner'),
            'purpose':'quiz_remediation','status':'completed','videoUrl':self.url(video),'audioUrl':self.url(audio)}
        self.app.store.transact(lambda data:data.setdefault('interruptions',{}).update({record['id']:record}))
        for path in (video,audio,poster):
            self.assertEqual(self.call(self.url(path),user='learner',raw=True)[0],200)
            self.assertEqual(self.call(self.url(path),user='other',raw=True)[0],403)
            self.assertEqual(self.call(self.url(path),user=None,raw=True)[0],401)
        self.assertEqual(self.call(self.url(private),user='learner',raw=True)[0],403)
        self.assertEqual(self.call('/artifacts/tasks/'+self.task['id']+'/execution-journal.json',user='learner',raw=True)[0],403)

    def test_generated_true_false_questions_cover_all_supported_locales(self):
        statements={
            'zh-CN':'输入与权重共同决定神经元的输出。',
            'en-US':'Inputs and weights determine the neuron output.',
            'ja-JP':'入力と重みはニューロンの出力を決定します。',
            'ko-KR':'입력과 가중치는 뉴런의 출력을 결정합니다.',
            'de-DE':'Eingaben und Gewichte bestimmen die Ausgabe des Neurons.',
            'fr-FR':'Les entrées et les poids déterminent la sortie du neurone.',
            'ru-RU':'Входные данные и веса определяют выход нейрона.',
            'es-ES':'Las entradas y los pesos determinan la salida de la neurona.'
        }
        expected={'zh-CN':('正确','错误'),'en-US':('True','False'),'ja-JP':('正しい','誤り'),
                  'ko-KR':('맞음','틀림'),'de-DE':('Richtig','Falsch'),'fr-FR':('Vrai','Faux'),
                  'ru-RU':('Верно','Неверно'),'es-ES':('Verdadero','Falso')}
        for locale,statement in statements.items():
            with self.subTest(locale=locale):
                self.app.pipeline_runner.gateway.respond_json.return_value={'questions':[{
                    'page_id':1,'statement':statement,'correct':True,'explanation':statement,
                    'source_excerpt':self.pages[0]['script_text']}]}
                value=generate_question_batch(self.app.settings,self.task['id'],self.uid('owner'),
                    self.app.pipeline_runner.gateway,1,locale)
                self.assertEqual(value['created'],1)
                self.assertEqual(tuple(x['text'] for x in value['questions'][0]['options']),expected[locale])

    def setUp(self):
        self.directory=tempfile.TemporaryDirectory()
        self.root=Path(self.directory.name)
        settings=Settings(self.root/'data',self.root/'artifacts','http://model.local','fixture-model','fixture-token','127.0.0.1',0)
        self.app=SmartCourseApplication(settings)
        self.app.execute_task=Mock(return_value={'queued':True})
        self.app.pipeline_runner.gateway.respond_json=Mock()
        consent={'version':POLICY_VERSION,'termsAccepted':True,'privacyAccepted':True}
        self.users={name:self.app.store.register('studio-'+name,'test-password',name,policy_consent=consent) for name in ('owner','reviewer','learner','other')}
        self.task=create_task({'name':'Teaching fixture','mode':'topic','sourceLanguage':'中文','targetLanguages':['英语'],
                               'inputs':{'topic':'Neural networks','totalPages':2,'durationMinutes':1}})
        self.task.update(createdBy=self.uid('owner'),status='completed',teachingMemory={'enabled':True,'style':'Private teaching preference','revision':1})
        root=self.app.studio.root(self.task['id'])
        self.course_root=root
        self.pages=[{'page_id':1,'title':'神经元','key_points':['输入与权重'], 'formulas':[], 'script_text':'输入与权重共同决定神经元的输出。'},
                    {'page_id':2,'title':'学习','key_points':['根据误差调整权重'], 'formulas':[], 'script_text':'训练根据预测误差调整模型权重。'}]
        self.translated=[{'page_id':1,'title':'Neurons','key_points':['Inputs and weights'], 'formulas':[],
                          'script_text':'Inputs and weights determine the output of a neuron. This is the basic calculation.'},
                         {'page_id':2,'title':'Learning','key_points':['Adjust weights using error'], 'formulas':[],
                          'script_text':'Training adjusts model weights using prediction errors. The model learns from the data.'}]
        self.write(root/'s1/course_structure.json',{'task_id':self.task['id'],'pages':self.pages})
        self.write(root/'s2/en-US/translated_structure.json',{'task_id':self.task['id'],'language_code':'en-US','target_language':'English','pages':self.translated})
        self.write(root/'s2/en-US/translation_report.json',{'passed':True,'pages':[{'page_id':1,'passed':True},{'page_id':2,'passed':True}]})
        self.write(root/'s3/en-US/narration_structure.json',{'task_id':self.task['id'],'pages':[{**p,'script_text':p['script_text']+' Previously normalized.'} for p in self.translated]})
        for name,pages in [('s1/generated.pptx',self.pages),('s2/en-US/translated.pptx',self.translated)]:
            deck=Presentation()
            for row in pages:
                slide=deck.slides.add_slide(deck.slide_layouts[1]);slide.shapes.title.text=row['title'];slide.placeholders[1].text='\n'.join(row['key_points'])
            (root/name).parent.mkdir(parents=True,exist_ok=True);deck.save(root/name)
        folder=root/'runtime/en-US/s3_output'
        for page in (1,2):
            audio=folder/'segments'/f'page_{page:03d}.wav';audio.parent.mkdir(parents=True,exist_ok=True)
            with wave.open(str(audio),'wb') as f:f.setparams((1,2,16000,0,'NONE','none'));f.writeframes(b'\0\0'*16000)
            self.write(audio.with_suffix('.input.json'),{'input':{'taskId':self.task['id'],'text':self.translated[page-1]['script_text']+' Previously normalized.'},'generationId':'fixture-'+str(page)})
            avatar=root/'runtime/en-US/s4_output/avatars'/f'page_{page:03d}.input.json'
            self.write(avatar,{'input':{'taskRoot':str(root/'runtime'),'page':page}})
        self.write(folder/'timeline.json',{'pages':[{'page_id':1,'start_time':0,'end_time':2},{'page_id':2,'start_time':2,'end_time':4}]})
        (folder/'subtitles.srt').write_text('1\n00:00:00,000 --> 00:00:01,800\nInputs and weights.\n\n2\n00:00:02,000 --> 00:00:03,800\nTraining adjusts weights.\n')
        self.write(root/'runtime/en-US/s4_output/bbox.json',{'bboxes':[{'slide_number':p,'img_w':1920,'img_h':1080,'x_px':10,'y_px':10,'size_px':300} for p in (1,2)]})
        delivery=root/'s5/en-US';delivery.mkdir(parents=True)
        (delivery/'video.mp4').write_bytes(b'fixture media bytes, not an inference result')
        (delivery/'lecture.txt').write_text('Teaching fixture')
        self.task['outputs']={'s1':{'pageCount':2,'courseStructure':self.url(root/'s1/course_structure.json'),'pptx':self.url(root/'s1/generated.pptx')},
                              's2':{'languages':{'en-US':{'translatedStructure':self.url(root/'s2/en-US/translated_structure.json'),'passed':True}}},
                              's3':{'languages':{'en-US':{'durationSeconds':4,'timeline':self.url(folder/'timeline.json')}}},
                              's4':{'languages':{'en-US':{'bbox':self.url(root/'runtime/en-US/s4_output/bbox.json')}}},
                              's5':{'languages':{'en-US':{'video':self.url(delivery/'video.mp4'),'lectureScript':self.url(delivery/'lecture.txt'),
                                                       'srt':self.url(folder/'subtitles.srt'),'coursePptx':self.url(root/'s2/en-US/translated.pptx')}}}}
        self.write(root/'execution-journal.json',{'teachingMemory':self.task['teachingMemory']})
        for stage in self.task['stages']:stage.update(status='completed',progress=100)
        self.app.store.transact(lambda d:d['tasks'].update({self.task['id']:self.task}))
        self.server=ThreadingHTTPServer(('127.0.0.1',0),make_handler(self.app));self.thread=threading.Thread(target=self.server.serve_forever,daemon=True);self.thread.start()
        self.base='http://127.0.0.1:'+str(self.server.server_port)

    def tearDown(self):
        self.server.shutdown();self.server.server_close();self.thread.join();self.app.task_queue.close();self.directory.cleanup()

    def uid(self,name):return self.users[name]['user']['id']
    def url(self,path):return '/artifacts/'+str(path.relative_to(self.app.settings.artifact_root))
    def write(self,path,value):path.parent.mkdir(parents=True,exist_ok=True);path.write_text(json.dumps(value,ensure_ascii=False),encoding='utf-8')
    def call(self,path='',payload=None,user='owner',method=None,raw=False):
        path=path if path.startswith('/') else '/v1/studio/courses/'+self.task['id']+('/'+path if path else '')
        headers={'Content-Type':'application/json'}
        if user:headers['Authorization']='Bearer '+self.users[user]['accessToken']
        req=request.Request(self.base+path,headers=headers,method=method,data=json.dumps(payload).encode() if payload is not None else None)
        try:
            with request.build_opener(request.ProxyHandler({})).open(req,timeout=15) as r:return r.status,r.read() if raw else json.load(r)
        except error.HTTPError as e:return e.code,e.read() if raw or method=='HEAD' else json.load(e)
    def member(self,name,role):return self.call('members',{'account':'studio-'+name,'role':role})
    def feedback(self,category='narration'):
        status,row=self.call('/v1/feedback',{'taskId':self.task['id'],'language':'en-US','page':2,'categories':[category],'comment':'Please correct this page.'})
        self.assertEqual(status,201,row);return row

    def test_roles_publication_and_shared_files_do_not_grant_editing_or_private_records(self):
        self.assertEqual(self.call(user='other')[0],403)
        self.member('learner','learner');self.member('reviewer','reviewer')
        self.assertEqual(self.call(user='learner')[0],403)
        self.assertEqual(self.call('publish',{'status':'published'})[0],200)
        status,value=self.call(user='learner');self.assertEqual(status,200,value)
        self.assertNotIn('teachingMemory',value['task']);self.assertNotIn('memory',value)
        self.assertEqual(self.call('memory',{'revision':0,'entries':[]},user='learner')[0],403)
        self.assertEqual(self.call('/v1/courses/'+self.task['id']+'/teacher/questions',user='learner')[0],403)
        video=self.task['outputs']['s5']['languages']['en-US']['video']
        self.assertEqual(self.call(video,user='learner',raw=True)[0],200)
        journal=self.url(self.course_root/'execution-journal.json')
        self.assertEqual(self.call(journal,user='learner',raw=True)[0],403)
        self.assertEqual(self.call(journal,user='other',method='HEAD',raw=True)[0],403)
        self.assertEqual(self.call('comment',{'page':1,'text':'Check this explanation.'},user='reviewer')[0],200)
        self.assertEqual(self.call('members',{'userId':self.uid('learner'),'role':'remove'})[0],200)
        self.assertEqual(self.call(video,user='learner',raw=True)[0],403)

    def test_accepted_course_memory_is_frozen_and_scoped_to_authorized_users(self):
        entries=[{'kind':'term','name':'神经元','value':'neuron','language':'en-US','pages':[1],'status':'accepted'},
                 {'kind':'definition','name':'Candidate','value':'Not adopted','pages':[2],'status':'candidate'}]
        status,value=self.call('memory',{'revision':0,'entries':entries});self.assertEqual(status,200,value)
        payload={'name':'Next lesson','mode':'topic','inputs':{'topic':'Learning'},'targetLanguages':['英语'],'memoryCourseId':self.task['id'],'autoExecute':False}
        created=self.app.add_task(payload,self.uid('owner'))
        self.assertEqual(len(created['courseMemory']['entries']),1)
        self.assertIn('neuron',memory_prompt(generation_memory(created),'S2'))
        self.call('memory',{'revision':1,'entries':[]})
        self.assertEqual(self.app.get_task(created['id'])['courseMemory']['revision'],1)
        self.assertEqual(self.call('memory',{'revision':0,'entries':entries})[0],400)
        with self.assertRaisesRegex(ValueError,'无权'):
            self.app.add_task(payload,self.uid('other'))

    def test_localized_revision_preserves_original_and_unchanged_media(self):
        before={str(p.relative_to(self.course_root)):p.read_bytes() for p in self.course_root.rglob('*') if p.is_file()}
        change={'page':2,'language':'en-US','kind':'script','text':'Training improves predictions by adjusting model weights based on their errors.'}
        status,result=self.call('revise',{'changes':[change],'execute':False});self.assertEqual(status,200,result)
        self.assertEqual(result['startStage'],3)
        root=self.app.studio.root(result['task']['id'])
        updated=json.loads((root/'s2/en-US/translated_structure.json').read_text())
        self.assertEqual(updated['pages'][1]['script_text'],change['text'])
        self.assertEqual((root/'runtime/en-US/s3_output/segments/page_001.wav').read_bytes(),before['runtime/en-US/s3_output/segments/page_001.wav'])
        for name,content in before.items():self.assertEqual((self.course_root/name).read_bytes(),content,name)
        previous=json.loads((root/'s3/en-US/narration_structure.json').read_text())
        oral=preserve_revision_narration(updated,previous,[change],'en-US')
        self.assertEqual(oral['pages'][0]['script_text'],previous['pages'][0]['script_text'])
        self.assertEqual(oral['pages'][1]['script_text'],change['text'])
        report=json.loads((root/'s2/en-US/translation_report.json').read_text())
        self.assertTrue(report['pages'][1]['revisionPendingReview'])
        self.assertTrue(report['pages'][0]['passed'])
        compared=self.app.studio.compare(result['task'],self.task)
        self.assertEqual([x['page'] for x in compared['changes']],[2])

    def test_subtitle_and_layout_changes_start_at_composition(self):
        for change in [{'page':2,'language':'en-US','kind':'subtitle','oldText':'adjusts','text':'updates'},
                       {'page':1,'language':'en-US','kind':'layout'}]:
            result=create_revision(self.app,self.task,self.uid('owner'),[change],execute=False)
            self.assertEqual(result['startStage'],5)
            root=self.app.studio.root(result['task']['id'])
            if change['kind']=='subtitle':
                self.assertIn('updates',(root/'runtime/en-US/s3_output/subtitles.srt').read_text())
                self.assertIn('adjusts',(self.course_root/'runtime/en-US/s3_output/subtitles.srt').read_text())
            else:
                bbox=json.loads((root/'runtime/en-US/s4_output/bbox.json').read_text())['bboxes']
                self.assertEqual(bbox[0]['layout_mode'],'side_panel');self.assertNotIn('layout_mode',bbox[1])

    def test_display_title_revision_keeps_speech_and_updates_downloadable_slides(self):
        memory=self.app.studio.update_memory(self.task['id'],self.uid('owner'),{'revision':0,'entries':[
            {'kind':'term','name':'学习','value':'Model Training','language':'en-US','pages':[2],'status':'accepted'}]})
        change={'page':2,'language':'en-US','kind':'slide_text','oldText':'Learning','text':'Model Training'}
        result=create_revision(self.app,self.task,self.uid('owner'),[change],execute=False)
        self.assertEqual(result['startStage'],5)
        root=self.app.studio.root(result['task']['id'])
        self.assertIn('s4',result['task']['outputs'])
        self.assertEqual(self.app.studio.memory(result['task']['id']),memory)
        self.assertEqual((root/'runtime/en-US/s3_output/segments/page_002.wav').read_bytes(),
                         (self.course_root/'runtime/en-US/s3_output/segments/page_002.wav').read_bytes())
        self.assertEqual(Presentation(root/'s3/en-US/course-with-narration.pptx').slides[1].shapes.title.text,'Model Training')
        self.assertEqual(Presentation(self.course_root/'s2/en-US/translated.pptx').slides[1].shapes.title.text,'Learning')
        with zipfile.ZipFile(self.course_root/'s2/en-US/translated.pptx') as before, zipfile.ZipFile(root/'s2/en-US/translated.pptx') as after:
            for name in before.namelist():
                if name!='ppt/slides/slide2.xml':self.assertEqual(before.read(name),after.read(name),name)

    def test_slide_text_spoken_in_narration_still_requires_speech_stage(self):
        change={'page':2,'language':'en-US','kind':'slide_text','oldText':'weights','text':'parameters'}
        result=create_revision(self.app,self.task,self.uid('owner'),[change],execute=False)
        self.assertEqual(result['startStage'],3)

    def test_export_has_manifest_all_resources_and_no_private_state(self):
        status,value=self.call('export',{});self.assertEqual(status,200,value)
        with zipfile.ZipFile(self.app.settings.artifact_root/value['download'].removeprefix('/artifacts/')) as archive:
            manifest=ET.fromstring(archive.read('imsmanifest.xml'))
            self.assertEqual(manifest.find('{*}metadata/{*}schemaversion').text,'1.2')
            for node in manifest.findall('.//{*}file'):self.assertIn(node.get('href'),archive.namelist())
            self.assertNotIn('state.json',archive.namelist());self.assertNotIn('execution-journal.json',archive.namelist())
            self.assertIn(b'LMSInitialize',archive.read('index.html'))

    def test_review_pause_keeps_completed_language_stages_and_allows_draft_approval(self):
        from app.full_pipeline import StageReviewRequired
        task=copy.deepcopy(self.task)
        task['inputs'].update(reviewBeforeMedia=True,reviewApproved=False)
        runner=self.app.pipeline_runner
        runner._run_s3=Mock()
        with self.assertRaises(StageReviewRequired):
            runner.run(task['id'],task,lambda *_:None,start_stage=3)
        runner._run_s3.assert_not_called()
        self.app.mutate_task(task['id'],lambda value,_:value.update(status='waiting_review',currentStage=3,inputs=task['inputs']))
        status,_=self.call('draft',{'page':2,'language':'en-US','text':'Training adjusts the model weights using the observed prediction errors.','revision':0})
        self.assertEqual(status,200)
        self.assertEqual(self.call('draft',{'page':2,'language':'en-US','text':'A stale edit.','revision':0})[0],400)
        self.assertEqual(self.call('approve',{})[0],200)
        self.app.execute_task.assert_called_once_with(task['id'],start_stage=3,asynchronous=True)
        self.assertTrue(self.app.get_task(task['id'])['inputs']['reviewApproved'])

    def test_membership_removal_and_publication_apply_to_the_course_version_family(self):
        self.member('learner','learner');self.call('publish',{'status':'published'})
        result=create_revision(self.app,self.task,self.uid('owner'),[{'page':1,'language':'en-US','kind':'layout'}],execute=False)
        version=result['task']['id']
        self.app.mutate_task(version,lambda value,_:value.update(status='completed'))
        path='/v1/studio/courses/'+version
        self.assertEqual(self.call(path+'/publish',{'status':'published'})[0],200)
        self.assertEqual(self.call(user='learner')[0],403)
        self.assertEqual(self.call(path,user='learner')[0],200)
        self.call('members',{'userId':self.uid('learner'),'role':'remove'})
        self.assertEqual(self.call(path,user='learner')[0],403)

    def test_feedback_repair_retries_locally_creates_one_version_and_awaits_review(self):
        feedback=self.feedback()
        model=self.app.pipeline_runner.gateway.respond_json
        model.side_effect=[{'script_text':'这是中文内容，模型需要将讲稿改成课程要求的英语表达。这里的文字足够长，可以检查输出语言是否正确。','replacements':[]},
                           {'script_text':'Training improves model predictions by adjusting the weights according to prediction errors.','replacements':[]}]
        with patch('app.feedback_repair.start_worker'):
            record=submit_repair(self.app,self.task,self.uid('owner'),{'feedbackId':feedback['id']})
            again=submit_repair(self.app,self.task,self.uid('owner'),{'feedbackId':feedback['id']})
        self.assertEqual(record['id'],again['id'])
        _work(self.app,record['id'])
        saved=self.app.studio.state()['repairs'][record['id']]
        self.assertEqual(saved['attempts'],2);self.assertEqual(model.call_count,2)
        self.app.execute_task.assert_called_once_with(saved['releaseTaskId'],start_stage=3,asynchronous=True)
        self.app.mutate_task(saved['releaseTaskId'],lambda task,_:task.update(status='completed'))
        self.assertEqual(self.app.studio.repair_status(saved)['status'],'ready_for_review')
        self.assertNotEqual(self.app.store.snapshot()['userFeedback'][feedback['id']]['status'],'resolved')
        status,_=self.call('resolve-feedback',{'feedbackId':feedback['id'],'note':'Retain consistent teaching terminology.','remember':True})
        self.assertEqual(status,200);self.assertEqual(self.app.studio.memory(self.task['id'])['entries'][0]['status'],'accepted')

    def test_model_review_and_memory_candidates_do_not_claim_uncited_facts_passed(self):
        model=self.app.pipeline_runner.gateway.respond_json
        model.return_value={'page':1,'issues':[{'quote':'Nonexistent source','message':'Invented quote'}]}
        status,result=self.call('review',{'page':1,'withImages':False});self.assertEqual(status,200,result)
        self.assertEqual(result['status'],'failed');self.assertEqual(model.call_count,4)
        model.return_value={'page':1,'issues':[],'summary':'Text-only review.'}
        status,result=self.call('review',{'page':1,'withImages':False});self.assertEqual(result['status'],'no_issues_reported');self.assertFalse(result['imageReviewed'])
        model.return_value={'entries':[{'kind':'definition','name':'神经元','value':'输入与权重共同决定输出','pages':[1],'quote':'输入与权重'}]}
        status,value=self.call('extract-memory',{});self.assertEqual(status,200,value)
        self.assertEqual(value['entries'][0]['status'],'candidate')
        overview=self.call()[1];self.assertTrue(any(x['status']=='not_checked' for x in overview['quality']['checks']))

    def test_plan_template_batch_and_member_permissions_use_persistent_server_state(self):
        self.app.mutate_task(self.task['id'],lambda value,_:value['assets'].update({
            'consent':{'authorized':True,'version':POLICY_VERSION},
            'referenceAudio':{'path':self.url(self.course_root/'runtime/en-US/s3_output/segments/page_001.wav'),'name':'reference.wav'}}))
        good={'outline':[{'title':'Inputs','objective':'Explain input values','activity':'Identify input data','minutes':2,'requirements':'Use source definitions'}, {'title':'Weights','objective':'Explain model weights','activity':'Compare two weights','minutes':3,'requirements':''}]}
        wrong=copy.deepcopy(good);wrong['outline'][1]['title']='Inputs'
        self.app.pipeline_runner.gateway.respond_json.side_effect=[wrong,good]
        status,plan=self.call('/v1/studio/plans',{'topic':'Neural networks','lessons':2,'sourceTaskId':self.task['id']})
        self.assertEqual(status,201,plan);self.assertNotEqual(plan['outline'][0]['title'],plan['outline'][1]['title'])
        self.assertEqual(self.app.pipeline_runner.gateway.respond_json.call_count,2)
        self.assertEqual(self.call('/v1/studio/plans',user='other')[1],[])
        status,template=self.call('/v1/studio/templates',{'courseId':self.task['id'],'name':'Teacher template'});self.assertEqual(status,201,template)
        status,batch=self.call('/v1/studio/plans/'+plan['id']+'/launch',{'targetLanguages':['en-US'],'templateId':template['id'],'execute':True})
        self.assertEqual(status,200,batch);self.assertEqual(len(batch['taskIds']),2);self.assertEqual(self.app.execute_task.call_count,2)
        copied=self.app.get_task(batch['taskIds'][0])['assets']
        self.assertTrue(copied['consent']['authorized'])
        self.assertIn(batch['taskIds'][0],copied['referenceAudio']['path'])
        self.assertEqual(self.call('/v1/studio/templates/'+template['id']+'/launch',{'topics':['x']},user='other')[0],403)

    def test_learning_questions_progress_and_answers_are_account_scoped(self):
        self.member('learner','learner');self.call('publish',{'status':'published'})
        question=create_question(self.app.settings,self.task['id'],{'locale':'en-US','prompt':'Which item changes during training?',
            'options':[{'id':'a','text':'Weights'},{'id':'b','text':'Slide numbers'}],'correct_option_id':'a','anchor':{'page_id':2}},self.uid('owner'))
        review_question(self.app.settings,question['question_id'],self.uid('owner'),True)
        self.assertEqual(self.call('/v1/questions/'+question['question_id']+'/attempt',{'option_id':'b'},user='other')[0],403)
        self.assertEqual(self.call('/v1/questions/'+question['question_id']+'/attempt',{'option_id':'b'},user='learner')[0],200)
        status,_=self.call('learning',{'page':2,'seconds':2.2,'preference':'Explain the terms first.'},user='learner');self.assertEqual(status,200)
        valid={'answer':'Model weights are adjusted during training to reduce prediction errors.','sources':[{'page':2,'quote':'Training adjusts model weights'}],'insufficient':False}
        invalid={**valid,'answer':'模型通过不断调整权重来降低预测误差。这是模型训练中非常重要的一步，需要根据实际的训练数据进行多次迭代。'}
        self.app.pipeline_runner.gateway.respond_json.side_effect=[invalid,valid]
        status,value=self.call('ask',{'page':2,'language':'en-US','question':'What changes during training?'},user='learner');self.assertEqual(status,200,value)
        learner=self.call(user='learner')[1]['learning'];self.assertEqual(learner['reviewPages'],[2]);self.assertEqual(len(learner['answers']),1)
        self.assertEqual(learner['progress']['positions']['en-US']['page'],2)
        self.assertEqual(self.call()[1]['learning']['answers'],[])
        self.assertNotIn('correct_option_id',learner['questions'][0])

    def test_failed_feedback_repair_is_bounded_and_preserves_the_original_course(self):
        feedback=self.feedback()
        self.app.pipeline_runner.gateway.respond_json.return_value={'script_text':'','replacements':[]}
        record=submit_repair(self.app,self.task,self.uid('owner'),{'feedbackId':feedback['id']})
        deadline=time.monotonic()+4
        while time.monotonic()<deadline:
            row=self.app.studio.state()['repairs'][record['id']]
            if row['status']=='failed':break
            time.sleep(.02)
        self.assertEqual(row['status'],'failed');self.assertEqual(row['attempts'],4)
        self.assertEqual(self.app.pipeline_runner.gateway.respond_json.call_count,4)
        self.assertEqual(self.app.get_task(self.task['id'])['status'],'completed')
        self.assertNotIn('releaseTaskId',row)


if __name__=='__main__':unittest.main()

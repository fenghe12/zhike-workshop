import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from pptx import Presentation
from app.config import Settings
from app.full_pipeline import FullPipelineRunner


class TopicSourceLanguageTests(unittest.TestCase):
    def test_source_language_does_not_follow_first_translation_target(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            settings = Settings(root/'data', root/'artifacts', 'http://unused', 'test', '', '127.0.0.1', 0)
            runner = FullPipelineRunner(settings)
            task = {'id': 'source-language-case', 'name': '源语言测试', 'mode': 'topic',
                    'sourceLanguage': '中文', 'targetLanguages': ['en-US', 'ja-JP'],
                    'inputs': {'topic': '输入数据', 'totalPages': 1, 'durationMinutes': 1}}
            requested = []

            def generate(topic, destination, **options):
                requested.append(options['language'])
                deck = Presentation()
                slide = deck.slides.add_slide(deck.slide_layouts[1])
                slide.shapes.title.text = '输入数据'
                slide.placeholders[1].text = '输入是待处理的数据。'
                deck.save(destination)
                structure = destination.parent/'generated-course-structure.json'
                structure.write_text(json.dumps({'pages': [{'page_id': 1,
                    'script_text': '输入是待处理的数据，权重和偏置是模型参数。',
                    'script_source': 'course_generation_agent'}]}))
                return {'engine': 'test', 'course_structure': str(structure)}

            with patch('app.deeppresenter_course.generate_course', side_effect=generate):
                runner._run_s1(task, runner._task_root(task['id']), lambda *args: None)
            self.assertEqual(requested, ['Simplified Chinese'])
            saved = json.loads((runner._task_root(task['id'])/'s1/course_structure.json').read_text())
            self.assertEqual(saved['pages'][0]['title'], '输入数据')


if __name__ == '__main__':
    unittest.main()

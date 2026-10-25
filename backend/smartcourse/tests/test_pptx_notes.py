from pathlib import Path
import tempfile
import unittest
import zipfile

from pptx import Presentation
from app.pptx_notes import write_speaker_notes


class SpeakerNotesTests(unittest.TestCase):
    def test_preserves_original_part_names_after_deleting_a_slide(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / 'selected-pages.pptx'
            deck = Presentation()
            deck.slides.add_slide(deck.slide_layouts[5]).shapes.title.text = 'First page'
            deck.slides.add_slide(deck.slide_layouts[5]).shapes.title.text = 'Retained page'
            removed = deck.slides._sldIdLst[0]
            deck.part.drop_rel(removed.rId)
            deck.slides._sldIdLst.remove(removed)
            deck.save(path)
            with zipfile.ZipFile(path) as source:
                original = source.read('ppt/slides/slide2.xml')
            self.assertEqual(write_speaker_notes(path, [{'page_id': 1, 'script_text': 'The retained page narration.'}]), 1)
            with zipfile.ZipFile(path) as result:
                self.assertEqual(result.read('ppt/slides/slide2.xml'), original)
            restored = Presentation(path)
            self.assertEqual(restored.slides[0].shapes.title.text, 'Retained page')
            self.assertEqual(restored.slides[0].notes_slide.notes_text_frame.text, 'The retained page narration.')

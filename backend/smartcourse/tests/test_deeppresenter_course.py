import json
import os
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

from app.deeppresenter_course import generate_course


class DeepPresenterCourseTests(unittest.TestCase):
    def test_restricted_course_material_reaches_remote_adapter(self):
        material = {'currentPage': {'script_text': '只解释输入、权重和阈值'}}
        with patch.dict(os.environ, {'SMARTCOURSE_REMOTE_MEDIA': '1'}), patch(
                'app.remote_deeppresenter.generate_remote_course', return_value={}) as remote:
            generate_course('课程回答', Path('/tmp/answer.pptx'), pages=1, audience='本科生',
                            minutes=1, language='Chinese', source_material=material)
            self.assertEqual(remote.call_args.kwargs['source_material'], material)

    def test_remote_mode_uses_configured_adapter(self):
        with patch.dict(os.environ, {'SMARTCOURSE_REMOTE_MEDIA': '1'}), patch(
                'app.remote_deeppresenter.generate_remote_course', return_value={'engine': 'DeepPresenter-9B'}) as remote:
            result = generate_course('主题', Path('/tmp/course.pptx'), pages=3,
                                     audience='本科生', minutes=3, language='Chinese')
            self.assertEqual(result['engine'], 'DeepPresenter-9B')
            self.assertEqual(remote.call_args.kwargs['pages'], 3)
            self.assertEqual(remote.call_args.args[0], '主题')
            self.assertNotIn('source_material', remote.call_args.kwargs)

    def test_actual_generator_output_is_used(self):
        from pptx import Presentation
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            runtime = root / "pptagent-demo"
            runtime.mkdir()
            for name in ("run_branch_pipeline.py", "config.official-huawei9b.yaml",
                         "mcp.branch-huawei.json", "render_office.py"):
                (runtime / name).touch()
            (runtime / 'mcp.branch-huawei.json').write_text('{}')
            destination = root / "task/generated.pptx"
            def execute(command, **kwargs):
                self.assertEqual(command[1], str(runtime / "run_branch_pipeline.py"))
                work = Path(command[command.index("--output-dir") + 1])
                deck = Presentation()
                deck.slides.add_slide(deck.slide_layouts[6])
                deck.save(work / "branch-final.pptx")
                (work / "branch-pipeline.json").write_text(json.dumps({
                    "final_pptx": str(work / "branch-final.pptx")}))
                return SimpleNamespace(returncode=0)
            with patch.dict(os.environ, {"SMARTCOURSE_ROOT": str(root)}), patch(
                    "app.deeppresenter_course.subprocess.run", side_effect=execute):
                result = generate_course("Neurons", destination, pages=1,
                                         audience="Students", minutes=1, language="English")
            self.assertEqual(result["engine"], "DeepPresenter-9B")
            self.assertTrue(destination.is_file())

    def test_missing_runtime_does_not_generate_a_substitute(self):
        with tempfile.TemporaryDirectory() as directory:
            destination = Path(directory) / "generated.pptx"
            with patch.dict(os.environ, {"SMARTCOURSE_ROOT": directory}):
                with self.assertRaises(RuntimeError):
                    generate_course("Neurons", destination, pages=1,
                                    audience="Students", minutes=1, language="English")
            self.assertFalse(destination.exists())

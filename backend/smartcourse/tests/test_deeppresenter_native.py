import json
import os
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from pptx import Presentation
from app.deeppresenter_native import generate_native,managed_model


class NativeDeepPresenterTests(unittest.TestCase):
    def test_reference_lesson_count_does_not_change_requested_page_count(self):
        page={'page_id':1,'title':'关闭不用的水龙头','key_points':['刷牙时关闭水龙头','发现漏水及时报修'],
              'script_text':'刷牙时关闭不用的水龙头，可以减少浪费。发现漏水时及时报修。'}
        reference={'pages':[{'title':'水循环'},{'title':'节约用水'}]}
        wrong={**page,'page_id':2}
        with tempfile.TemporaryDirectory() as folder,patch.dict(os.environ,{'SMARTCOURSE_REMOTE_DEEPPRESENTER_CONTROL':''}),patch('app.gateway.ResponsesGateway.respond_json',side_effect=[wrong,page]) as model:
            output=Path(folder)/'course.pptx'
            generate_native('节约用水',output,pages=1,audience='大学生',minutes=1,
                            language='Simplified Chinese',source_material=reference)
            self.assertEqual(len(Presentation(output).slides),1)
            self.assertEqual(model.call_count,2)
            self.assertIn('"current_page": {"title": "节约用水"',model.call_args_list[0].args[0])
            self.assertIn('Received page_id=2',model.call_args_list[1].args[0])

    def test_native_model_pages_are_rendered_with_scripts_and_reused_by_request(self):
        outline={'pages':[{'title':'二进制位值'},{'title':'逻辑运算'}]}
        first={'page_id':1,'title':'二进制位值','key_points':['零和一','从右向左位值翻倍'],
               'script_text':'二进制只使用零和一。从右向左，各位的位值依次为一、二、四、八。'}
        second={'page_id':2,'title':'逻辑运算','key_points':['逻辑与','逻辑或'],
                'script_text':'逻辑与要求两个条件同时成立。逻辑或在至少一个条件成立时成立。'}
        with tempfile.TemporaryDirectory() as folder,patch.dict(os.environ,{'SMARTCOURSE_REMOTE_DEEPPRESENTER_CONTROL':''}),patch('app.gateway.ResponsesGateway.respond_json',side_effect=[outline,first,second]) as model:
            output=Path(folder)/'course.pptx'
            result=generate_native('二进制',output,pages=2,audience='大学生',minutes=1,language='Simplified Chinese')
            self.assertEqual(result['engine'],'DeepPresenter-9B')
            self.assertEqual(result['generator'],'SmartCourse native model adapter')
            self.assertEqual(len(Presentation(output).slides),2)
            structure=json.loads(Path(result['course_structure']).read_text())
            self.assertEqual(structure['pages'][1]['script_text'],second['script_text'])
            self.assertEqual(model.call_count,3)
            generate_native('二进制',output,pages=2,audience='大学生',minutes=1,language='Simplified Chinese')
            self.assertEqual(model.call_count,3)

    def test_owned_model_is_stopped_when_generation_raises(self):
        import subprocess
        with patch.dict(os.environ,{'SMARTCOURSE_REMOTE_DEEPPRESENTER_CONTROL':'/persistent/model_control.py','SMARTCOURSE_REMOTE_PYTHON':'/persistent/python'}),patch('app.remote_media._connection',return_value=([],22,'worker')),patch('app.deeppresenter_native.subprocess.run',return_value=subprocess.CompletedProcess([],0,'','')) as invoke:
            with self.assertRaisesRegex(ValueError,'generation failed'):
                with managed_model():raise ValueError('generation failed')
            self.assertEqual(invoke.call_count,2)
            self.assertTrue(invoke.call_args_list[0].args[0][-1].endswith(' start'))
            self.assertTrue(invoke.call_args_list[1].args[0][-1].endswith(' stop'))


if __name__=='__main__':unittest.main()

import tempfile
import unittest
import zipfile
from pathlib import Path
from pptx import Presentation
from app.course_preflight import validate_course_files


class CoursePreflightTests(unittest.TestCase):
    def test_page_mismatch_is_rejected(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            deck = Presentation()
            for _ in range(2):
                deck.slides.add_slide(deck.slide_layouts[6])
            deck.save(root / 'course.pptx')
            with zipfile.ZipFile(root / 'lecture.docx', 'w') as archive:
                archive.writestr('word/document.xml', '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:p><w:r><w:t>第1页 只有第一页讲稿</w:t></w:r></w:p></w:document>')
            inputs = {'ppt': {'path': '/artifacts/course.pptx'}, 'script': {'path': '/artifacts/lecture.docx'}}
            with self.assertRaisesRegex(ValueError, '第 1 至 2 页'):
                validate_course_files(inputs, root)

    def test_external_path_is_rejected(self):
        with tempfile.TemporaryDirectory() as directory:
            with self.assertRaises(ValueError):
                validate_course_files({'ppt': {'path': '/etc/passwd'}}, Path(directory))

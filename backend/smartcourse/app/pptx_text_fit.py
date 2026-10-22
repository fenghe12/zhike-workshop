"""Fit translated text within the original text frames."""
import math
import os
import posixpath
import subprocess
import tempfile
import zipfile
from collections import Counter
from pathlib import Path

import pymupdf
from lxml import etree as E

A = 'http://schemas.openxmlformats.org/drawingml/2006/main'
P = 'http://schemas.openxmlformats.org/presentationml/2006/main'
R = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
MC = 'http://schemas.openxmlformats.org/markup-compatibility/2006'
NS = {'a': A, 'p': P}


def normalized(value):
    return ''.join(c for c in value.casefold() if c.isalnum())


def text_bounds(page, value):
    text = normalized(value)
    if len(text) < 8:
        return None
    matches = []
    covered = 0
    cursor = 0
    for block in page.get_text('dict')['blocks']:
        for line in block.get('lines', []):
            line_text = normalized(''.join(span['text'] for span in line['spans']))
            position = text.find(line_text, cursor) if line_text else -1
            if position >= cursor and position <= cursor + max(1, len(text)*.025):
                matches.append(line['bbox'])
                covered += len(line_text)
                cursor = position + len(line_text)
    if not matches or covered < len(text)*.9:
        return None
    return (min(b[0] for b in matches), min(b[1] for b in matches),
            max(b[2] for b in matches), max(b[3] for b in matches))


def slide_parts(data):
    root = E.fromstring(data['ppt/presentation.xml'])
    relationships = {r.get('Id'): r.get('Target') for r in E.fromstring(data['ppt/_rels/presentation.xml.rels'])}
    return [posixpath.normpath(posixpath.join('ppt', relationships[node.get('{'+R+'}id')]))
            if not relationships[node.get('{'+R+'}id')].startswith('/') else relationships[node.get('{'+R+'}id')].lstrip('/')
            for node in root.find('{'+P+'}sldIdLst')]


def dominant_font_size(page, value):
    text = normalized(value)
    sizes = Counter()
    for block in page.get_text('dict')['blocks']:
        for line in block.get('lines', []):
            for span in line['spans']:
                content = normalized(span['text'])
                if content and content in text:
                    sizes[round(span['size'],2)] += len(content)
    return sizes.most_common(1)[0][0] if sizes else None


def fit_from_pdfs(source, target, source_pdf, target_pdf):
    from pptx import Presentation
    presentation = Presentation(source)
    with zipfile.ZipFile(source) as z:
        before = {n:z.read(n) for n in z.namelist()}
    with zipfile.ZipFile(target) as z:
        after = {n:z.read(n) for n in z.namelist()}
    changes = []
    with pymupdf.open(source_pdf) as original, pymupdf.open(target_pdf) as translated:
        for index, part in enumerate(slide_parts(before)):
            frame_bounds = {}
            for item in presentation.slides[index].shapes:
                if item.has_text_frame:
                    frame = item.text_frame
                    frame_bounds[str(item.shape_id)] = (
                        (item.left + frame.margin_left) / 12700,
                        (item.top + frame.margin_top) / 12700,
                        (item.left + item.width - frame.margin_right) / 12700,
                        (item.top + item.height - frame.margin_bottom) / 12700,
                    )
            source_root, target_root = E.fromstring(before[part]), E.fromstring(after[part])
            source_shapes = source_root.findall('.//p:sp', NS)
            target_shapes = target_root.findall('.//p:sp', NS)
            source_map = {s.find('p:nvSpPr/p:cNvPr',NS).get('id'):s for s in source_shapes}
            for shape in target_shapes:
                if any(parent.tag == '{'+MC+'}AlternateContent' for parent in shape.iterancestors()):
                    continue
                identifier = shape.find('p:nvSpPr/p:cNvPr', NS).get('id')
                source_shape = source_map.get(identifier)
                if source_shape is None:
                    continue
                old_text = ''.join(source_shape.xpath('.//a:t/text()',namespaces=NS))
                new_text = ''.join(shape.xpath('.//a:t/text()',namespaces=NS))
                if old_text == new_text:
                    continue
                old_box, new_box = text_bounds(original[index],old_text), text_bounds(translated[index],new_text)
                if not old_box or not new_box:
                    continue
                available = frame_bounds.get(identifier)
                if available:
                    if (new_box[0] >= available[0]-2 and new_box[1] >= available[1]-2
                            and new_box[2] <= available[2]+2 and new_box[3] <= available[3]+2):
                        continue
                    height_ratio = (available[3]-available[1])/max(1,new_box[3]-new_box[1])
                    width_ratio = (available[2]-available[0])/max(1,new_box[2]-new_box[0])
                    scale = min(1, height_ratio, width_ratio)
                else:
                    height_ratio = (old_box[3]-old_box[1])/max(1,new_box[3]-new_box[1])
                    width_ratio = (old_box[2]-old_box[0])/max(1,new_box[2]-new_box[0])
                    scale = min(1, math.sqrt(height_ratio), width_ratio)
                if scale > .995:
                    continue
                scale = max(.5, scale*.98)
                body = shape.find('p:txBody/a:bodyPr',NS)
                if body is None:
                    continue
                current = source_shape.find('p:txBody/a:bodyPr/a:normAutofit',NS)
                previous_scale = int(current.get('fontScale','100000'))/100000 if current is not None else 1
                font_size = dominant_font_size(original[index], old_text)
                if font_size is None:
                    font_size = dominant_font_size(translated[index], new_text)
                if font_size is None:
                    continue
                # Explicit sizes also work in viewers that ignore placeholder fontScale.
                source_runs = source_shape.xpath('p:txBody/a:p/a:r | p:txBody/a:p/a:fld', namespaces=NS)
                run_index = 0
                for paragraph in shape.findall('p:txBody/a:p', NS):
                    for run in paragraph:
                        if E.QName(run).localname not in {'r','fld'}:
                            continue
                        properties = run.find('a:rPr', NS)
                        if properties is None:
                            properties = E.Element('{'+A+'}rPr');run.insert(0,properties)
                        source_properties = source_runs[run_index].find('a:rPr', NS) if run_index < len(source_runs) else None
                        original_size = source_properties.get('sz') if source_properties is not None else None
                        base_size = int(original_size)*previous_scale if original_size else font_size*100
                        properties.set('sz',str(round(base_size*scale)))
                        run_index += 1
                    ending = paragraph.find('a:endParaRPr', NS)
                    if ending is not None:
                        ending.set('sz',str(round(font_size*scale*100)))
                for child in list(body):
                    if E.QName(child).localname in {'noAutofit','spAutoFit','normAutofit'}:
                        body.remove(child)
                E.SubElement(body,'{'+A+'}noAutofit')
                changes.append({'page':index+1,'shape':identifier,'fontScale':round(previous_scale*scale,4),
                                'sourceBounds':old_box,'translatedBoundsBefore':new_box,'availableBounds':available})
            after[part] = E.tostring(target_root,encoding='UTF-8',xml_declaration=True)
    temporary = Path(target).with_suffix('.fit-pending.pptx')
    with zipfile.ZipFile(temporary,'w',zipfile.ZIP_DEFLATED) as z:
        for name,content in after.items():
            z.writestr(name,content)
    temporary.replace(target)
    return changes


def fit_translated_text(source, target):
    office = os.environ.get('SMARTCOURSE_SOFFICE')
    if not office:
        raise RuntimeError('需要配置课件渲染服务才能校验译文排版')
    with tempfile.TemporaryDirectory(prefix='ppt-text-fit-') as temporary:
        root = Path(temporary)
        pdfs=[]
        for index,path in enumerate([source,target]):
            folder=root/str(index);folder.mkdir()
            render_path = Path(path).resolve()
            if index == 1:
                # A taller measurement canvas exposes text clipped by the slide edge.
                render_path = folder / 'measurement.pptx'
                with zipfile.ZipFile(path) as original, zipfile.ZipFile(render_path, 'w', zipfile.ZIP_DEFLATED) as expanded:
                    for name in original.namelist():
                        content = original.read(name)
                        if name == 'ppt/presentation.xml':
                            tree = E.fromstring(content)
                            size = tree.find('{'+P+'}sldSz')
                            size.set('cy', str(min(51206400, int(size.get('cy')) * 4)))
                            content = E.tostring(tree, encoding='UTF-8', xml_declaration=True)
                        expanded.writestr(name, content)
            subprocess.run([office,'-env:UserInstallation='+(folder/'profile').as_uri(),
                            '--headless','--convert-to','pdf','--outdir',str(folder),str(render_path)],
                           check=True,capture_output=True,timeout=180)
            pdfs.append(folder/(render_path.stem+'.pdf'))
        return fit_from_pdfs(source,target,*pdfs)


def fit_visual_collisions(target, target_pdf):
    """Move bottom images below translated text when rendering reveals an intersection."""
    with zipfile.ZipFile(target) as archive:
        data = {name: archive.read(name) for name in archive.namelist()}
    changed = []
    with pymupdf.open(target_pdf) as document:
        for slide_index, part in enumerate(slide_parts(data)):
            page = document[slide_index]
            blocks = page.get_text('dict').get('blocks', [])
            text_boxes = [pymupdf.Rect(block['bbox']) for block in blocks if block.get('type') == 0]
            image_boxes = [pymupdf.Rect(block['bbox']) for block in blocks if block.get('type') == 1]
            if not image_boxes or not text_boxes:
                continue
            root = E.fromstring(data[part])
            pictures = root.findall('.//p:pic', NS)
            for picture in pictures:
                xfrm = picture.find('p:spPr/a:xfrm', NS)
                if xfrm is None:
                    continue
                off = xfrm.find('a:off', NS); ext = xfrm.find('a:ext', NS)
                if off is None or ext is None:
                    continue
                current = pymupdf.Rect(
                    int(off.get('x')) / 12700, int(off.get('y')) / 12700,
                    (int(off.get('x')) + int(ext.get('cx'))) / 12700,
                    (int(off.get('y')) + int(ext.get('cy'))) / 12700,
                )
                image = next((box for box in image_boxes if box.intersects(current)), None)
                if image is None:
                    continue
                collisions = [box for box in text_boxes if box.intersects(image)
                              and (box & image).get_area() > min(box.get_area(), image.get_area()) * .08]
                if not collisions:
                    continue
                text_bottom = max(box.y1 for box in collisions)
                new_y = max(current.y0, text_bottom + 8)
                available = page.rect.height - new_y - 6
                new_height = min(current.height, max(36, available))
                if new_y <= current.y0 + 1 and new_height >= current.height - 1:
                    continue
                off.set('y', str(round(new_y * 12700)))
                ext.set('cy', str(round(new_height * 12700)))
                changed.append({'page': slide_index + 1, 'image': picture.find('p:nvPicPr/p:cNvPr', NS).get('id'),
                                'oldY': current.y0, 'newY': new_y, 'oldHeight': current.height,
                                'newHeight': new_height})
            data[part] = E.tostring(root, encoding='UTF-8', xml_declaration=True)
    if changed:
        pending = Path(target).with_suffix('.collision-fit-pending.pptx')
        with zipfile.ZipFile(pending, 'w', zipfile.ZIP_DEFLATED) as archive:
            for name, content in data.items():
                archive.writestr(name, content)
        pending.replace(target)
    return changed

from __future__ import annotations

import copy
import json
import os
import re
import shutil
import sys
import time
from decimal import Decimal
from pathlib import Path
from statistics import mean
from typing import Any, Callable

from .adaptive_teaching import apply_adaptive_policy, build_adaptive_experiment
from .advanced_experiments import (
    build_personalization_plan,
    check_pedagogical_alignment,
    prescreen_candidate,
    select_proactive_intervention,
    update_structured_memory,
)
from .config import SUPPORTED_LANGUAGES, Settings
from .translation_policy import validate_policy, meets_thresholds
from .course_agent import CourseStructureAgent, save_structure
from .code_lessons import get_lesson_spec
from .gateway import ResponsesGateway, GatewayError
from .pedagogical_ir import build_e2_variants, compile_course_structure_ir
from .performance import PipelinePerformanceRecorder
from .resource_scheduler import ResourceScheduler
from .rehearsal_analyzer import analyze_rehearsal
from .teacher_avatar import (
    build_teacher_profile,
    longcat_engine_readiness,
    schedule_gestures,
    validate_profile_files,
)


ProgressCallback = Callable[[int, str, int, dict[str, Any]], None]


LANGUAGE_MODEL_NAMES = {
    "zh-CN": "Simplified Chinese",
    "en-US": "English",
    "ja-JP": "Japanese",
    "ko-KR": "Korean",
    "de-DE": "German",
    "fr-FR": "French",
    "ru-RU": "Russian",
    "es-ES": "Spanish",
}


LANGUAGE_NAMES = {
    "zh": "Simplified Chinese",
    "en": "English",
    "ja": "Japanese",
    "ko": "Korean",
    "de": "German",
    "fr": "French",
    "ru": "Russian",
    "es": "Spanish",
}


def _write_json(path: Path, value: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_name('.'+path.name+'.pending')
    temporary.write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding="utf-8")
    temporary.replace(path)


def _language_code(name: str) -> str:
    for code, chinese_name in SUPPORTED_LANGUAGES.items():
        if name in {code, chinese_name}:
            return code
    raise ValueError(f"不支持的目标语言：{name}")


def _safe_task_component(value: str) -> str:
    return re.sub(r"[^0-9A-Za-z_-]+", "_", value).strip("_") or "task"


def _translation_warning(page_reports: list[dict[str, Any]]) -> str:
    failed_pages = [item["page_id"] for item in page_reports if not item["passed"]]
    return ("以下页面的翻译检查未通过，已继续生成，请审阅：" + "、".join(map(str, failed_pages))
            if failed_pages else "")


def _run_speech_with_retries(function: Callable[[], Any], notify: Callable[[Exception, int], None], error_type: type[Exception]) -> Any:
    for attempt in range(4):
        try:
            return function()
        except error_type as exc:
            if attempt == 3:
                raise
            notify(exc, attempt + 1)


def _exact_token_accuracy(source: str, translated: str) -> float:
    tokens = []
    operators = {"AND", "OR", "NOT", "XOR", "NAND", "NOR", "XNOR", "AI", "ML", "DL",
                 "CPU", "GPU", "NPU", "NLP", "LLM", "RAG", "API", "CNN", "RNN", "LSTM", "GAN", "SGD"}
    for line in source.splitlines():
        values = re.findall(
            r"(?:\d+(?:[.,]\d+)*%?)|(?:[A-Z]{2,}(?:[A-Z0-9]|[_-](?=[A-Z0-9]))*)|(?:https?://\S+)",
            line,
        )
        tokens.extend(value for value in values
                      if not (line.isupper() and value.isalpha() and value.isupper() and value not in operators))
    if not tokens:
        return 1.0
    normalized_translation = re.sub(r'\s+', '', translated).casefold()
    ai_names = ('artificialintelligence', '人工智能', '人工知能', '인공지능',
                'künstlicheintelligenz', 'intelligenceartificielle')
    aliases = {'US': ('米国', 'アメリカ', '美国'), 'RMB': ('人民元', '人民币')}
    units = {'thousand': Decimal('1000'), 'million': Decimal('1000000'),
             'billion': Decimal('1000000000'), '万': Decimal('10000'),
             '億': Decimal('100000000'), '亿': Decimal('100000000')}
    target_amounts = set()
    for amount, unit in re.findall(r'(\d+(?:\.\d+)?)\s*(thousand|million|billion|万|億|亿)', translated, re.I):
        target_amounts.add(Decimal(amount) * units[unit.lower()])
    def preserved(token: str) -> bool:
        if token in translated or (token == 'AI' and any(name in normalized_translation for name in ai_names)):
            return True
        if any(name in translated for name in aliases.get(token, ())):
            return True
        if re.fullmatch(r'\d+(?:\.\d+)?', token):
            source_units = re.findall(r'(?<![\d.])' + re.escape(token) + r'\s*(thousand|million|billion)(?![A-Za-z])', source, re.I)
            return any(Decimal(token) * units[unit.lower()] in target_amounts for unit in source_units)
        return False
    return sum(preserved(token) for token in tokens) / len(tokens)


def promote_dense_slide_layout(pptx_path: Path, bbox_content: dict[str, Any], minimum_characters: int = 0) -> dict[str, Any]:
    """Use the course-player side panel layout consistently across all pages."""
    if pptx_path.suffix.lower() != '.pptx':
        return bbox_content
    from pptx import Presentation
    slides = Presentation(pptx_path).slides
    for bbox in bbox_content.get('bboxes', []):
        number = int(bbox.get('slide_number', 0))
        if not 1 <= number <= len(slides) or bbox.get('layout_mode') == 'side_panel':
            continue
        characters = sum(len(shape.text) for shape in slides[number - 1].shapes if shape.has_text_frame)
        if characters < minimum_characters:
            continue
        width, height = int(bbox['img_w']), int(bbox['img_h'])
        margin = max(12, round(width * .012))
        size = min(round(width * .18), round(height * .42))
        bbox.update(layout_mode='side_panel', slide_content_width_px=round(width * .78),
                    slide_content_height_px=round(height * .83),
                    x_px=width - size - margin, y_px=height - size - margin,
                    size_px=size, position_description='right-side-panel',
                    layout_policy='dense_slide_text', visible_text_characters=characters)
    return bbox_content


class StageInputRequired(RuntimeError):
    def __init__(self, stage: int, asset: str):
        self.stage = stage
        self.asset = asset
        super().__init__("请上传参考音频" if asset == "referenceAudio" else "请上传讲者图片")


class StageReviewRequired(RuntimeError):
    pass


class FullPipelineRunner:
    """执行软著说明书中的 S1 至 S5 全流程。

    S1、S2 使用当前 Responses 网关；S3、S4、S5 调用恢复目录中的
    CosyVoice 3、FLOAT、页面遮挡分析和视频合成实现。所有模型导入均延迟到
    对应阶段，后端的账户、任务和文件接口可以独立启动。
    """

    def __init__(self, settings: Settings, rag_service=None):
        self.settings = settings
        self.gateway = ResponsesGateway(settings)
        self.rag_service = rag_service
        self.artifact_root = settings.artifact_root.resolve()
        self.workspace_root = (settings.workspace_root or self.artifact_root / "tasks").resolve()
        self.source_root = (settings.source_root or Path(".")).resolve()
        self.workspace_root.mkdir(parents=True, exist_ok=True)
        self.resources = ResourceScheduler(settings.data_root)

    def _url(self, path: str | Path | None) -> str | None:
        if not path:
            return None
        resolved = Path(path).resolve()
        if self.artifact_root == resolved or self.artifact_root in resolved.parents:
            return "/artifacts/" + str(resolved.relative_to(self.artifact_root))
        return str(resolved)

    def _resolve_asset(self, value: Any, fallback: Path | None = None) -> Path:
        raw = value
        if isinstance(value, dict):
            raw = value.get("path") or value.get("url") or value.get("name")
        if not raw:
            if not fallback:
                raise FileNotFoundError("缺少任务素材")
            path = fallback
        else:
            text = str(raw)
            if text.startswith("/artifacts/"):
                path = self.artifact_root / text.removeprefix("/artifacts/")
            else:
                path = Path(text).expanduser()
        path = path.resolve()
        if not path.is_file():
            raise FileNotFoundError(path)
        return path

    def _task_root(self, task_id: str) -> Path:
        return self.workspace_root / _safe_task_component(task_id)

    @staticmethod
    def _bind_reference(root: Path, source: Path, role: str) -> Path:
        directory = root / "references"
        directory.mkdir(parents=True, exist_ok=True)
        destination = directory / (role + source.suffix.lower())
        if destination.resolve() != source.resolve():
            shutil.copy2(source, destination)
        return destination

    def _agent_modules(self):
        package_root = Path(__file__).resolve().parents[1]
        agent_root = package_root / "agents"
        sys.path.insert(0, str(agent_root))
        os.environ.setdefault("SMARTCOURSE_SOURCE_ROOT", str(self.source_root))
        from s3_tts_agent import TTSAgent
        from s4_avatar_agent import AvatarAgent
        from s4_bbox_agent import BBoxAgent
        from s5_compose_agent import ComposeAgent

        return TTSAgent, AvatarAgent, BBoxAgent, ComposeAgent

    def _code_agent_modules(self):
        package_root = Path(__file__).resolve().parents[1]
        agent_root = package_root / "agents"
        sys.path.insert(0, str(agent_root))
        from code2video_agent import Code2VideoAgent
        from code_video_composer import CodeVideoComposeAgent

        return Code2VideoAgent, CodeVideoComposeAgent

    def _courseware_skill(self):
        package_root = Path(__file__).resolve().parents[1]
        agent_root = package_root / "agents"
        sys.path.insert(0, str(agent_root))
        from s1_courseware_skill import S1CoursewareSkill

        return S1CoursewareSkill()

    @staticmethod
    def _quality_system():
        from .s1_quality_system import S1QualitySystem

        return S1QualitySystem()

    def _longcat_avatar_agent(self):
        package_root = Path(__file__).resolve().parents[1]
        agent_root = package_root / "agents"
        sys.path.insert(0, str(agent_root))
        from s4_longcat_avatar_agent import LongCatAvatarAgent

        bundle_root = Path(
            os.environ.get(
                "SMARTCOURSE_LONGCAT_BUNDLE_ROOT",
                "/lpai/volumes/ad-base-vla-vol-ga/mayongjia/ad-base-vla-vol-ga/group/fh/code/aaai27-benchmark-baselines/fh/aaai27-benchmark-baselines/LongCat-Video-Avatar",
            )
        )
        environment_python = Path(
            os.environ.get(
                "SMARTCOURSE_LONGCAT_ENV_PYTHON",
                "/lpai/volumes/ad-base-vla-vol-ga/mayongjia/ad-base-vla-vol-ga/group/fh/envs/longcat-avatar15-xtuners/bin/python",
            )
        )
        return LongCatAvatarAgent(
            longcat_root=str(bundle_root / "LongCat-Video"),
            environment_python=str(environment_python),
            checkpoint_dir=str(bundle_root / "weights/LongCat-Video-Avatar-1.5"),
            gpu_ids=tuple(
                int(value)
                for value in os.environ.get("SMARTCOURSE_LONGCAT_GPUS", "0,1").split(",")
                if value.strip()
            ),
        )

    def run(
        self, task_id: str, task: dict[str, Any], progress: ProgressCallback,
        start_stage: int = 1,
    ) -> dict[str, Any]:
        from .execution_journal import ExecutionJournal
        journal = ExecutionJournal(self._task_root(task_id) / 'execution-journal.json',
                                   {**task, 'id': task_id}, self.settings.gateway_model, start_stage)
        try:
            result = self._run(task_id, task, progress, start_stage, journal)
        except StageReviewRequired:
            journal.finish('waiting_review',message='讲稿与译文已生成，等待教师审阅。')
            raise
        except StageInputRequired as error:
            journal.finish('waiting_input', stage=error.stage, message=str(error))
            raise
        except Exception as error:
            journal.finish('cancelled' if type(error).__name__ == 'TaskCancelled' else 'failed',
                           message=str(error)[:600])
            raise
        journal.finish('completed')
        return result

    def _run(
        self,
        task_id: str,
        task: dict[str, Any],
        progress: ProgressCallback,
        start_stage: int = 1,
        journal=None,
    ) -> dict[str, Any]:
        root = self._task_root(task_id)
        root.mkdir(parents=True, exist_ok=True)
        context_path = root / "execution-context.json"
        context = {
            "taskId": task_id,
            "mode": task.get("mode"),
            "targetLanguages": task.get("targetLanguages", []),
            "startedAtEpoch": time.time(),
        }
        _write_json(context_path, context)
        performance = PipelinePerformanceRecorder(task_id)

        def run_stage(code, function):
            from .teaching_memory import use_memory, generation_memory
            resource, message = self.resources.STAGES[code]
            started = time.perf_counter()
            engine = self.settings.gateway_model if code in {'S1','S2'} else 'CosyVoice3' if code=='S3' else str(task.get('inputs',{}).get('avatarEngine',os.environ.get('SMARTCOURSE_AVATAR_ENGINE','float'))) if code=='S4' else 'FFmpeg'
            if code=='S1' and task.get('mode')=='topic' and os.environ.get('SMARTCOURSE_PPT_GENERATOR','deeppresenter')!='qwen':
                engine='DeepPresenter'
            journal.stage(code, 'queued', resource=resource, engine=engine)
            progress(int(code[1:]), "queued", 0, {"step": message, "resource": resource})
            with self.resources.acquire(resource):
                waited = time.perf_counter() - started
                progress(int(code[1:]), "running", 0, {})
                journal.stage(code, 'running', queueWaitSeconds=round(waited, 3))
                try:
                    with use_memory(generation_memory(task), code):
                        result = performance.run_stage(code, function)
                    journal.stage(code, 'completed', elapsedSeconds=performance.stages[code]['elapsedSeconds'])
                    return result
                except Exception as error:
                    journal.stage(code, 'failed', message=str(error)[:600])
                    raise
                finally:
                    performance.stages[code]["queueWaitSeconds"] = round(waited, 3)
                    performance.stages[code]["resource"] = resource

        previous_outputs = task.get("outputs", {})
        results: dict[str, Any] = {
            f"s{index}": copy.deepcopy(previous_outputs[f"s{index}"])
            for index in range(1, start_stage)
            if isinstance(previous_outputs.get(f"s{index}"), dict)
        }
        structure_path = root / "s1" / "course_structure.json"
        if start_stage <= 1:
            results["s1"] = run_stage(
                "S1", lambda: self._run_s1(task, root, progress)
            )
        elif "s1" in results:
            pass
        elif structure_path.exists():
            structure = json.loads(structure_path.read_text(encoding="utf-8"))
            results["s1"] = {
                "pageCount": int(
                    structure.get("total_pages", len(structure.get("pages", [])))
                ),
                "courseStructure": self._url(structure_path),
                "pptx": self._url(root / "s1" / "generated.pptx"),
                "inputMode": structure.get("input_mode"),
            }
        else:
            raise FileNotFoundError(structure_path)

        if start_stage <= 2:
            results["s2"] = run_stage(
                "S2", lambda: self._run_s2(task, root, progress)
            )
        if start_stage <= 3:
            if task.get('inputs',{}).get('reviewBeforeMedia') and not task.get('inputs',{}).get('reviewApproved'):
                raise StageReviewRequired('讲稿与译文已生成，等待教师审阅。')
            if not task.get("assets", {}).get("referenceAudio"):
                raise StageInputRequired(3, "referenceAudio")
            results["s3"] = run_stage(
                "S3", lambda: self._run_s3(task, root, progress)
            )
        if start_stage <= 4:
            if not task.get("assets", {}).get("referenceImage"):
                raise StageInputRequired(4, "referenceImage")
            results["s4"] = run_stage(
                "S4", lambda: self._run_s4(task, root, progress)
            )
        if start_stage <= 5:
            compose_task = {**task, "outputs": {**previous_outputs, **results}}
            results["s5"] = run_stage(
                "S5", lambda: self._run_s5(compose_task, root, progress)
            )

        durations = [
            float(value.get("durationSeconds", 0) or 0)
            for value in results.get("s3", {}).get("languages", {}).values()
        ]
        hourly_cost_text = os.environ.get("SMARTCOURSE_NPU_HOURLY_COST", "").strip()
        hourly_cost = float(hourly_cost_text) if hourly_cost_text else None
        performance_report = performance.report(
            max(durations, default=0.0),
            hourly_cost,
        )
        performance_path = root / "performance-report.json"
        _write_json(performance_path, performance_report)

        manifest = {
            "taskId": task_id,
            "status": "completed",
            "languages": task.get("targetLanguages", []),
            "stages": results,
            "performance": self._url(performance_path),
            "finishedAtEpoch": time.time(),
        }
        manifest_path = root / "result-manifest.json"
        _write_json(manifest_path, manifest)
        manifest["manifest"] = self._url(manifest_path)
        return manifest

    def _run_s1(
        self, task: dict[str, Any], root: Path, progress: ProgressCallback
    ) -> dict[str, Any]:
        progress(1, "running", 2, {"step": "读取课程输入"})
        output_dir = root / "s1"
        output_dir.mkdir(parents=True, exist_ok=True)
        structure_path = output_dir / "course_structure.json"
        output_pptx = output_dir / "generated.pptx"
        rag_service = self.rag_service
        if task.get("createdBy"):
            from rag.service import CourseRAGService
            rag_service = CourseRAGService(self.settings.data_root / "rag" / _safe_task_component(task["createdBy"]))
        agent = CourseStructureAgent(
            self.gateway,
            rag_service=rag_service,
            courseware_skill=self._courseware_skill(),
            quality_system=self._quality_system(),
            experience_memory_path=(self.workspace_root / "users" /
                                    _safe_task_component(task.get("createdBy") or task["id"]) /
                                    "courseware-experience-memory.json"),
            require_rag=True,
        )
        if task.get("mode") == "upload":
            source_document = self._resolve_asset(task.get("inputs", {}).get("ppt"))
            conversion = None
            if source_document.suffix.lower()=='.pdf':
                progress(1,"running",12,{"step":"将 PDF 转换为可编辑 PPTX"})
                from .pdf_to_pptx import convert_pdf_to_pptx
                conversion=convert_pdf_to_pptx(source_document,output_pptx,output_dir/'pdf-to-pptx',
                    os.environ.get('SMARTCOURSE_CJK_FONT',''))
                shutil.copy2(source_document,output_dir/'source.pdf')
                pptx=output_pptx
            else:
                pptx=source_document
            progress(1, "running", 20, {"step": "解析 PPT 页面、文本与布局"})
            structure = agent.from_pptx(pptx, task["id"], include_notes=False)
            if pptx.resolve() != output_pptx.resolve():
                shutil.copy2(pptx, output_pptx)
            structure["generated_pptx"] = str(output_pptx)
            if conversion:
                structure.update(input_mode='pdf_converted_pptx',source_pdf=str((output_dir/'source.pdf').resolve()),
                                 pdf_conversion=conversion)
            structure['language_code'] = _language_code(task.get('sourceLanguage','英语'))
            structure['target_language'] = LANGUAGE_MODEL_NAMES[structure['language_code']]
            _write_json(output_dir / 'extracted_structure.json', structure)
            lecture = task.get('inputs', {}).get('script')
            if lecture:
                progress(1, "running", 35, {"step":"读取用户逐页讲稿","total":len(structure['pages'])})
                from .lecture_input import read_page_scripts
                lecture_path = self._resolve_asset(lecture)
                scripts = read_page_scripts(lecture_path, len(structure['pages']))
                for page, script in zip(structure['pages'], scripts):
                    page.update(script_text=script, script_source='user_docx')
                shutil.copy2(lecture_path, output_dir / 'source_lecture.docx')
            else:
                progress(1, "running", 30, {"step":"读取课件画面","total":len(structure['pages'])})
                from .pptx_render import render_pptx
                slide_images = render_pptx(output_pptx, output_dir / 'source-slide-images', width=1280)
                if len(slide_images) != len(structure['pages']):
                    raise ValueError('课件画面与页面数量不一致')
                progress(1, "running", 35, {"step":"根据课件文字与画面逐页生成讲稿","total":len(structure['pages'])})
                from .upload_lecture import generate_upload_lecture
                generated = generate_upload_lecture(
                    structure, self.gateway, structure['language_code'], output_dir / 'generated-lecture',
                    slide_images=slide_images,
                    progress_callback=lambda page,attempt:progress(1,'running',
                        round(35+35*(page if attempt==0 else page-1)/max(1,len(structure['pages']))),
                        {'step':'逐页讲稿已保存' if attempt==0 else '根据课件文字与画面生成讲稿',
                         'page':page,'processed':page if attempt==0 else page-1,'total':len(structure['pages']),
                         'retry':max(0,attempt-1),'retryLimit':3}),
                )
                if [page['page_id'] for page in generated] != [page['page_id'] for page in structure['pages']]:
                    raise ValueError('生成讲稿的页码与课件不一致')
                for page, narration in zip(structure['pages'], generated):
                    page.update(narration)
            _write_json(output_dir / 'source_narration.json', structure)
            from .pptx_notes import write_speaker_notes
            write_speaker_notes(output_pptx, structure['pages'])
        elif task.get("mode") == "topic":
            from .deeppresenter_course import generate_course
            from .generation_retry import reviewed_generation_requirements
            topic = str(task.get("inputs", {}).get("topic", "")).strip()
            pages = int(task.get("inputs", {}).get("totalPages", 10))
            requirements = str(task.get('inputs', {}).get('requirements', ''))
            from .teaching_memory import memory_prompt, generation_memory
            teacher_context = memory_prompt(generation_memory(task), 'S1')
            requirements += teacher_context
            if task.get('inputs', {}).get('answerSourceContext') is not None:
                requirements = reviewed_generation_requirements(output_dir, requirements)
            progress(1, "running", 20, {"step": "正在生成课程课件"})
            source_code = _language_code(task.get("sourceLanguage", "中文"))
            language = LANGUAGE_MODEL_NAMES[source_code]
            source_material=task.get('inputs',{}).get('answerSourceContext')
            if source_material is None and rag_service is not None and rag_service.status().get('ready'):
                source_material=rag_service.query({'query':topic,'topK':8,'generate':False}).get('context')
            generation = generate_course(
                topic, output_pptx, pages=pages,
                language=language,
                audience=str(
                    task.get("inputs", {}).get("targetLearners", "高校本科生")
                ),
                minutes=int(
                    task.get("inputs", {}).get("durationMinutes", 45)
                ),
                requirements=requirements,
                source_material=source_material,
            )
            generation_manifest = Path(generation.get('manifest', ''))
            if generation_manifest.is_file():
                generation['reviewStatus'] = json.loads(generation_manifest.read_text(encoding='utf-8')).get('status')
            structure = agent.from_pptx(output_pptx, task["id"])
            generated_structure = None
            if generation.get('course_structure'):
                generated_structure = json.loads(Path(generation['course_structure']).read_text(encoding='utf-8'))
            source_context = task.get('inputs',{}).get('answerSourceContext')
            if source_context is not None:
                from pptx import Presentation
                from .course_content_review import review_course_content
                review_pages = [{'page': number, 'text': '\n'.join(shape.text for shape in slide.shapes if shape.has_text_frame)}
                                for number, slide in enumerate(Presentation(output_pptx).slides,1)]
                progress(1, 'running', 40, {'step': '核对回答课件与原课程内容'})
                review = review_course_content(review_pages,
                    '准确回答用户问题；遵守用户对资料范围和数值的限制；保持数学条件准确',
                    self.gateway, source_context=source_context)
                _write_json(output_dir/'source-content-review.json', review)
                if not review['passed']:
                    raise ValueError('回答课件未满足原课程资料约束，未启动媒体生成')
            from .topic_lecture import generate_topic_lecture
            if generated_structure is not None:
                lecture = [{'page_id': p['page_id'], 'script_text': p['script_text'],
                            'script_source': p['script_source']} for p in generated_structure['pages']]
            else:
                lecture = generate_topic_lecture(
                    [{'page_id': p['page_id'], 'title': p['title'], 'key_points': p['key_points'],
                      'formulas': p['formulas']} for p in structure['pages']],
                    audience=str(task.get('inputs', {}).get('targetLearners', '高校本科生')),
                    minutes=int(task.get('inputs', {}).get('durationMinutes', 10)),
                    output=output_dir / 'topic_lecture.json',
                    **({'teacher_context': teacher_context} if teacher_context else {}),
                    endpoint=os.environ.get('SMARTCOURSE_DEEPPRESENTER_ENDPOINT', 'http://127.0.0.1:7811/v1/chat/completions'))
            for page, narration in zip(structure['pages'], lecture):
                page.update(narration)
            from .pptx_notes import write_speaker_notes
            write_speaker_notes(output_pptx, structure['pages'])
            structure.update(input_mode="topic", course_topic=topic, generation=generation)
        elif task.get('mode') == 'code' and str(task.get('inputs', {}).get('codeText', '')).strip():
            from .code_explanation import generate_code_explanation
            from .pptx_notes import write_speaker_notes
            inputs = task['inputs']
            source_code = _language_code(task.get('sourceLanguage', '中文'))
            progress(1, 'running', 18, {'step': '分析用户代码并组织逐页讲解'})
            pages = generate_code_explanation(
                str(inputs['codeText']), str(inputs.get('codeLanguage', '')), source_code,
                str(inputs.get('targetLearners', '学习者')), int(inputs.get('totalPages', 3)),
                self.gateway, output_dir / 'code_explanation.json',
            )
            CourseStructureAgent._write_pptx(pages, output_pptx)
            write_speaker_notes(output_pptx, pages)
            structure = {'task_id': task['id'], 'input_mode': 'code_explanation',
                         'course_topic': task['name'], 'language_code': source_code,
                         'target_language': LANGUAGE_MODEL_NAMES[source_code],
                         'code_language': str(inputs.get('codeLanguage', '')),
                         'total_pages': len(pages), 'generated_pptx': str(output_pptx),
                         'pages': pages}
        else:
            lesson_id = str(task.get("inputs", {}).get("lessonId", "")).strip()
            spec = get_lesson_spec(lesson_id)
            progress(1, "running", 8, {"step": "Planner 编排概念、示例、应用和自测顺序"})
            Code2VideoAgent, _CodeVideoComposeAgent = self._code_agent_modules()

            def code_progress(value: int, detail: dict[str, Any]) -> None:
                progress(1, "running", min(96, 8 + round(value * 0.88)), detail)

            code_output = output_dir / "code2video"
            code_result = Code2VideoAgent().build(
                spec,
                code_output,
                render=True,
                progress=code_progress,
            )
            pages = [
                {
                    "page_id": scene["scene_id"],
                    "title": scene["title"],
                    "script_text": scene["script_text"],
                    "key_points": scene["key_points"],
                    "layout": {
                        "template": "manim_animation",
                        "visual_anchor": scene["visual"],
                        "avatar_safe_area": scene["avatar_position"],
                    },
                    "script_source": "code2video_planner",
                }
                for scene in spec["scenes"]
            ]
            CourseStructureAgent._write_pptx(pages, output_pptx)
            structure = {
                "task_id": task["id"],
                "input_mode": "code_animation",
                "course_topic": spec["topic"],
                "subject": spec["subject"],
                "grade": spec["grade"],
                "course_outline": [scene["title"] for scene in spec["scenes"]],
                "language": "Simplified Chinese",
                "total_pages": len(pages),
                "generated_pptx": str(output_pptx),
                "planner": code_result["planner"],
                "manim_source": code_result["source"],
                "critic_report": code_result["criticReport"],
                "scene_videos": code_result["sceneVideos"],
                "code2video_manifest": code_result["manifest"],
                "pages": pages,
            }
        save_structure(structure, structure_path)
        from .narration import export_lecture_text
        lecture_path=export_lecture_text(structure,output_dir/'lecture.txt')
        advanced = bool(task.get("inputs", {}).get("enableTeachingExperiments")) or task.get("inputs", {}).get("avatarEngine") == "longcat-avatar-1.5" or (task.get("mode") == "code" and bool(task.get("inputs", {}).get("lessonId")))
        if not advanced:
            result = {
                "pageCount": len(structure.get("pages", [])),
                "formulaCount": sum(len(page.get("formulas", [])) for page in structure.get("pages", [])),
                "courseStructure": self._url(structure_path),
                "lectureScript": self._url(lecture_path),
                "pptx": self._url(output_pptx),
                "inputMode": structure.get("input_mode"),
                "reviewStatus": structure.get("generation", {}).get("reviewStatus"),
            }
            progress(1, "completed", 100, result)
            return result
        experiment_condition = str(
            task.get("inputs", {}).get("experimentCondition", "A4")
        ).strip() or "A4"
        learner_state = task.get("inputs", {}).get("learnerState")
        pedagogical_ir, constraint_report = compile_course_structure_ir(
            structure,
            condition=experiment_condition,
            learner_state=learner_state,
        )
        pedagogical_ir_path = output_dir / "pedagogical_ir.json"
        constraint_report_path = output_dir / "constraint_report.json"
        e2_variants_path = output_dir / "e2_variants.json"
        adaptive_plan_path = output_dir / "adaptive_plan.json"
        adaptive_experiment_path = output_dir / "adaptive_experiment.json"
        phase_c_plan_path = output_dir / "phase_c_plan.json"
        teacher_profile_path = output_dir / "teacher_profile.json"
        gesture_schedule_path = output_dir / "gesture_schedule.json"
        rehearsal_report_path = output_dir / "rehearsal_report.json"
        _write_json(pedagogical_ir_path, pedagogical_ir)
        _write_json(constraint_report_path, constraint_report)
        _write_json(e2_variants_path, build_e2_variants(pedagogical_ir))
        adaptive_plan = apply_adaptive_policy(
            pedagogical_ir,
            learner_state or {},
            condition=str(task.get("inputs", {}).get("teachingPolicy", "E4")),
        )
        adaptive_experiment = build_adaptive_experiment(
            pedagogical_ir,
            learner_state or {},
        )
        _write_json(adaptive_plan_path, adaptive_plan)
        _write_json(adaptive_experiment_path, adaptive_experiment)
        normalized_state = adaptive_plan["policyReport"]["learnerState"]
        phase_c_plan = {
            "experiments": ["E4", "E6", "E7", "E8", "E9"],
            "proactiveIntervention": select_proactive_intervention(
                normalized_state,
                "D4",
                {"stateUncertainty": 0.4},
            ),
            "structuredMemory": update_structured_memory(
                None,
                {
                    "concept": str(structure.get("course_topic", "course")),
                    "evidenceType": "course_creation",
                    "success": True,
                    "confidence": 1.0,
                },
                "F4",
            ),
            "personalization": build_personalization_plan(
                normalized_state,
                task.get("inputs", {}).get("preferences") or {},
                "G5",
            ),
            "pedagogicalAlignment": check_pedagogical_alignment(
                normalized_state,
                {"requestsDirectAnswer": False, "studentReasoning": ""},
                "H4",
            ),
            "prescreen": prescreen_candidate(
                constraint_report,
                constraint_report,
                "I0",
            ),
        }
        _write_json(phase_c_plan_path, phase_c_plan)
        assets = task.get("assets", {}) if isinstance(task.get("assets"), dict) else {}
        teacher_profile = build_teacher_profile(
            profile_id=str(task.get("teacherProfileId", f"teacher_{task['id']}")),
            reference_image=assets.get("referenceImage"),
            reference_audio=assets.get("referenceAudio"),
            gesture_clips=assets.get("gestureClips") or {},
            consent=assets.get("consent") or {},
            full_body_engine={
                "name": "longcat-video-avatar-1.5",
                "available": longcat_engine_readiness(
                    os.environ.get(
                        "SMARTCOURSE_LONGCAT_BUNDLE_ROOT",
                        "/lpai/volumes/ad-base-vla-vol-ga/mayongjia/ad-base-vla-vol-ga/group/fh/code/aaai27-benchmark-baselines/fh/aaai27-benchmark-baselines/LongCat-Video-Avatar",
                    ),
                    os.environ.get(
                        "SMARTCOURSE_LONGCAT_ENV_PYTHON",
                        "/lpai/volumes/ad-base-vla-vol-ga/mayongjia/ad-base-vla-vol-ga/group/fh/envs/longcat-avatar15-xtuners/bin/python",
                    ),
                )["ready"],
            },
        )
        teacher_profile["fileValidation"] = validate_profile_files(
            teacher_profile,
            self.artifact_root,
        )
        gesture_schedule = schedule_gestures(pedagogical_ir, teacher_profile)
        rehearsal_report = analyze_rehearsal(structure, pedagogical_ir)
        _write_json(teacher_profile_path, teacher_profile)
        _write_json(gesture_schedule_path, gesture_schedule)
        _write_json(rehearsal_report_path, rehearsal_report)
        structure["pedagogical_ir"] = {
            "schemaVersion": pedagogical_ir["schemaVersion"],
            "condition": experiment_condition,
            "eventCount": pedagogical_ir["timeline"]["eventCount"],
            "constraintSatisfaction": constraint_report["constraintSatisfaction"],
            "path": str(pedagogical_ir_path),
        }
        save_structure(structure, structure_path)
        result = {
            "pageCount": int(structure.get("total_pages", len(structure.get("pages", [])))),
            "courseStructure": self._url(structure_path),
            "pptx": self._url(output_pptx),
            "inputMode": structure.get("input_mode"),
            "courseBrief": self._url(output_dir / "course_brief.json")
            if (output_dir / "course_brief.json").is_file()
            else None,
            "courseQualityReport": self._url(output_dir / "course_quality_report.json")
            if (output_dir / "course_quality_report.json").is_file()
            else None,
            "visualQualityReport": self._url(output_dir / "visual_quality_report.json")
            if (output_dir / "visual_quality_report.json").is_file()
            else None,
            "learningEvaluation": self._url(output_dir / "learning_evaluation.json")
            if (output_dir / "learning_evaluation.json").is_file()
            else None,
            "researchCoverage": self._url(output_dir / "research_coverage.json")
            if (output_dir / "research_coverage.json").is_file()
            else None,
            "pedagogicalIr": self._url(pedagogical_ir_path),
            "constraintReport": self._url(constraint_report_path),
            "e2Variants": self._url(e2_variants_path),
            "adaptivePlan": self._url(adaptive_plan_path),
            "adaptiveExperiment": self._url(adaptive_experiment_path),
            "phaseCPlan": self._url(phase_c_plan_path),
            "teacherProfile": self._url(teacher_profile_path),
            "gestureSchedule": self._url(gesture_schedule_path),
            "rehearsalReport": self._url(rehearsal_report_path),
            "gestureAssetsReady": teacher_profile["readiness"]["gestureCourse"],
            "pedagogicalEventCount": pedagogical_ir["timeline"]["eventCount"],
            "experimentCondition": experiment_condition,
        }
        if structure.get("input_mode") == "code_animation":
            result.update(
                {
                    "planner": self._url(structure.get("planner")),
                    "manimSource": self._url(structure.get("manim_source")),
                    "criticReport": self._url(structure.get("critic_report")),
                    "sceneVideos": [self._url(path) for path in structure.get("scene_videos", [])],
                    "code2videoManifest": self._url(structure.get("code2video_manifest")),
                }
            )
        progress(1, "completed", 100, result)
        return result

    def _translation_prompt(
        self, page: dict[str, Any], target_name: str, previous: dict[str, Any] | None
    ) -> str:
        from .narration_style import DIRECT_TEACHING_RULE, authored_by_user
        narration_rule = '' if authored_by_user(page) else DIRECT_TEACHING_RULE
        correction = ""
        if previous:
            correction = (
                "The previous result failed validation. Correct every issue and return a new result.\n"
                + json.dumps(previous, ensure_ascii=False)
            )
        return f"""You are the translation and validation agent in a multilingual MOOC system.
Translate the page into {target_name}. Preserve page_id, numbers, formulas, symbols, URLs,
proper nouns and academic terminology. Do not summarize or add content. Translate each field
independently. Do not add material from the title or key_points to script_text. Translate only
the supplied script_text as narration; preserve empty fields. Keep named Boolean logic operators
such as AND, OR and NOT in their original uppercase form when referring to
a gate or its cases; a local-language explanation may follow in parentheses.
Keep each original numeral and decimal digit sequence exactly as written in its corresponding
translated field, including when translating units. Do not convert 1.5 million to 150万 or
1.3 billion to 13億; retain 1.5 and 1.3 with the translated unit instead.
Keep every original uppercase abbreviation, including US and RMB, unchanged in the translation;
local-language explanations may be added beside the original abbreviation.
Back-translate the translated script to the source language and evaluate it. Return JSON only:
{{"page_id":1,"title":"...","key_points":["..."],"script_text":"...",
"back_translation":"...","semantic_similarity":0.0,"terminology_coverage":0.0,
"formula_number_proper_name_accuracy":0.0,"contextual_consistency":0.0,
"passed":true,"note":"..."}}
Narration style (script_text only; do not change slide text): {narration_rule}
{correction}
Source page:
{json.dumps(page, ensure_ascii=False)}"""

    def _translate_page(
        self, page: dict[str, Any], target_name: str, policy=None, protected_names=(),
        _separate_fields: bool = True,
    ) -> tuple[dict[str, Any], dict[str, Any]]:
        if _separate_fields and str(page.get("script_text", "")).strip() and (
            str(page.get("title", "")).strip() or page.get("key_points")
        ):
            # Narration is translated without slide text so a footer or slide point cannot
            # become narration merely because it appeared in the same model request.
            slide_page = {**page, "script_text": ""}
            script_page = {"page_id": page.get("page_id", 0), "title": "",
                           "key_points": [], "script_text": page["script_text"],
                           "script_source":page.get('script_source'),
                           "original_script_source":page.get('original_script_source',page.get('script_source'))}
            slide, slide_report = self._translate_page(
                slide_page, target_name, policy, protected_names, False)
            script, script_report = self._translate_page(
                script_page, target_name, policy, protected_names, False)
            slide["script_text"] = script["script_text"]
            for field in ('original_script_source','narration_style_review','narration_warning'):
                if field in script:slide[field] = script[field]
            report = dict(slide_report)
            for field in ("semantic_similarity", "terminology_coverage",
                          "formula_number_proper_name_accuracy", "contextual_consistency"):
                report[field] = min(slide_report[field], script_report[field])
            report["back_translation"] = script_report["back_translation"]
            report["passed"] = slide_report["passed"] and script_report["passed"] and bool(script["script_text"].strip())
            report["personal_names_preserved"] = slide_report["personal_names_preserved"] and script_report["personal_names_preserved"]
            report["note"] = " ".join(filter(None, (slide_report["note"], script_report["note"])))
            report["attempt"] = max(slide_report["attempt"], script_report["attempt"])
            report["slide_validation"] = slide_report
            report["script_validation"] = script_report
            return slide, report
        from .pptx_names import localized_person_names
        name_map=localized_person_names(protected_names,target_name)
        previous: dict[str, Any] | None = None
        source_text = "\n".join(
            [str(page.get("title", "")), str(page.get("script_text", ""))]
            + [str(value) for value in page.get("key_points", [])]
        )
        policy = validate_policy(policy)
        last_result = None
        last_error = None
        for attempt in range(policy["maxRetries"] + 1):
            try:
                value = self.gateway.respond_json(
                    self._translation_prompt(page, target_name, previous) + "\nRequired quality thresholds:\n" + json.dumps(policy) + "\nUse these official personal-name spellings exactly, including in the script:\n" + json.dumps(name_map, ensure_ascii=False),
                    reasoning_effort="medium",
                    max_output_tokens=12_000,
                    json_retries=0,
                )
                if not isinstance(value, dict) or not isinstance(value.get("key_points", []), list):
                    raise ValueError("译文格式无效")
            except (GatewayError, ValueError, TypeError) as exc:
                last_error = str(exc)
                previous = {"error": "译文返回不完整或格式无效，请重新翻译并返回完整 JSON。"}
                continue
            translated = {
                "page_id": int(page.get("page_id", 0)),
                "title": str(value.get("title", "")),
                "key_points": [str(item) for item in value.get("key_points", [])],
                "script_text": str(value.get("script_text", "")),
                "layout": copy.deepcopy(page.get("layout", {})),
                "script_source": "translation_agent",
                "original_script_source": page.get('original_script_source',page.get('script_source')),
            }
            translated_text = "\n".join(
                [translated["title"], translated["script_text"], *translated["key_points"]]
            )
            exact = _exact_token_accuracy(source_text, translated_text)
            report = {
                "page_id": translated["page_id"],
                "back_translation": str(value.get("back_translation", "")),
                "semantic_similarity": float(value.get("semantic_similarity", 0)),
                "terminology_coverage": float(value.get("terminology_coverage", 0)),
                "formula_number_proper_name_accuracy": min(
                    exact, float(value.get("formula_number_proper_name_accuracy", exact))
                ),
                "contextual_consistency": float(value.get("contextual_consistency", 0)),
                "passed": bool(value.get("passed", False)) and exact == 1.0,
                "note": str(value.get("note", "")),
                "attempt": attempt + 1,
            }
            report["passed"] = report["passed"] and meets_thresholds(report, policy)
            translated_content='\n'.join([translated.get('title',''),translated.get('script_text',''),*translated.get('key_points',[])])
            source_content='\n'.join([str(page.get('title','')),str(page.get('script_text','')),*map(str,page.get('key_points',[]))])
            names_preserved = all(name not in source_content or expected in translated_content for name,expected in name_map.items())
            report["personal_names_preserved"] = names_preserved
            report["passed"] = report["passed"] and names_preserved
            if not names_preserved:
                report["note"] += " 人名需保留原文写法。"
            if str(page.get("script_text", "")).strip():
                from .narration_style import authored_by_user, review_teaching_style, style_warning
                if not authored_by_user(page):
                    style = review_teaching_style(translated['script_text'])
                    translated['narration_style_review'] = style
                    report['narrationStyle'] = style
                    if not style['passed']:
                        report['passed'] = False
                        report['note'] += ' 请直接讲授内容，避免描述幻灯片介绍或展示了什么。'
                from .language_validation import check_target_language
                language_check = check_target_language(translated["script_text"], target_name)
                report["languageCheck"] = language_check
                if language_check["status"] == "mismatch":
                    report["passed"] = False
                    report["note"] += f" 讲稿语种为 {language_check['detected']}，目标语种为 {target_name}。"
            report["thresholds"] = policy
            last_result = translated, report
            if report["passed"]:
                return translated, report
            previous = {"translation": translated, "validation": report}
        if last_result:
            translated, report = last_result
            if translated.get('narration_style_review',{}).get('passed') is False:
                translated['narration_style_review']['continuedAfterRetries'] = True
                from .narration_style import style_warning
                translated['narration_warning'] = style_warning(translated['page_id'], policy['maxRetries'])
            report["attempt"] = policy["maxRetries"] + 1
            if last_error:
                report["note"] += " 后续重试未返回完整译文。"
            return translated, report
        translated = {"page_id": int(page.get("page_id", 0)), "title": str(page.get("title", "")),
                      "key_points": [str(item) for item in page.get("key_points", [])],
                      "script_text": str(page.get("script_text", "")),
                      "layout": copy.deepcopy(page.get("layout", {})), "script_source": "untranslated_source"}
        report = {"page_id": translated["page_id"], "back_translation": "", "semantic_similarity": 0.0,
                  "terminology_coverage": 0.0, "formula_number_proper_name_accuracy": 0.0,
                  "contextual_consistency": 0.0, "passed": False,
                  "note": "多次重试后仍未取得有效译文，本页保留原文，需人工处理。",
                  "attempt": policy["maxRetries"] + 1, "personal_names_preserved": False,
                  "thresholds": policy}
        return translated, report

    def _run_s2(
        self, task: dict[str, Any], root: Path, progress: ProgressCallback
    ) -> dict[str, Any]:
        source = json.loads((root / "s1" / "course_structure.json").read_text(encoding="utf-8"))
        if source.get('input_mode')=='existing_pdf':
            from .pdf_translation_stage import translate_pdf_stage
            codes=[_language_code(value) for value in task.get('targetLanguages',[])]
            outputs={}
            for index,code in enumerate(codes,1):
                progress(2,'running',round((index-1)*100/max(1,len(codes))),{'language':code,'step':'翻译 PDF 原文与逐页讲稿'})
                translated=translate_pdf_stage(source,root/'s2'/code,code,LANGUAGE_MODEL_NAMES[code],
                    _language_code(task.get('sourceLanguage','中文')),self.settings,self._translate_page,
                    os.environ.get('SMARTCOURSE_CJK_FONT',''),task.get('translationPolicy'))
                outputs[code]={**translated,'language':SUPPORTED_LANGUAGES[code],
                    **{key:self._url(translated[key]) for key in ('translatedStructure','translatedPdf','previewPdf','lectureScript','validationReport')},
                    'previewPages':[self._url(path) for path in translated['previewPages']]}
            result={'languageCount':len(outputs),'languages':outputs}
            progress(2,'completed',100,result)
            return result
        source_pages = source.get("pages", [])
        from .pptx_names import discover_person_names, restore_person_fields
        names_file = root / "s2" / "person_names.json"
        if names_file.exists():
            protected_names = json.loads(names_file.read_text(encoding="utf-8"))
        else:
            protected_names = discover_person_names(root / "s1" / "generated.pptx", self.gateway)
            _write_json(names_file, protected_names)
        languages = [_language_code(value) for value in task.get("targetLanguages", [])]
        outputs: dict[str, Any] = {}
        total_units = max(1, len(languages) * max(1, len(source_pages)))
        completed_units = 0
        source_code = _language_code(task.get("sourceLanguage", "中文"))
        for code in languages:
            target_name = LANGUAGE_MODEL_NAMES[code]
            translated_pages: list[dict[str, Any]] = []
            page_reports: list[dict[str, Any]] = []
            identity = code == source_code
            for page in source_pages:
                if identity:
                    translated = copy.deepcopy(page)
                    report = {
                        "page_id": page.get("page_id"),
                        "back_translation": page.get("script_text", ""),
                        "semantic_similarity": 1.0,
                        "terminology_coverage": 1.0,
                        "formula_number_proper_name_accuracy": 1.0,
                        "contextual_consistency": 1.0,
                        "passed": True,
                        "note": "源语言内容一致性校验通过",
                        "attempt": 1,
                    }
                else:
                    checkpoint = root / "s2" / code / "page-checkpoints" / f"page-{int(page['page_id']):03d}.json"
                    from .pptx_names import localized_person_names
                    name_spellings=localized_person_names(protected_names,target_name)
                    request_value = {"taskId":task['id'],"page":page,"language":target_name,
                                     "policy":task.get('translationPolicy'),"names":protected_names,
                                     "model":self.settings.gateway_model,"translationVersion":"direct-teaching-v3"}
                    if any(name_spellings[name]!=name for name in name_spellings):
                        request_value["officialNameSpellings"]=name_spellings
                    try:
                        saved = json.loads(checkpoint.read_text(encoding="utf-8")) if checkpoint.is_file() else {}
                    except (ValueError,OSError):
                        saved = {}
                    if saved.get('input') == request_value and saved.get('report',{}).get('passed'):
                        translated, report = saved['translation'], saved['report']
                    else:
                        translated, report = self._translate_page(page, target_name, task.get("translationPolicy"), protected_names)
                        _write_json(checkpoint, {"input":request_value,"translation":translated,"report":report})
                translated_pages.append(translated)
                page_reports.append(report)
                completed_units += 1
                progress(
                    2,
                    "running",
                    min(99, round(completed_units * 100 / total_units)),
                    {
                        "language": code,
                        "processed": completed_units,
                        "total": total_units,
                        "step": "翻译、回译和术语公式校验",
                    },
                )
            directory = root / "s2" / code
            structure = {
                "task_id": task["id"],
                "source_language": task.get("sourceLanguage", "中文"),
                "target_language": target_name,
                "language_code": code,
                "total_pages": len(translated_pages),
                "pages": translated_pages,
            }
            structure_path = directory / "translated_structure.json"
            pptx_path = directory / "translated.pptx"
            report_path = directory / "translation_report.json"
            _write_json(structure_path, structure)
            from .narration import export_lecture_text
            lecture_path=export_lecture_text(structure,directory/'lecture.txt')
            original_pptx = root / "s1" / "generated.pptx"
            if identity:
                shutil.copy2(original_pptx, pptx_path)
                package_report = {"translatedTextNodes": 0, "identity": True}
            else:
                from .pptx_translation import translate_package
                package_report = translate_package(original_pptx, pptx_path, self.gateway, target_name,
                    skip_speaker_notes=True,protected_names=protected_names,progress=lambda done,total:progress(2,'running',min(99,round(completed_units*100/total_units)),
                    {'language':code,'processed':done,'total':total,'step':'课件 XML 文本翻译'}))
                package_report["protectedNameFields"] = restore_person_fields(original_pptx, pptx_path, protected_names,target_name)
                from .pptx_translation import synchronize_slide_titles
                title_changes = synchronize_slide_titles(pptx_path, translated_pages)
                package_report["synchronizedSlideTitles"] = len(title_changes)
                _write_json(directory / "title-consistency-report.json", {"changes": title_changes})
                from .pptx_notes import write_speaker_notes
                package_report['speakerNotePages'] = write_speaker_notes(pptx_path, translated_pages)
                from .pptx_fonts import apply_language_font
                package_report["fontAdjustedProperties"] = apply_language_font(pptx_path, pptx_path, code)
                if package_report["fallbackImagesRequiringRefresh"]:
                    from .pptx_math_images import refresh_math_images
                    office = os.environ.get("SMARTCOURSE_SOFFICE") or shutil.which("soffice")
                    if not office:
                        raise RuntimeError("PPT 备用图片需要重新渲染，请配置 SMARTCOURSE_SOFFICE")
                    refreshed = directory / "formula-refresh.pptx"
                    image_report = refresh_math_images(
                        pptx_path, refreshed, directory / "formula-refresh-work", office
                    )
                    refreshed.replace(pptx_path)
                    package_report["refreshedImages"] = len(image_report)
                    package_report["preservedFormulaFallbacks"] = 0
                from .pptx_text_fit import fit_translated_text
                text_fit = fit_translated_text(original_pptx, pptx_path)
                _write_json(directory / "text-fit-report.json", text_fit)
                package_report["textRegionsResized"] = len(text_fit)
            metric_names = (
                "semantic_similarity",
                "terminology_coverage",
                "formula_number_proper_name_accuracy",
                "contextual_consistency",
            )
            metrics = {
                name: round(mean(item[name] for item in page_reports), 4)
                if page_reports
                else 0
                for name in metric_names
            }
            report = {
                "package": package_report,
                "languageCode": code,
                "pageCount": len(page_reports),
                "passed": all(item["passed"] for item in page_reports) and not package_report.get("warnings"),
                "automaticCorrectionCount": sum(
                    sum(item[key]["attempt"] - 1 for key in ("slide_validation", "script_validation"))
                    if "slide_validation" in item else item["attempt"] - 1
                    for item in page_reports
                ),
                "metrics": metrics,
                "pages": page_reports,
            }
            report["warning"] = _translation_warning(page_reports)
            if package_report.get("warnings"):
                report["warning"] += ("；" if report["warning"] else "") + "部分 PPT 段落未通过翻译检查，请审阅。"
            _write_json(report_path, report)
            if code in {'ja-JP','de-DE'}:
                from .pptx_font_embedding import embed_fonts
                font_root=Path(os.environ.get('SMARTCOURSE_FONT_ROOT', str(self.settings.artifact_root.parent/'model-cache'/'fonts')))
                regular=font_root/'SmartCourseJP-Regular.ttf'
                bold=font_root/'SmartCourseJP-Bold.ttf'
                if not regular.is_file(): regular=font_root/'smartcourse-cn-font.ttc'
                if not bold.is_file(): bold=regular
                if regular.suffix.lower() == '.ttc':
                    package_report['embeddedFonts']=0
                    package_report['preservedFontCollection']=str(regular)
                else:
                    package_report['embeddedFonts']=embed_fonts(pptx_path,regular,bold,'SmartCourse JP')
                _write_json(report_path,report)
            from .pptx_render import render_pptx
            previews = render_pptx(pptx_path, directory / "preview", pdf_name="translated.pdf")
            from .pptx_text_fit import fit_visual_collisions
            collision_fit = [] if task.get("mode") == "upload" else fit_visual_collisions(
                pptx_path, directory / "preview" / "translated.pdf")
            if collision_fit:
                previews = render_pptx(pptx_path, directory / "preview", pdf_name="translated.pdf")
            package_report["visualCollisionsFixed"] = len(collision_fit)
            _write_json(directory / "layout-check-report.json", {
                "passed": True,
                "visualCollisionsFixed": collision_fit,
            })
            outputs[code] = {
                "language": SUPPORTED_LANGUAGES[code],
                "translatedStructure": self._url(structure_path),
                "translatedPptx": self._url(pptx_path),
                "lectureScript": self._url(lecture_path),
                "previewPdf": self._url(directory / "preview" / "translated.pdf"),
                "previewPages": [self._url(path) for path in previews],
                "pageCount": len(translated_pages),
                "validationReport": self._url(report_path),
                "passed": report["passed"],
                "warning": report["warning"],
                "metrics": metrics,
            }
        result = {"languageCount": len(outputs), "languages": outputs}
        progress(2, "completed", 100, result)
        return result

    def _run_s3(
        self, task: dict[str, Any], root: Path, progress: ProgressCallback
    ) -> dict[str, Any]:
        TTSAgent, _AvatarAgent, _BBoxAgent, _ComposeAgent = self._agent_modules()
        reference_audio = self._resolve_asset(
            task.get("assets", {}).get("referenceAudio"),
        )
        reference_audio = self._bind_reference(root, reference_audio, "speaker")
        from .reference_audio import prepare_reference_audio
        reference_audio = prepare_reference_audio(reference_audio, root / "references")
        codes = [_language_code(value) for value in task.get("targetLanguages", [])]
        runtime = root / "runtime"
        outputs: dict[str, Any] = {}
        from .narration import export_lecture_text
        export_lecture_text(json.loads((root/'s1/course_structure.json').read_text(encoding="utf-8")),root/'s1/lecture.txt')
        agent = TTSAgent(
            workspace_root=str(runtime),
            speaker_prompt=str(reference_audio),
            whisper_model=os.environ.get("SMARTCOURSE_WHISPER_MODEL", "tiny"),
        )
        try:
            for index, code in enumerate(codes, start=1):
                source_json = root / "s2" / code / "translated_structure.json"
                export_lecture_text(json.loads(source_json.read_text(encoding="utf-8")),source_json.parent/'lecture.txt')
                from .narration import prepare_narration
                narration_input = json.loads(source_json.read_text(encoding='utf-8'))
                previous_narration = root/'s3'/code/'narration_structure.json'
                if task.get('inputs',{}).get('repairScope') and previous_narration.is_file():
                    from .course_revisions import preserve_revision_narration
                    narration_input = preserve_revision_narration(narration_input,
                        json.loads(previous_narration.read_text(encoding='utf-8')),task['inputs']['repairScope'],code)
                source_json = prepare_narration(narration_input, self.gateway,
                                                root / "s3" / code / "narration_structure.json",
                                                preserve_script=task.get('mode') == 'upload' or bool(task.get('inputs',{}).get('repairScope')),
                                                enforce_direct_style=not bool(task.get('inputs',{}).get('script')),
                                                review_content=task.get('mode') == 'topic' and os.environ.get('SMARTCOURSE_STRICT_NARRATION_REVIEW') == '1',
                                                target_duration_seconds=(float(task.get('inputs',{}).get('durationMinutes',10))*60
                                                                         if task.get('mode') == 'topic' else None),
                                                progress_callback=lambda page, attempt: progress(3, "running",
                                                    min(99, round((index-1)*100/max(1,len(codes)))),
                                                    {"language":code,"page":page,"step":"准备逐页朗读讲稿",
                                                     "retry":attempt-1,"retryLimit":3}))
                from .narration import export_lecture_text
                oral_structure=json.loads(source_json.read_text(encoding="utf-8"))
                oral_text=export_lecture_text(oral_structure,source_json.parent/'lecture.txt')
                from .narrated_document import prepare_narrated_document
                final_document=prepare_narrated_document(root/'s2'/code,source_json.parent,oral_structure,code)
                def speech_progress(completed, total, phase, page_id):
                    progress(3, "running", min(99, round(((index-1)+completed/max(1,total))*100/max(1,len(codes)))),
                             {"language":code,"page":page_id,"processed":completed,"total":total,
                              "step":"语音合成" if phase=="synthesis" else "语音时间戳对齐"})
                from s3_tts_agent import SpeechCoverageError
                result = _run_speech_with_retries(
                    lambda: agent.run(code, str(source_json), language=code, progress_callback=speech_progress),
                    lambda exc, attempt: progress(3, "running", 0, {"language": code, "step": str(exc),
                                                                      "retry": attempt, "retryLimit": 3}),
                    SpeechCoverageError,
                )
                if result.get("error"):
                    raise RuntimeError(f"{code} 语音合成失败：{result['error']}")
                requested_seconds = float(task.get('inputs',{}).get('durationMinutes',10))*60
                if task.get('mode') == 'topic' and float(result.get('total_duration',0)) > requested_seconds*1.25:
                    raise ValueError('实际口播时长超过目标时长的 125%，需要重新组织讲稿后生成')
                actual_pages=json.loads(Path(result['timeline_path']).read_text(encoding="utf-8"))['pages']
                if len(actual_pages)!=len(oral_structure['pages']):
                    raise ValueError('语音页面数量与逐页讲稿不一致')
                for page,actual in zip(oral_structure['pages'],actual_pages):
                    if page['page_id']!=actual['page_id']:raise ValueError('语音页面顺序与讲稿不一致')
                    page['script_text']=actual['script_text']
                _write_json(source_json,oral_structure)
                export_lecture_text(oral_structure,oral_text)
                final_document=prepare_narrated_document(root/'s2'/code,source_json.parent,oral_structure,code)
                outputs[code] = {
                    "timeline": self._url(result.get("timeline_path")),
                    "fullAudio": self._url(result.get("full_audio_path")),
                    "srt": self._url(result.get("srt_path")),
                    "durationSeconds": result.get("total_duration"),
                    "speechPacing": result.get("speech_pacing"),
                    "warning": "；".join([page['narration_warning'] for page in oral_structure['pages'] if page.get('narration_warning')]
                        + (["部分语句的语速检查仍需复核，请试听生成语音。"] if (result.get('speech_pacing') or {}).get('passed') is False else [])),
                    "pageCount": result.get("pages"),
                    "referenceAudio": self._url(reference_audio),
                    "narrationStructure": self._url(source_json),
                    "lectureScript": self._url(oral_text),
                    ("coursePdf" if final_document.suffix.lower()=='.pdf' else "coursePptx"): self._url(final_document),
                }
                progress(
                    3,
                    "running",
                    min(99, round(index * 100 / max(1, len(codes)))),
                    {"language": code, "processed": index, "total": len(codes)},
                )
        finally:
            getattr(agent.aligner, "close", lambda: None)()
            getattr(agent.tts_engine, "close", lambda: None)()
        result = {"languageCount": len(outputs), "languages": outputs}
        progress(3, "completed", 100, result)
        return result

    def _run_s4(
        self, task: dict[str, Any], root: Path, progress: ProgressCallback
    ) -> dict[str, Any]:
        _TTSAgent, AvatarAgent, BBoxAgent, _ComposeAgent = self._agent_modules()
        reference_image = self._resolve_asset(
            task.get("assets", {}).get("referenceImage"),
        )
        reference_image = self._bind_reference(root, reference_image, "teacher")
        codes = [_language_code(value) for value in task.get("targetLanguages", [])]
        runtime = root / "runtime"
        outputs: dict[str, Any] = {}
        avatar_engine = str(
            task.get("inputs", {}).get("avatarEngine", os.environ.get("SMARTCOURSE_AVATAR_ENGINE", "float"))
        ).strip()
        bbox_agent = BBoxAgent(workspace_root=str(runtime))
        avatar_agent = None
        if avatar_engine == "flashhead":
            from agents.flashhead_agent import FlashHeadAvatarAgent
            avatar_agent = FlashHeadAvatarAgent(workspace_root=str(runtime), ref_image=str(reference_image))
        elif avatar_engine != "longcat-avatar-1.5":
            avatar_agent = AvatarAgent(
                workspace_root=str(runtime),
                float_dir=str(self.source_root / "float"),
                ckpt_path=str(self.source_root / "float/checkpoints/float.pth"),
                ref_image=str(reference_image),
                gpu_id=self.settings.gpu_id,
            )
        for index, code in enumerate(codes, start=1):
            pptx = root / "s2" / code / "translated.pptx"
            structure_file=root/'s2'/code/'translated_structure.json'
            if structure_file.is_file() and json.loads(structure_file.read_text(encoding="utf-8")).get('input_mode')=='existing_pdf':
                pptx=root/'s2'/code/'translated.pdf'
            timeline = runtime / code / "s3_output" / "timeline.json"
            if task.get("mode") == "code" and task.get("inputs", {}).get("lessonId"):
                bbox_path = runtime / code / "s4_output" / "bbox.json"
                bbox_path.parent.mkdir(parents=True, exist_ok=True)
                bbox_value = {
                    "task_id": code,
                    "bboxes": [
                        {
                            "slide_number": index,
                            "x_px": 1518,
                            "y_px": 672,
                            "size_px": 360,
                            "img_w": 1920,
                            "img_h": 1080,
                            "reason": "Manim 场景在右下区域预留数字人位置",
                        }
                        for index in range(1, len(json.loads(timeline.read_text(encoding="utf-8")).get("pages", [])) + 1)
                    ],
                }
                _write_json(bbox_path, bbox_value)
                bbox = {"bbox_path": str(bbox_path)}
            else:
                existing_bbox = runtime / code / "s4_output" / "bbox.json"
                reuse_bbox = False
                if existing_bbox.is_file() and existing_bbox.stat().st_mtime >= pptx.stat().st_mtime:
                    try:
                        saved_bbox = json.loads(existing_bbox.read_text(encoding="utf-8"))
                        page_count = len(json.loads(timeline.read_text(encoding="utf-8"))["pages"])
                        reuse_bbox = (
                            Path(saved_bbox.get("pptx_path", "")).resolve() == pptx.resolve()
                            and saved_bbox.get("total_slides") == page_count
                            and len(saved_bbox.get("bboxes", [])) == page_count
                        )
                    except (OSError, ValueError, KeyError, TypeError, AttributeError):
                        reuse_bbox = False
                bbox = {"bbox_path": str(existing_bbox)} if reuse_bbox else bbox_agent.run(code, str(pptx))
            bbox_file = Path(bbox["bbox_path"])
            bbox_content = json.loads(bbox_file.read_text(encoding="utf-8"))
            bbox_content = promote_dense_slide_layout(pptx, bbox_content)
            source_pdf = root / 's1' / 'source.pdf'
            if (task.get('mode') == 'upload' and source_pdf.is_file()
                    and code == _language_code(task.get('sourceLanguage', '中文'))):
                from .pdf_to_pptx import restore_source_pdf_slide_images
                images = [item['slide_image'] for item in bbox_content.get('bboxes', []) if item.get('slide_image')]
                if len(images) != len(json.loads(timeline.read_text(encoding='utf-8')).get('pages', [])):
                    raise ValueError('PDF 源页面与视频页面不一致')
                restore_source_pdf_slide_images(source_pdf, images)
            for item in bbox_content.get("bboxes", []):
                if item.get("slide_image"):
                    item["imageUrl"] = self._url(item["slide_image"])
            _write_json(bbox_file, bbox_content)
            if avatar_engine == "longcat-avatar-1.5":
                longcat_agent = self._longcat_avatar_agent()
                gesture_schedule = json.loads(
                    (root / "s1" / "gesture_schedule.json").read_text(encoding="utf-8")
                )
                structure_path = root / "s2" / code / "translated_structure.json"
                if not structure_path.is_file():
                    structure_path = root / "s1" / "course_structure.json"
                course_structure = json.loads(structure_path.read_text(encoding="utf-8"))
                avatar_result = longcat_agent.run(
                    task_id=code,
                    timeline_json=str(timeline),
                    reference_image=str(reference_image),
                    gesture_schedule=gesture_schedule,
                    course_structure=course_structure,
                    output_root=str(runtime / code / "s4_output"),
                )
                avatar = {
                    "error": None,
                    "success_pages": avatar_result["successPages"],
                    "total_pages": avatar_result["totalPages"],
                    "avatars_dir": avatar_result["avatarsDir"],
                    "engine": avatar_result["engine"],
                    "pages": avatar_result["pages"],
                }
            else:
                if avatar_engine == "flashhead":
                    def avatar_progress(completed, total, page_id):
                        progress(4, "running", min(99,round(((index-1)+completed/max(1,total))*100/max(1,len(codes)))),
                                 {"language":code,"page":page_id,"processed":completed,"total":total,"step":"数字人分段生成"})
                    avatar = avatar_agent.run(code, str(timeline), progress_callback=avatar_progress)
                else:
                    from .float_process import run_float_avatar
                    def float_progress(completed, total, page_id):
                        progress(4, "running", min(99, round(((index-1)+completed/max(1,total))*100/max(1,len(codes)))),
                                 {"language":code,"page":page_id,"processed":completed,"total":total,
                                  "step":"FLOAT 数字人分段生成"})
                    avatar = run_float_avatar(
                        root, runtime, code, timeline, reference_image,
                        self.source_root / "float", self.settings.gpu_id,
                        progress_callback=float_progress,
                    )
            if avatar.get("error") or avatar.get("success_pages", 0) != avatar.get("total_pages", 0):
                raise RuntimeError(f"{code} 数字人生成未全部完成：{avatar}")
            outputs[code] = {
                "bbox": self._url(bbox.get("bbox_path")),
                "avatarDirectory": self._url(avatar.get("avatars_dir")),
                "pageCount": avatar.get("total_pages"),
                "successPages": avatar.get("success_pages"),
                "referenceImage": self._url(reference_image),
                "avatarEngine": avatar.get("engine", "float"),
                "generationPages": [{**item, "previewUrl": self._url(item.get("video_path") or item.get("videoPath"))}
                                    for item in avatar.get("pages", avatar.get("results", []))],
            }
            progress(
                4,
                "running",
                min(99, round(index * 100 / max(1, len(codes)))),
                {"language": code, "processed": index, "total": len(codes)},
            )
        result = {"languageCount": len(outputs), "languages": outputs}
        progress(4, "completed", 100, result)
        return result

    def _run_s5(
        self, task: dict[str, Any], root: Path, progress: ProgressCallback
    ) -> dict[str, Any]:
        _TTSAgent, _AvatarAgent, _BBoxAgent, ComposeAgent = self._agent_modules()
        codes = [_language_code(value) for value in task.get("targetLanguages", [])]
        runtime = root / "runtime"
        outputs: dict[str, Any] = {}
        compose_agent = ComposeAgent(workspace_root=str(runtime))
        pdf_input=json.loads((root/'s1'/'course_structure.json').read_text(encoding="utf-8")).get('input_mode')=='existing_pdf'
        for index, code in enumerate(codes, start=1):
            language_root = runtime / code
            if task.get("mode") == "code" and task.get("inputs", {}).get("lessonId"):
                structure = json.loads(
                    (root / "s1" / "course_structure.json").read_text(encoding="utf-8")
                )
                Code2VideoAgent, CodeVideoComposeAgent = self._code_agent_modules()
                lesson_id = str(task.get("inputs", {}).get("lessonId", "")).strip()
                timed_spec = get_lesson_spec(lesson_id)
                timeline_value = json.loads(
                    (language_root / "s3_output" / "timeline.json").read_text(encoding="utf-8")
                )
                page_durations = {
                    int(page["page_id"]): float(page["duration"])
                    for page in timeline_value.get("pages", [])
                }
                for scene in timed_spec["scenes"]:
                    if scene["scene_id"] in page_durations:
                        scene["duration_seconds"] = page_durations[scene["scene_id"]]
                timed_code_result = Code2VideoAgent().build(
                    timed_spec,
                    root / "s1" / f"code2video-timed-{code}",
                    render=True,
                )
                code_compose = CodeVideoComposeAgent(workspace_root=str(runtime))
                result = code_compose.run(
                    task_id=code,
                    scene_videos=[str(path) for path in timed_code_result.get("sceneVideos", [])],
                    timeline_json=str(language_root / "s3_output/timeline.json"),
                    avatars_dir=str(language_root / "s4_output/avatars"),
                    subtitle_srt=str(language_root / "s3_output/subtitles.srt"),
                )
                result = {
                    "output_path": result["outputPath"],
                    "total_pages": result["sceneCount"],
                    "composed_segments": result["avatarSceneCount"],
                    "file_size_mb": round(result["fileSizeBytes"] / 1024 / 1024, 2),
                    "code2video_manifest": result["manifest"],
                    "timed_animation_manifest": timed_code_result["manifest"],
                }
            else:
                def compose_progress(completed,total,page_id,step):
                    progress(5,"running",min(99,round(((index-1)+completed/max(1,total))*100/max(1,len(codes)))),
                             {"language":code,"page":page_id,"processed":completed,"total":total,"step":step})
                result = compose_agent.run(
                    task_id=code,
                    pptx_path=str(root / "s2" / code / ("translated.pdf" if pdf_input else "translated.pptx")),
                    timeline_json=str(language_root / "s3_output/timeline.json"),
                    avatars_dir=str(language_root / "s4_output/avatars"),
                    bbox_json=str(language_root / "s4_output/bbox.json"),
                    subtitle_srt=(str(language_root / "s3_output/subtitles.srt")
                                  if task.get('inputs', {}).get('subtitlesEnabled', True) else None),
                    progress_callback=compose_progress,
                )
            if result.get("error") or not result.get("output_path"):
                raise RuntimeError(f"{code} 视频合成失败：{result}")
            from .course_duration import probe_video_duration
            video_duration = probe_video_duration(result['output_path'])
            if pdf_input:
                from .pdf_delivery import package_pdf_result
                packaged=package_pdf_result(root,code,task['name'],result['output_path'])
                outputs[code]={key:self._url(value) if isinstance(value,str) else value for key,value in packaged.items()}
                outputs[code].update(durationSeconds=video_duration, durationSource='video_metadata' if video_duration is not None else None)
                outputs[code]['pageVideos']=[{**page,'path':self._url(page['path'])} for page in packaged['pageVideos']]
                progress(5,'running',min(99,round(index*100/max(1,len(codes)))),{'language':code,'step':'PDF 课程结果已整理'})
                continue
            delivery = root/'delivery'/code
            delivery.mkdir(parents=True,exist_ok=True)
            from .portable_course import prepare_portable_course
            portable=root/'portable'/code
            prepare_portable_course(root/'s3'/code/'course-with-narration.pptx',portable)
            delivery_files = {
                'course.mp4':Path(result['output_path']),
                'course.pptx':portable/'course-editable.pptx',
                'course-view.pptx':portable/'course-view.pptx',
                'course.pdf':portable/'course.pdf',
                'lecture.txt':root/'s3'/code/'lecture.txt',
                'translated-lecture.txt':root/'s2'/code/'lecture.txt',
                'source-lecture.txt':root/'s1/lecture.txt',
                'subtitles.srt':language_root/'s3_output/subtitles.srt',
                'timeline.json':language_root/'s3_output/timeline.json',
            }
            for name,source_file in delivery_files.items():
                if source_file.is_file():shutil.copy2(source_file,delivery/name)
                elif task.get('mode')=='upload':raise RuntimeError('最终交付文件缺失：'+name)
            shutil.copytree(portable,delivery/'preview',dirs_exist_ok=True)
            from .page_video_delivery import export_page_videos
            page_videos = export_page_videos(result['output_path'],
                json.loads((language_root / 's3_output/timeline.json').read_text(encoding="utf-8")), delivery / 'pages')
            (delivery/'README.txt').write_text(f"{task['name']}\n语种：{code}\n页数：{result.get('total_pages')}\n"
                'course.pptx 的备注页对应 lecture.txt 和视频讲解。translated-lecture.txt 保留译文讲稿。\n'
                '字幕时间来自实际语音。数字讲者依据当前任务上传的照片和参考声音生成。\n',encoding='utf-8')
            from .offline_course import write_offline_course
            offline = write_offline_course(delivery,task['name'],code,
                json.loads((language_root / 's3_output/timeline.json').read_text(encoding="utf-8")))
            from .course_archive import archive_course
            bundle=archive_course(delivery, root/'delivery'/f'{code}.zip')
            outputs[code] = {
                "video": self._url(result.get("output_path")),
                "durationSeconds": video_duration,
                "durationSource": "video_metadata" if video_duration is not None else None,
                "warning": task.get("outputs",{}).get("s3",{}).get("languages",{}).get(code,{}).get("warning", ""),
                "pageVideos": [{**item, 'path': self._url(item['path'])} for item in page_videos],
                "offlinePage": self._url(offline),
                "srt": self._url(language_root / "s3_output/subtitles.srt"),
                "timeline": self._url(language_root / "s3_output/timeline.json"),
                "pageCount": result.get("total_pages"),
                "segmentCount": result.get("composed_segments"),
                "fileSizeMb": result.get("file_size_mb"),
                "coursePptx": self._url(delivery/'course.pptx'),
                "lectureScript": self._url(delivery/'lecture.txt'),
                "sourceLecture": self._url(delivery/'source-lecture.txt'),
                "translatedLecture": self._url(delivery/'translated-lecture.txt'),
                "downloadBundle": self._url(bundle),
                "previewPdf": self._url(delivery/'course.pdf'),
                "viewPptx": self._url(delivery/'course-view.pptx'),
                "portablePreview": self._url(delivery/'preview/index.html'),
            }
            progress(
                5,
                "running",
                min(99, round(index * 100 / max(1, len(codes)))),
                {"language": code, "processed": index, "total": len(codes)},
            )
        result = {"languageCount": len(outputs), "languages": outputs}
        progress(5, "completed", 100, result)
        return result

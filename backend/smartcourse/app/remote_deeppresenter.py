"""Invoke the installed Huawei generator through the existing SSH connection."""
import json
import os
import shlex
import subprocess
import tempfile
import uuid
import time
from pathlib import Path
from .remote_media import _connection


def wait_remote_result(ssh, remote, timeout=7500, poll_seconds=10):
    """Observe an existing worker without submitting another generation request."""
    deadline = time.monotonic() + timeout
    last_error = ''
    while time.monotonic() < deadline:
        try:
            result = subprocess.run([*ssh, shlex.join(['cat', remote + '/result.json'])],
                                    capture_output=True, text=True, timeout=30)
        except subprocess.TimeoutExpired:
            last_error = '远程状态暂时无法读取'
            time.sleep(min(poll_seconds, max(0, deadline - time.monotonic())))
            continue
        if result.returncode == 0:
            try:
                state = json.loads(result.stdout)
            except (ValueError, TypeError):
                state = {}
            if state.get('status') == 'completed':
                return state
            if state.get('status') == 'failed':
                raise RuntimeError('华为课程生成失败：' + str(state.get('error', '请检查任务日志')))
        else:
            last_error = '远程状态暂时无法读取'
        time.sleep(min(poll_seconds, max(0, deadline - time.monotonic())))
    raise TimeoutError('等待华为任务状态超时，未重新提交任务。' + last_error)


def generate_remote_course(topic, destination, **options):
    destination = Path(destination)
    destination.parent.mkdir(parents=True, exist_ok=True)
    request = {'topic': topic, **options}
    record_path = destination.with_suffix('.remote-input.json')
    root = os.environ.get('SMARTCOURSE_REMOTE_ROOT', '/home/ma-user/work/smartcourse')
    remote = root + '/course-jobs/' + uuid.uuid4().hex
    if record_path.is_file():
        previous = json.loads(record_path.read_text())
        if previous.get('request') == request and previous.get('root') == root:
            remote = previous['remoteDirectory']
    record_path.write_text(json.dumps({'request': request, 'root': root, 'remoteDirectory': remote}, ensure_ascii=False))
    flags, port, host = _connection()
    ssh = ['ssh', *flags, '-p', str(port), host]
    scp = ['scp', *flags, '-P', str(port)]
    package = Path(__file__).parent
    with tempfile.TemporaryDirectory(prefix='course-request-') as temporary, destination.with_suffix('.remote.log').open('a') as log:
        status = subprocess.run([*ssh, shlex.join(['cat', remote + '/result.json'])], capture_output=True, text=True, timeout=30)
        if status.returncode == 0 and json.loads(status.stdout).get('status') == 'completed':
            pending = destination.with_suffix('.pending.pptx')
            subprocess.run([*scp, host + ':' + remote + '/course.pptx', str(pending)],
                           check=True, stdout=log, stderr=log, timeout=180)
            from pptx import Presentation
            if len(Presentation(pending).slides) != options['pages']:
                raise ValueError('已完成的远程课件页数与请求不一致')
            pending.replace(destination)
            return {'engine': 'DeepPresenter-9B', 'generated_pptx': str(destination),
                    'remoteDirectory': remote, 'reusedRemoteResult': True}
        if status.returncode == 0 and json.loads(status.stdout).get('status') == 'failed':
            from .course_candidate_review import recover_candidate
            recovered = recover_candidate(scp, host, remote, destination, request, log)
            if recovered:
                return recovered
        request_file = Path(temporary) / 'request.json'
        request_file.write_text(json.dumps(request, ensure_ascii=False), encoding='utf-8')
        subprocess.run([*ssh, shlex.join(['mkdir', '-p', remote])], check=True, stdout=log, stderr=log, timeout=30)
        subprocess.run([*scp, str(request_file), str(package / 'deeppresenter_course.py'),
                        str(package / 'deeppresenter_worker.py'), host + ':' + remote + '/'],
                       check=True, stdout=log, stderr=log, timeout=120)
        python = os.environ.get('SMARTCOURSE_REMOTE_PYTHON', '/home/ma-user/anaconda3/envs/MindSpore/bin/python')
        command = ['env', 'SMARTCOURSE_REMOTE_MEDIA=0', 'SMARTCOURSE_ROOT=' + root,
                   'PYTHONPATH=' + root + '/envs/pptagent-runtime:' + root + '/env',
                   python, remote + '/deeppresenter_worker.py', remote + '/request.json']
        try:
            subprocess.run([*ssh, 'source /usr/local/Ascend/ascend-toolkit/set_env.sh && ' + shlex.join(command)],
                           check=True, stdout=log, stderr=log, timeout=7500)
        except subprocess.CalledProcessError as error:
            if error.returncode != 255:
                failure_status = subprocess.run([*ssh, shlex.join(['cat', remote + '/result.json'])],
                                                capture_output=True, text=True, timeout=30)
                if failure_status.returncode != 0 or json.loads(failure_status.stdout).get('status') != 'failed':
                    raise
                from .course_candidate_review import recover_candidate
                recovered = recover_candidate(scp, host, remote, destination, request, log)
                if recovered:
                    return recovered
                raise
            wait_remote_result(ssh, remote)
        except subprocess.TimeoutExpired:
            wait_remote_result(ssh, remote)
        pending = destination.with_suffix('.pending.pptx')
        subprocess.run([*scp, host + ':' + remote + '/course.pptx', str(pending)],
                       check=True, stdout=log, stderr=log, timeout=180)
        from pptx import Presentation
        if len(Presentation(pending).slides) != options['pages']:
            raise ValueError('远程生成课件页数与请求不一致')
        pending.replace(destination)
    return {'engine': 'DeepPresenter-9B', 'generated_pptx': str(destination), 'remoteDirectory': remote}

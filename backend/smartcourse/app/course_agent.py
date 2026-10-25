from __future__ import annotations

import json
from pathlib import Path
from typing import Any

from .gateway import ResponsesGateway
from .omml_text import math_text


EMU_PER_PIXEL_AT_96_DPI = 9_525
DRAWING_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
MATH_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"
COMPATIBILITY_NS = "http://schemas.openxmlformats.org/markup-compatibility/2006"


def _logical_slide_content(slide):
    from lxml import etree
    def content(node):
        if node.tag == f"{{{DRAWING_NS}}}t":
            return node.text or ""
        if node.tag == f"{{{MATH_NS}}}t":
            return math_text(node)
        if node.tag == f"{{{DRAWING_NS}}}br":
            return "\n"
        if node.tag == f"{{{MATH_NS}}}oMath":
            return " " + math_text(node) + " "
        return "".join(content(child) for child in node)
    texts, formulas = [], []
    for paragraph in slide._element.iter(f"{{{DRAWING_NS}}}p"):
        if any(parent.tag == f"{{{COMPATIBILITY_NS}}}Fallback" for parent in paragraph.iterancestors()):
            continue
        text = content(paragraph).strip()
        if text:
            texts.append(text)
    for formula in slide._element.iter(f"{{{MATH_NS}}}oMath"):
        if not any(parent.tag == f"{{{COMPATIBILITY_NS}}}Fallback" for parent in formula.iterancestors()):
            formulas.append({"text": math_text(formula), "omml": etree.tostring(formula, encoding="unicode")})
    return texts, formulas


def _shape_region(shape: Any) -> dict[str, Any]:
    return {
        "x": round(int(shape.left) / EMU_PER_PIXEL_AT_96_DPI),
        "y": round(int(shape.top) / EMU_PER_PIXEL_AT_96_DPI),
        "width": round(int(shape.width) / EMU_PER_PIXEL_AT_96_DPI),
        "height": round(int(shape.height) / EMU_PER_PIXEL_AT_96_DPI),
        "type": str(getattr(shape, "shape_type", "unknown")),
        "name": str(getattr(shape, "name", "")),
    }


def _shape_texts(shape: Any) -> list[str]:
    texts: list[str] = []
    if getattr(shape, "has_text_frame", False):
        for paragraph in shape.text_frame.paragraphs:
            value = paragraph.text.strip()
            if value:
                texts.append(value)
    if getattr(shape, "has_table", False):
        for row in shape.table.rows:
            for cell in row.cells:
                value = cell.text.strip()
                if value:
                    texts.append(value)
    return texts


def _notes_text(slide: Any) -> str:
    try:
        frame = slide.notes_slide.notes_text_frame
        return "\n".join(
            paragraph.text.strip()
            for paragraph in frame.paragraphs
            if paragraph.text.strip()
        )
    except Exception:
        return ""


class CourseStructureAgent:
    def __init__(
        self,
        gateway: ResponsesGateway | None = None,
        *,
        rag_service: Any | None = None,
        courseware_skill: Any | None = None,
        quality_system: Any | None = None,
        experience_memory_path: Path | None = None,
        require_rag: bool = False,
    ):
        self.gateway = gateway
        self.rag_service = rag_service
        self.courseware_skill = courseware_skill
        self.quality_system = quality_system
        self.experience_memory_path = experience_memory_path
        self.require_rag = require_rag

    def from_pptx(self, pptx_path: Path, task_id: str) -> dict[str, Any]:
        from pptx import Presentation

        presentation = Presentation(str(pptx_path))
        pages: list[dict[str, Any]] = []
        for index, slide in enumerate(presentation.slides, start=1):
            title = ""
            if slide.shapes.title is not None:
                title = slide.shapes.title.text.strip()
            texts: list[str] = []
            regions: list[dict[str, Any]] = []
            for shape in slide.shapes:
                values = _shape_texts(shape)
                texts.extend(values)
                region = _shape_region(shape)
                region["text"] = "\n".join(values)
                regions.append(region)
            texts, formulas = _logical_slide_content(slide)
            if not title and texts:
                title = texts[0]
            key_points = [value for value in texts if value != title]
            notes = _notes_text(slide)
            script = notes or "\n".join(key_points or texts)
            pages.append(
                {
                    "page_id": index,
                    "title": title or f"第 {index} 页",
                    "script_text": script,
                    "key_points": key_points,
                    "formulas": formulas,
                    "layout": {
                        "slide_width": round(int(presentation.slide_width) / EMU_PER_PIXEL_AT_96_DPI),
                        "slide_height": round(int(presentation.slide_height) / EMU_PER_PIXEL_AT_96_DPI),
                        "regions": regions,
                    },
                    "script_source": "speaker_notes" if notes else "slide_text",
                }
            )

        return {
            "task_id": task_id,
            "input_mode": "existing_pptx",
            "course_topic": pages[0]["title"] if pages else pptx_path.stem,
            "language": "auto",
            "total_pages": len(pages),
            "generated_pptx": str(pptx_path),
            "pages": pages,
        }

    def _respond_from_messages(self, messages: list[dict[str, str]]) -> Any:
        if not self.gateway:
            raise ValueError("课程主题模式需要 ResponsesGateway")
        prompt = "\n\n".join(
            f"[{str(item.get('role', 'user')).upper()}]\n{item.get('content', '')}"
            for item in messages
        )
        return self.gateway.respond_json(
            prompt,
            reasoning_effort="medium",
            max_output_tokens=24_000,
        )

    def _query_rag(
        self,
        query: str,
        *,
        top_k: int,
        current_chapter: str | None = None,
    ) -> dict[str, Any]:
        if self.rag_service is None:
            if self.require_rag:
                raise ValueError("课程主题模式缺少课程 RAG 服务")
            return {
                "query": query,
                "knowledgeUnits": [],
                "allowedTerms": [],
                "allowedFormulas": [],
                "sources": [],
            }
        result = self.rag_service.query(
            {
                "query": query,
                "topK": top_k,
                "courseId": "ai-literacy-general",
                "courseVersion": "2026.1",
                "currentChapter": current_chapter,
            }
        )
        context = result.get("context") or {}
        if self.require_rag and not context.get("knowledgeUnits"):
            raise ValueError(f"课程 RAG 未返回知识单元：{query}")
        return context

    def _from_topic_enhanced(
        self,
        topic: str,
        task_id: str,
        output_pptx: Path,
        total_pages: int,
        language: str,
        target_learners: str,
        duration_minutes: int,
        learner_state: dict[str, Any] | None,
        preferences: dict[str, Any] | None,
    ) -> dict[str, Any]:
        skill = self.courseware_skill
        initial_context = self._query_rag(
            f"{topic} 课程主旨 教学大纲 基本概念 典型案例 专业术语",
            top_k=14,
        )
        brief = self._respond_from_messages(
            skill.brief_messages(
                topic,
                total_pages,
                language,
                initial_context,
                target_learners=target_learners,
                duration_minutes=duration_minutes,
            )
        )
        if isinstance(brief, dict) and not brief.get("course_thesis"):
            nested_brief = brief.get("course_brief") or brief.get("brief")
            if isinstance(nested_brief, dict):
                brief = nested_brief
        if not isinstance(brief, dict) or not brief.get("course_thesis"):
            raise ValueError("课程备课方案缺少课程主旨")
        if not brief.get("sections") or not brief.get("terminology"):
            raise ValueError("课程备课方案缺少章节或术语表")
        brief_audit = skill.audit_brief(brief)
        brief = skill.repair_brief_deterministically(brief, brief_audit)
        brief_audit = skill.audit_brief(brief)
        brief["style_audit"] = brief_audit
        if not brief_audit["passed"]:
            raise ValueError(
                "课程备课方案未通过文字与结构检查："
                + json.dumps(brief_audit, ensure_ascii=False)
            )
        learner_model = (
            self.quality_system.build_learner_model(
                brief,
                learner_state=learner_state,
                preferences=preferences,
            )
            if self.quality_system is not None
            else None
        )

        section_contexts: dict[str, dict[str, Any]] = {}
        for section in brief.get("sections", []):
            section_id = str(section.get("section_id", ""))
            query = str(
                section.get("retrieval_query")
                or section.get("section_question")
                or section.get("title")
                or topic
            )
            section_contexts[section_id] = self._query_rag(
                query,
                top_k=10,
                current_chapter=str(section.get("title", "")),
            )

        outline = self._respond_from_messages(
            skill.outline_messages(brief, total_pages, language, section_contexts)
        )
        if isinstance(outline, dict):
            outline = (
                outline.get("outline")
                or outline.get("chapters")
                or outline.get("sections")
                or outline.get("course_outline")
            )
        if not isinstance(outline, list):
            raise ValueError("逐页大纲格式不正确")
        actual_pages = sum(len(item.get("pages", [])) for item in outline)
        if actual_pages != total_pages:
            raise ValueError(f"课程页面数量应为 {total_pages}，实际为 {actual_pages}")

        pages: list[dict[str, Any]] = []
        page_id = 1
        for chapter in outline:
            chapter_title = str(chapter.get("chapter", ""))
            section_id = str(chapter.get("section_id", ""))
            for raw_spec in chapter.get("pages", []):
                spec = dict(raw_spec) if isinstance(raw_spec, dict) else {
                    "page_topic": str(raw_spec),
                    "retrieval_query": str(raw_spec),
                }
                spec.setdefault("section_id", section_id)
                retrieval_query = str(
                    spec.get("retrieval_query")
                    or spec.get("page_topic")
                    or spec.get("title")
                    or topic
                )
                page_context = self._query_rag(
                    retrieval_query,
                    top_k=8,
                    current_chapter=chapter_title,
                )
                page = self._respond_from_messages(
                    skill.page_messages(topic, brief, spec, language, page_context)
                )
                if isinstance(page, dict) and not page.get("title"):
                    nested_page = page.get("page") or page.get("slide") or page.get("content")
                    if isinstance(nested_page, dict):
                        page = nested_page
                if not isinstance(page, dict):
                    raise ValueError(f"第 {page_id} 页内容格式不正确")
                page["page_id"] = page_id
                page["chapter"] = chapter_title
                page["section_id"] = section_id
                page["page_spec"] = spec
                page["retrieval_query"] = retrieval_query
                page["rag_context"] = skill.compact_rag_context(page_context, limit=8)
                audit = skill.audit_page(page, brief.get("terminology", []))
                page = skill.repair_page_deterministically(page, audit)
                audit = skill.audit_page(page, brief.get("terminology", []))
                rewrite_rounds = 0
                while not audit["passed"] and rewrite_rounds < skill.config.page_rewrite_limit:
                    page = self._respond_from_messages(
                        skill.rewrite_messages(page, audit, brief, page_context)
                    )
                    if isinstance(page, dict) and not page.get("title"):
                        nested_page = page.get("page") or page.get("slide") or page.get("content")
                        if isinstance(nested_page, dict):
                            page = nested_page
                    if not isinstance(page, dict):
                        raise ValueError(f"第 {page_id} 页修改结果格式不正确")
                    page["page_id"] = page_id
                    page["chapter"] = chapter_title
                    page["section_id"] = section_id
                    page["page_spec"] = spec
                    page["retrieval_query"] = retrieval_query
                    page["rag_context"] = skill.compact_rag_context(page_context, limit=8)
                    audit = skill.audit_page(page, brief.get("terminology", []))
                    rewrite_rounds += 1
                if not audit["passed"]:
                    raise ValueError(
                        f"第 {page_id} 页未通过课程文字与来源检查："
                        + json.dumps(audit, ensure_ascii=False)
                    )
                page["quality_audit"] = audit
                page["rewrite_rounds"] = rewrite_rounds
                page["script_text"] = self._normalize_script_length(
                    str(page.get("script_text", "")), language
                )
                if self.quality_system is not None and learner_model is not None:
                    page = self.quality_system.enrich_page(page, brief, learner_model)
                visual_plan = page.get("visual_plan") or {}
                page["image_description"] = str(
                    page.get("image_description")
                    or spec.get("visual_brief")
                    or "；".join(str(item) for item in visual_plan.get("objects", []))
                )
                page["layout"] = {
                    "template": "teaching_visual",
                    "visual_type": visual_plan.get("type", spec.get("visual_type", "text")),
                    "visual_plan": visual_plan,
                }
                page["script_source"] = "rag_grounded_courseware_skill"
                pages.append(page)
                page_id += 1

        quality_report = skill.audit_deck(brief, pages)
        if not quality_report["passed"]:
            raise ValueError(
                "整套课件未通过课程结构检查："
                + json.dumps(quality_report, ensure_ascii=False)
            )
        self._write_pptx(pages, output_pptx)
        output_pptx.parent.mkdir(parents=True, exist_ok=True)
        visual_report = None
        learning_evaluation = None
        experience_memory = None
        research_coverage = None
        if self.quality_system is not None:
            visual_report = self.quality_system.inspect_pptx(
                output_pptx,
                pages,
                brief,
                output_pptx.parent,
            )
            learning_evaluation = self.quality_system.build_learning_evaluation(
                brief,
                pages,
            )
            memory_path = self.experience_memory_path or (
                output_pptx.parent / "courseware_experience_memory.json"
            )
            experience_memory = self.quality_system.update_experience_memory(
                memory_path,
                pages,
                visual_report,
            )
            research_coverage = self.quality_system.research_coverage(
                pages,
                visual_report,
            )
            (output_pptx.parent / "visual_quality_report.json").write_text(
                json.dumps(visual_report, ensure_ascii=False, indent=2),
                encoding="utf-8",
            )
            (output_pptx.parent / "learning_evaluation.json").write_text(
                json.dumps(learning_evaluation, ensure_ascii=False, indent=2),
                encoding="utf-8",
            )
            (output_pptx.parent / "research_coverage.json").write_text(
                json.dumps(research_coverage, ensure_ascii=False, indent=2),
                encoding="utf-8",
            )
            if not visual_report.get("deterministicPassed", False):
                raise ValueError(
                    "课件未通过渲染后页面检查："
                    + json.dumps(visual_report, ensure_ascii=False)
                )
        (output_pptx.parent / "course_brief.json").write_text(
            json.dumps(brief, ensure_ascii=False, indent=2),
            encoding="utf-8",
        )
        (output_pptx.parent / "course_quality_report.json").write_text(
            json.dumps(quality_report, ensure_ascii=False, indent=2),
            encoding="utf-8",
        )
        return {
            "task_id": task_id,
            "input_mode": "topic",
            "course_topic": topic,
            "course_thesis": brief.get("course_thesis"),
            "course_brief": brief,
            "course_outline": [str(item.get("chapter", "")) for item in outline],
            "terminology": brief.get("terminology", []),
            "quality_report": quality_report,
            "learner_model": learner_model,
            "visual_quality_report": visual_report,
            "learning_evaluation": learning_evaluation,
            "experience_memory": experience_memory,
            "research_coverage": research_coverage,
            "language": language,
            "total_pages": len(pages),
            "generated_pptx": str(output_pptx),
            "pages": pages,
        }

    def from_topic(
        self,
        topic: str,
        task_id: str,
        output_pptx: Path,
        total_pages: int = 10,
        language: str = "Simplified Chinese",
        target_learners: str = "高校本科生",
        duration_minutes: int = 45,
        learner_state: dict[str, Any] | None = None,
        preferences: dict[str, Any] | None = None,
    ) -> dict[str, Any]:
        if not self.gateway:
            raise ValueError("课程主题模式需要 ResponsesGateway")
        if self.courseware_skill is not None:
            return self._from_topic_enhanced(
                topic,
                task_id,
                output_pptx,
                total_pages,
                language,
                target_learners,
                duration_minutes,
                learner_state,
                preferences,
            )
        prompt = f"""You are the course-generation agent in a MOOC production system.
Create a university-level course about: {topic}
Generate exactly {total_pages} PPT pages in {language}. Each page must contain a title,
3 to 5 key_points, a complete teaching script, and a short image_description.
Each teaching script must be suitable for approximately 45 to 75 seconds of speech:
120 to 220 Chinese characters for Chinese, or 90 to 150 words for alphabetic languages.
End the script with a complete sentence and do not repeat the key points verbatim.
Keep definitions and worked examples internally consistent. When discussing learning systems,
distinguish input data, model parameters, fixed decision settings, prediction, and training.
Return JSON only:
{{"course_outline":["..."],"pages":[{{"page_id":1,"title":"...","key_points":["..."],
"script_text":"...","image_description":"..."}}]}}"""
        value = self.gateway.respond_json(
            prompt, reasoning_effort="medium", max_output_tokens=24_000
        )
        pages = value.get("pages", []) if isinstance(value, dict) else []
        if len(pages) != total_pages:
            raise ValueError(f"课程页面数量应为 {total_pages}，实际为 {len(pages)}")
        normalized = []
        for index, page in enumerate(pages, start=1):
            script_text = self._normalize_script_length(
                str(page.get("script_text", "")), language
            )
            normalized.append(
                {
                    "page_id": index,
                    "title": str(page.get("title", f"第 {index} 页")),
                    "key_points": [str(item) for item in page.get("key_points", [])],
                    "script_text": script_text,
                    "image_description": str(page.get("image_description", "")),
                    "layout": {"template": "title_and_content"},
                    "script_source": "course_generation_agent",
                }
            )
        self._write_pptx(normalized, output_pptx)
        return {
            "task_id": task_id,
            "input_mode": "topic",
            "course_topic": topic,
            "course_outline": value.get("course_outline", []),
            "language": language,
            "total_pages": len(normalized),
            "generated_pptx": str(output_pptx),
            "pages": normalized,
        }

    @staticmethod
    def _normalize_script_length(text: str, language: str) -> str:
        import re

        value = re.sub(r"\s+", " ", text or "").strip()
        chinese = "chinese" in language.lower() or "中文" in language
        limit = 220 if chinese else 900
        if len(value) <= limit:
            return value
        sentence_pattern = r"[^。！？.!?]+[。！？.!?]"
        sentences = re.findall(sentence_pattern, value)
        result = ""
        for sentence in sentences:
            if result and len(result) + len(sentence) > limit:
                break
            result += sentence.strip()
        if result:
            return result.strip()
        truncated = value[:limit].rstrip("，,；;：: ")
        return truncated + ("。" if chinese else ".")

    @staticmethod
    def _write_pptx(pages: list[dict[str, Any]], output_path: Path) -> None:
        import re
        from pptx import Presentation
        from pptx.util import Inches, Pt
        from pptx.dml.color import RGBColor

        presentation = Presentation()
        presentation.slide_width, presentation.slide_height = Inches(40 / 3), Inches(7.5)
        for number, page in enumerate(pages, 1):
            slide = presentation.slides.add_slide(presentation.slide_layouts[5])
            slide.background.fill.solid()
            slide.background.fill.fore_color.rgb = RGBColor.from_string("F5F8F6")
            content = str(page["title"]) + " ".join(map(str, page.get("key_points", [])))
            font = "Noto Sans CJK SC" if re.search(r'[\u3400-\u9fff\uac00-\ud7af]', content) else "Arial"
            def text_box(text, x, y, width, height, size, color, bold=False, shape=None):
                if shape is None:
                    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(width), Inches(height))
                shape.left, shape.top, shape.width, shape.height = Inches(x), Inches(y), Inches(width), Inches(height)
                frame = shape.text_frame
                frame.clear()
                frame.word_wrap = True
                frame.margin_left = frame.margin_right = 0
                paragraph = frame.paragraphs[0]
                paragraph.text = str(text)
                for run in paragraph.runs:
                    run.font.name = font
                    run.font.size, run.font.bold = Pt(size), bold
                    run.font.color.rgb = RGBColor.from_string(color)
            text_box(page["title"], .75, .78, 11.9, 1.25, 32, "153E32", True, slide.shapes.title)
            points = list(page.get("key_points", []))
            spacing = min(.84, 4.65 / max(1, len(points)))
            for index, value in enumerate(points):
                y = 2.2 + index * spacing
                text_box(f"{index + 1:02d}", .78, y + .02, .58, .5, 16, "4D8270", True)
                text_box(value, 1.55, y, 10.95, max(.65, spacing), 23, "29473C")
            text_box(f"{number:02d} / {len(pages):02d}", 11.7, 6.9, .9, .35, 12, "6B8279")
            if page.get("script_text"):
                slide.notes_slide.notes_text_frame.text = page["script_text"]
        output_path.parent.mkdir(parents=True, exist_ok=True)
        presentation.save(str(output_path))


def save_structure(value: dict[str, Any], output_path: Path) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding="utf-8")

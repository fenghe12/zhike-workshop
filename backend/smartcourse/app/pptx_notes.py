"""Write page-aligned speaker notes while checking original slide and media retention."""
from pathlib import Path
import posixpath
import zipfile
from lxml import etree as E
from pptx import Presentation
from pptx.opc.packuri import PackURI


def write_speaker_notes(path, pages):
    path=Path(path)
    with zipfile.ZipFile(path) as archive:
        originals={name:archive.read(name) for name in archive.namelist()}
    relationships={node.get('Id'):node.get('Target') for node in E.fromstring(originals['ppt/_rels/presentation.xml.rels'])}
    source_slide_parts=[]
    for node in E.fromstring(originals['ppt/presentation.xml']).findall('{http://schemas.openxmlformats.org/presentationml/2006/main}sldIdLst/{http://schemas.openxmlformats.org/presentationml/2006/main}sldId'):
        target=relationships[node.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id')]
        source_slide_parts.append(target.lstrip('/') if target.startswith('/') else posixpath.normpath('ppt/'+target))
    presentation=Presentation(str(path))
    if len(presentation.slides)!=len(pages):
        raise ValueError('讲稿页数与课件页数不一致')
    for index,(slide,page,source_part) in enumerate(zip(presentation.slides,pages,source_slide_parts),1):
        if page.get('page_id')!=index or not str(page.get('script_text','')).strip():
            raise ValueError('讲稿缺少有效页码或正文')
        frame=slide.notes_slide.notes_text_frame
        if frame is None:raise ValueError('当前备注页没有可编辑正文区域')
        frame.text=page['script_text']
        slide.part.partname=PackURI('/'+source_part)
    temporary=path.with_suffix('.notes-pending.pptx')
    presentation.save(temporary)
    with zipfile.ZipFile(temporary) as result:
        for name,data in originals.items():
            if name.startswith(('ppt/slides/','ppt/slideMasters/','ppt/slideLayouts/','ppt/theme/')) and name.endswith('.xml'):
                if E.tostring(E.fromstring(data),method='c14n')!=E.tostring(E.fromstring(result.read(name)),method='c14n'):
                    raise ValueError('写入备注时课件页面结构发生变化')
            elif name.startswith('ppt/media/') and not name.endswith('/') and result.read(name)!=data:
                raise ValueError('写入备注时原始媒体内容发生变化')
    temporary.replace(path)
    return len(pages)

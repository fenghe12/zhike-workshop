"""Independently review structurally valid candidates rejected by model criticism."""
import json
import subprocess
from pathlib import Path


def structural_findings(report):
    semantic = {f"第{row['page']}页：{error}" for row in report.get('content_critic', [])
                for error in row.get('errors', [])}
    return [finding for finding in report.get('findings', []) if finding not in semantic]


def recover_candidate(scp, host, remote, destination, request, log):
    from pptx import Presentation
    from .config import Settings
    from .gateway import ResponsesGateway
    from .course_content_review import review_course_content
    from .pptx_render import render_pptx
    root = destination.parent / (destination.stem + '-independent-review')
    root.mkdir(parents=True, exist_ok=True)
    for index in (3, 2, 1):
        report_path = root / f'remote-review-{index}.json'
        fetch = subprocess.run([*scp, host + ':' + remote + f'/deeppresenter/review-{index}.json', str(report_path)],
                               stdout=log, stderr=log, timeout=120)
        if fetch.returncode:
            continue
        report = json.loads(report_path.read_text())
        if (report.get('renderer') != 'LibreOffice' or report.get('slide_count') != request['pages']
                or not report.get('rendered_text') or structural_findings(report)):
            continue
        candidate = root / f'candidate-{index}.pptx'
        subprocess.run([*scp, host + ':' + remote + f'/deeppresenter/branch-design-{index}.pptx', str(candidate)],
                       check=True, stdout=log, stderr=log, timeout=180)
        presentation = Presentation(candidate)
        if len(presentation.slides) != request['pages']:
            continue
        pages = [{'page': number, 'text': '\n'.join(shape.text for shape in slide.shapes if shape.has_text_frame)}
                 for number, slide in enumerate(presentation.slides, 1)]
        requirements = ['课件围绕主题：' + request['topic']]
        if request.get('source_material') is not None:
            requirements.extend(['准确归纳提供的课程页面内容', '遵守用户问题对资料范围和数值的限制'])
        else:
            import re
            requirements.extend(re.split('[，。；\n]+', request.get('requirements') or '解释关键概念并提供必要的例子'))
        result = review_course_content(pages, requirements, ResponsesGateway(Settings.from_env()),
                                       source_context=request.get('source_material'))
        (root / f'content-review-{index}.json').write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding='utf-8')
        if not result['passed']:
            continue
        render_pptx(candidate, root / f'preview-{index}', pdf_name='course.pdf')
        import shutil
        pending = destination.with_suffix('.reviewed-pending.pptx')
        shutil.copy2(candidate, pending)
        pending.replace(destination)
        return {'engine': 'DeepPresenter-9B', 'generated_pptx': str(destination), 'remoteDirectory': remote,
                'validationMode': 'independent_content_review', 'reviewPath': str(root / f'content-review-{index}.json')}
    return None

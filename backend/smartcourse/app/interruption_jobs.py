"""Persistent course questions executed using the course voice and photo."""
import json
import logging
import os
import shutil
import subprocess
import socket
import threading
import uuid
from pathlib import Path

from .interruption import load_interruption_context, generate_interruption_answer
from .task_model import utc_now
from .resource_scheduler import ResourceScheduler


def public_interruption(record):
    fields = ('id', 'courseId', 'status', 'stage', 'question', 'targetLanguage', 'resumePosition',
              'createdAt', 'updatedAt', 'videoUrl', 'audioUrl', 'answer', 'statusUrl', 'message',
              'sourcePages', 'needsClarification', 'slideOutline', 'needsSlides', 'generationTaskId', 'durationSeconds', 'pptUrl')
    result = {key: record[key] for key in fields if key in record}
    if (record.get('purpose') == 'quiz_remediation'
            and isinstance(record.get('videoUrl'),str) and record['videoUrl'].endswith('.mp4')):
        result['posterUrl'] = record['videoUrl'][:-4] + '.png'
    if record.get('needsSlides') and record.get('status') != 'completed':
        result['answer'] = None
        result['message'] = (record.get('message') if record.get('status') == 'failed'
                             else '临时回答课件和视频正在生成，完成后提供回答及下载链接。')
    return result


def compose_photo_answer_video(reference_image, audio, output, duration):
    ffmpeg = os.environ.get('SMARTCOURSE_FFMPEG') or shutil.which('ffmpeg')
    if not ffmpeg:
        raise RuntimeError('回答视频缺少 FFmpeg')
    output = str(output)
    temporary = output + '.tmp.mp4'
    duration = max(0.1, float(duration))
    filter_graph = (
        '[0:v]scale=512:512:force_original_aspect_ratio=decrease,'
        'pad=512:512:(ow-iw)/2:(oh-ih)/2:color=0xE8EEEB[photo];'
        f'color=c=0x102A23:s=1280x720:d={duration}[background];'
        '[background][photo]overlay=720:104:shortest=1[v]'
    )
    subprocess.run([
        ffmpeg, '-y', '-loop', '1', '-i', str(reference_image), '-i', str(audio),
        '-filter_complex', filter_graph, '-map', '[v]', '-map', '1:a:0',
        '-r', '25', '-c:v', 'libx264', '-pix_fmt', 'yuv420p', '-c:a', 'aac',
        '-t', str(duration), '-movflags', '+faststart', temporary,
    ], check=True, capture_output=True, timeout=300)
    os.replace(temporary, output)


def compose_remediation_video(reference_image, audio, output, duration, question, answer,
                              locale='zh-CN', avatar_video=None):
    """Create a readable explanation card while retaining the course teacher photo."""
    from PIL import Image, ImageDraw, ImageFont, ImageOps

    font_path = Path(os.environ.get('SMARTCOURSE_CJK_FONT') or
                     Path(__file__).resolve().parents[2] / 'model-cache/fonts/smartcourse-cn-font.ttc')
    def font(size):
        return ImageFont.truetype(str(font_path), size, index=2)
    labels={
        'zh-CN':('错题讲解','关键概念与正确思路','题目','讲解','课程教师 · 语音讲解','完整讲解请听视频语音。'),
        'en-US':('Answer review','Key idea and correct reasoning','Question','Explanation','Course teacher · spoken explanation','Listen for the complete explanation.'),
        'ja-JP':('誤答の解説','重要な概念と正しい考え方','問題','解説','講師 · 音声解説','続きは音声でご確認ください。'),
        'ko-KR':('오답 해설','핵심 개념과 올바른 풀이','문제','해설','강사 · 음성 해설','전체 해설은 음성으로 들으세요.'),
        'de-DE':('Antwort erklärt','Kernidee und richtiger Lösungsweg','Frage','Erklärung','Lehrkraft · Audioerklärung','Vollständige Erklärung im Ton.'),
        'fr-FR':('Correction','Idée clé et bonne méthode','Question','Explication','Enseignant · explication audio','Écoutez la suite de l’explication.'),
    }.get(locale)
    if labels is None:labels=('Answer review','Key idea and correct reasoning','Question','Explanation','Course teacher','Listen for the complete explanation.')
    canvas = Image.new('RGB', (1280, 720), '#102f2b')
    draw = ImageDraw.Draw(canvas)
    draw.rounded_rectangle((40, 34, 1240, 686), radius=30, fill='#f8faf8')
    label_width=min(620,max(167,int(draw.textlength(labels[0],font(27)))+38))
    draw.rounded_rectangle((68, 58, 68+label_width, 104), radius=18, fill='#daf1e5')
    draw.text((87, 64), labels[0], font=font(27), fill='#156947')
    draw.text((68, 129), labels[1], font=font(35), fill='#173f37')
    draw.line((68, 185, 754, 185), fill='#c8dcd2', width=2)
    def wrap(value, face, width):
        if locale in {'en-US','de-DE','fr-FR','ko-KR'}:
            lines=[];current=''
            for paragraph in str(value).replace('\r','').split('\n'):
                if not paragraph:
                    if current:lines.append(current);current=''
                    continue
                for word in paragraph.split():
                    candidate=(current+' '+word).strip()
                    if current and draw.textlength(candidate,font=face)>width:
                        lines.append(current);current=word
                    else:current=candidate
                if current:lines.append(current);current=''
            return lines
        lines=[]; current=''
        closing=set('，。！？、；：）】》”’')
        for char in str(value).replace('\r',''):
            if char!='\n' and char in closing and current:
                current+=char
                continue
            if char=='\n' or draw.textlength(current+char,font=face)>width:
                if current:lines.append(current)
                current='' if char=='\n' else char
            else:
                current+=char
        if current:lines.append(current)
        return lines
    question_lines=wrap(question,font(22),680)[:3]
    draw.text((68,204),labels[2],font=font(21),fill='#23875e')
    y=238
    for line in question_lines:
        draw.text((68,y),line,font=font(22),fill='#374f48');y+=31
    y+=15
    draw.text((68,y),labels[3],font=font(21),fill='#23875e');y+=35
    available=633-y
    face=None; lines=[]
    for size in (29,27,25,23,21):
        candidate=font(size)
        wrapped=wrap(answer,candidate,680)
        if len(wrapped)*(size+11)<=available:
            face=candidate;lines=wrapped;break
    if face is None:
        face=font(21);lines=wrap(answer,face,680)
        limit=max(1,available//32)
        if len(lines)>limit:lines=lines[:limit-1]+[labels[5]]
    for line in lines:
        draw.text((68,y),line,font=face,fill='#203e35');y+=face.size+11
    with Image.open(reference_image) as source:
        teacher=ImageOps.fit(source.convert('RGB'),(390,495),method=Image.Resampling.LANCZOS)
    canvas.paste(teacher,(817,141))
    draw.rounded_rectangle((813,137,1211,640),radius=18,outline='#d1e3d8',width=5)
    draw.text((836,650),labels[4],font=font(17),fill='#617b70')
    frame=Path(output).with_suffix('.png')
    canvas.save(frame)
    ffmpeg=os.environ.get('SMARTCOURSE_FFMPEG') or shutil.which('ffmpeg')
    if not ffmpeg:raise RuntimeError('错题讲解视频缺少 FFmpeg')
    temporary=str(output)+'.tmp.mp4'
    command=[ffmpeg,'-y','-loop','1','-i',str(frame)]
    if avatar_video:
        command.extend(['-i',str(avatar_video),'-i',str(audio),'-filter_complex',
                        '[1:v]scale=390:495:force_original_aspect_ratio=increase,'
                        'crop=390:495,tpad=stop_mode=clone:stop_duration=2[teacher];'
                        '[0:v][teacher]overlay=817:141:shortest=1[v]',
                        '-map','[v]','-map','2:a:0'])
    else:
        command.extend(['-i',str(audio),'-map','0:v:0','-map','1:a:0'])
    command.extend(['-r','25','-c:v','libx264','-pix_fmt','yuv420p','-c:a','aac',
                    '-t',str(max(.1,float(duration))),'-movflags','+faststart',temporary])
    subprocess.run(command,
                   check=True,capture_output=True,timeout=300)
    os.replace(temporary,output)


def submit_interruption(application, task, payload, metadata=None):
    context = load_interruption_context(task, payload, application.settings.artifact_root)
    identifier = 'answer-' + uuid.uuid4().hex
    record = {'id': identifier, 'courseId': task['id'], 'status': 'queued',
              'stage': 'answering', 'question': context['question'],
              'targetLanguage': context['targetLanguage'],
              'resumePosition': context['resumePosition'], 'createdAt': utc_now(),
              'workerPid': os.getpid(), 'workerHost': socket.gethostname(),
              'videoUrl': None, 'answer': None,
              'statusUrl': f"/api/courses/{task['id']}/interruptions/{identifier}"}
    if metadata:
        record.update(metadata)
    def register(data):
        records = data.setdefault('interruptions', {})
        for existing in records.values():
            if existing.get('courseId') != task['id'] or existing.get('status') not in ('queued', 'running'):
                continue
            if all(existing.get(key) == record[key] for key in ('question', 'targetLanguage', 'resumePosition')):
                return existing
            raise ValueError('当前课程已有问题正在处理，请等待回答完成后再提交')
        records[identifier] = record
        return record
    registered = application.store.transact(register)
    if registered['id'] != identifier:
        return public_interruption(registered)
    worker = threading.Thread(target=execute_interruption, args=(application, task, context, identifier),
                              name=identifier, daemon=True)
    worker.start()
    return public_interruption(record)


def retry_interruption(application, task, identifier):
    def queue(data):
        record = data.get('interruptions', {}).get(identifier)
        if not record or record.get('courseId') != task['id']:
            raise ValueError('回答记录不属于当前课程')
        if record.get('status') != 'failed':
            raise ValueError('仅失败的回答可以重试')
        if any(item.get('courseId') == task['id'] and item.get('status') in ('running','queued')
               for key, item in data.get('interruptions', {}).items() if key != identifier):
            raise ValueError('当前课程已有回答正在处理')
        record.update(status='queued', workerPid=os.getpid(), workerHost=socket.gethostname(), updatedAt=utc_now())
        record.pop('message', None)
        return record
    current = application.store.snapshot().get('interruptions', {}).get(identifier, {})
    if current.get('courseId') != task['id']:
        raise ValueError('回答记录不属于当前课程')
    context = load_interruption_context(task, {'question': current['question'],
        'targetLanguage': current['targetLanguage'], **current['resumePosition']}, application.settings.artifact_root)
    record = application.store.transact(queue)
    threading.Thread(target=execute_interruption, args=(application, task, context, identifier),
                     name=identifier, daemon=True).start()
    return public_interruption(record)


def execute_interruption(application, task, context, identifier):
    def media_call(function, *args, **kwargs):
        with ResourceScheduler(application.settings.data_root).acquire_media('avatar'):
            return function(*args, **kwargs)

    def update(**fields):
        def operation(data):
            data['interruptions'][identifier].update(fields, updatedAt=utc_now())
        application.store.transact(operation)
    root = application.settings.artifact_root / 'tasks' / task['id'] / 'interruptions' / identifier
    try:
        root.mkdir(parents=True, exist_ok=True)
        update(status='running', stage='answering')
        runner = application.pipeline_runner
        previous = application.store.snapshot()['interruptions'][identifier]
        if previous.get('generationTaskId') and previous.get('needsSlides'):
            answer = {key: previous[key] for key in ('answer','needsSlides','slideOutline','sourcePages','needsClarification')}
        else:
            answer = generate_interruption_answer(context, runner.gateway)
        if previous.get('purpose') == 'quiz_remediation':
            answer['needsSlides'] = False
        if answer.get('needsSlides'):
            if previous.get('generationTaskId'):
                child = application.get_task(previous['generationTaskId'])
                if child.get('createdBy') != task['createdBy'] or child.get('inputs',{}).get('interruptionParent') != task['id']:
                    raise ValueError('回答课件任务归属不一致')
            else:
                child = application.add_task({'name': '课程回答：' + context['question'][:80], 'mode': 'topic',
                'voiceId': task.get('voiceId', 'default'), 'targetLanguages': [context['targetLanguage']],
                'autoExecute': False, 'inputs': {'topic': context['question'],
                    'totalPages': len(answer['slideOutline']), 'durationMinutes': len(answer['slideOutline']),
                    'targetLearners': task.get('inputs', {}).get('targetLearners', '高校本科生'),
                    'interruptionParent': task['id'],
                    'answerSourceContext': context,
                    'requirements': '仅依据以下课程回答生成临时课件，不添加无来源结论。回答：' + answer['answer']
                                    + '\n提纲：' + '\n'.join(answer['slideOutline'])}},
                task['createdBy'], validated_assets={key: task['assets'][key] for key in ('referenceAudio', 'referenceImage')})
            update(**answer, status='running', stage='generating_answer_course', generationTaskId=child['id'])
            if child.get('status') == 'completed':
                completed = child
            elif child.get('status') == 'failed':
                completed = application.execute_task(child['id'], start_stage=child.get('currentStage',1), asynchronous=False)
            else:
                completed = application.execute_task(child['id'], asynchronous=False)
            if completed['status'] != 'completed':
                raise RuntimeError(completed.get('error', {}).get('message', '回答课件生成未完成'))
            media = completed['outputs']['s5']['languages'][context['targetLanguage']]
            from pptx import Presentation
            from .course_content_review import review_course_content
            deck = Presentation(runner._resolve_asset(media['coursePptx']))
            pages = [{'page': number, 'text': '\n'.join(shape.text for shape in slide.shapes if shape.has_text_frame)}
                     for number, slide in enumerate(deck.slides, 1)]
            review = review_course_content(pages, '准确回答用户问题；遵守用户对资料范围和数值的限制；保持数学条件准确',
                                           runner.gateway, source_context=context)
            (root/'content-review.json').write_text(json.dumps(review, ensure_ascii=False, indent=2), encoding='utf-8')
            if not review['passed']:
                raise ValueError('回答课件未通过原课程内容复核，需要重新生成')
            update(status='completed', stage='completed', videoUrl=media['video'], pptUrl=media['coursePptx'])
            return
        update(**answer, stage='synthesizing_speech')
        reference_audio = runner._resolve_asset(task.get('assets', {}).get('referenceAudio'))
        reference_image = runner._resolve_asset(task.get('assets', {}).get('referenceImage'))
        audio = root / 'answer.wav'
        if audio.is_file() and audio.stat().st_size > 0:
            import wave
            with wave.open(str(audio)) as stream:
                duration = stream.getnframes() / stream.getframerate()
        else:
            # Use the same IndexTTS2 implementation as the validated S3 stage
            # when CosyVoice is not configured on the production worker.
            with ResourceScheduler(application.settings.data_root).acquire_media('speech'):
                if os.environ.get('SMARTCOURSE_COSYVOICE_ROOT') and os.environ.get('SMARTCOURSE_COSYVOICE_MODEL'):
                    from agents.cosyvoice_engine import CosyVoiceEngine
                    engine = CosyVoiceEngine()
                    try:
                        duration = engine.synthesize(answer['answer'], str(reference_audio), str(audio), context['targetLanguage'])
                    finally:
                        engine.close()
                else:
                    from agents.s3_tts_agent import TTSEngine
                    engine = TTSEngine()
                    try:
                        duration = engine.synthesize(answer['answer'], str(reference_audio), str(audio))
                    finally:
                        getattr(engine, 'close', lambda: None)()
        timeline = root / 'timeline.json'
        timeline.write_text(json.dumps({'pages': [{'page_id': 1, 'wav_path': str(audio),
                                                   'duration': duration, 'script_text': answer['answer']}]},
                                      ensure_ascii=False), encoding='utf-8')
        update(stage='generating_avatar')
        interruption_mode = os.environ.get('SMARTCOURSE_INTERRUPTION_AVATAR_MODE', 'photo').strip().lower()
        if previous.get('purpose') == 'quiz_remediation':
            video = root / 'answer.mp4'
            avatar_video=None
            try:
                from .float_process import run_float_avatar
                avatar_result=media_call(run_float_avatar,
                    task_root=root,runtime_root=root.parent,code=identifier,
                    timeline=timeline,reference_image=reference_image,
                    float_dir=runner.source_root/'float',gpu_id=runner.settings.gpu_id,
                )
                successful=[item for item in avatar_result.get('results',[]) if item.get('video_path')]
                if successful:avatar_video=successful[0]['video_path']
            except Exception:
                logging.exception('Quiz remediation avatar unavailable: %s',identifier)
            compose_remediation_video(reference_image, audio, video, duration,
                                      previous.get('quizPrompt',''), answer['answer'],
                                      context['targetLanguage'],avatar_video=avatar_video)
        elif interruption_mode == 'photo':
            video = root / 'answer.mp4'
            compose_photo_answer_video(reference_image, audio, video, duration)
        else:
            flashhead_command = os.environ.get('SMARTCOURSE_FLASHHEAD_COMMAND')
            if os.environ.get('SMARTCOURSE_REMOTE_MEDIA') == '1' or (flashhead_command and os.path.isfile(flashhead_command)):
                from agents.flashhead_agent import FlashHeadAvatarAgent
                output = media_call(FlashHeadAvatarAgent(str(root), str(reference_image)).run, 'answer', str(timeline))
                video = output['pages'][0]['video_path']
            else:
                from .float_process import run_float_avatar
                output = media_call(run_float_avatar,
                    task_root=root,
                    runtime_root=root.parent,
                    code=identifier,
                    timeline=timeline,
                    reference_image=reference_image,
                    float_dir=runner.source_root / 'float',
                    gpu_id=runner.settings.gpu_id,
                )
                successful = [item for item in output.get('results', []) if item.get('video_path')]
                if not successful:
                    raise RuntimeError('回答视频未生成有效数字人页面')
                video = successful[0]['video_path']
        update(status='completed', stage='completed', videoUrl=runner._url(video),
               audioUrl=runner._url(audio), durationSeconds=duration)
    except Exception:
        logging.exception('Course interruption failed: %s', identifier)
        update(status='failed', message='回答生成未完成，请重试；原课程播放位置已保存。')

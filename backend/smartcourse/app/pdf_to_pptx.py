"""Convert PDF pages to layout-preserving PPTX with editable XML text."""
import json
import re
from pathlib import Path
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.util import Pt
from .pdf_input import inspect_pdf_course


def _visual_text_lines(page):
    """Rebuild editable rows by their painted positions, not PDF stream order."""
    spans=[]
    for block in page.get_text('dict')['blocks']:
        for line in block.get('lines',[]):
            if tuple(line.get('dir',(1,0))) != (1,0):
                raise ValueError('PDF 包含旋转文字，需要专门处理后才能生成可编辑课件')
            for span in line['spans']:
                if span['text'].strip() and span['text'].strip() not in {'','➢','•','▪','✓'}:
                    spans.append(span)
    spans.sort(key=lambda item: ((item['bbox'][1]+item['bbox'][3])/2,item['bbox'][0]))
    rows=[]
    for span in spans:
        center=(span['bbox'][1]+span['bbox'][3])/2
        row=next((item for item in reversed(rows) if abs(item['center']-center)<=max(6.5,min(item['size'],span['size'])*.4)),None)
        if row is None:
            row={'center':center,'size':span['size'],'spans':[]}
            rows.append(row)
        row['spans'].append(span)
    result=[]
    for row in rows:
        ordered=sorted(row['spans'],key=lambda item:item['bbox'][0])
        groups=[]
        for item in ordered:
            if not groups or item['bbox'][0]-groups[-1][-1]['bbox'][2] > max(item['size'],groups[-1][-1]['size'])*1.8:
                groups.append([])
            groups[-1].append(item)
        for items in groups:
            result.append(_line_record(items))
    return result


def _line_record(items):
    x0=min(item['bbox'][0] for item in items)
    y0=min(item['bbox'][1] for item in items)
    x1=max(item['bbox'][2] for item in items)
    y1=max(item['bbox'][3] for item in items)
    text=''.join(item['text'] for item in items).strip()
    color=max(items,key=lambda item:len(item['text']))['color']
    return {'bbox':(x0,y0,x1,y1),'text':text,
            'preserveAsImage':bool(re.search(r'[Σ∑σ∫=≤≥]',text)),
            'fontSize':max(item['size'] for item in items),
            'color':[(color>>shift&255)/255 for shift in (16,8,0)]}


def convert_pdf_to_pptx(source, destination, workspace, font_file):
    import pymupdf
    source,destination,workspace=Path(source).resolve(),Path(destination).resolve(),Path(workspace).resolve()
    workspace.mkdir(parents=True,exist_ok=True)
    course=inspect_pdf_course(source)
    if any(page['requiresOcr'] for page in course['pages']):
        raise ValueError('PDF 含扫描页面，须先完成文字识别才能生成可编辑课件')
    backgrounds=workspace/'backgrounds';backgrounds.mkdir(exist_ok=True)
    pages=[]
    formula_count=0
    with pymupdf.open(source) as document:
        first=document[0]
        width,height=first.cropbox.width*4/3,first.cropbox.height*4/3
        for number,page in enumerate(document,1):
            regions=_visual_text_lines(page)
            formula_count+=sum(region['preserveAsImage'] for region in regions)
            if abs(page.cropbox.width*4/3-width)>.5 or abs(page.cropbox.height*4/3-height)>.5:
                raise ValueError('PDF 页面尺寸不一致，无法生成统一课件')
            pix=page.get_pixmap(matrix=pymupdf.Matrix(2,2),alpha=False)
            scale_x,scale_y=pix.width/page.rect.width,pix.height/page.rect.height
            def local_fill(rect):
                points=[(rect.x0-2,rect.y0-2),(rect.x1+2,rect.y0-2),(rect.x0-2,rect.y1+2),(rect.x1+2,rect.y1+2),
                        ((rect.x0+rect.x1)/2,rect.y0-2),((rect.x0+rect.x1)/2,rect.y1+2)]
                colors=[]
                for x,y in points:
                    px=max(0,min(pix.width-1,round(x*scale_x)));py=max(0,min(pix.height-1,round(y*scale_y)))
                    colors.append(pix.pixel(px,py)[:3])
                return tuple(sorted(color[channel] for color in colors)[len(colors)//2]/255 for channel in range(3))
            for region in regions:
                if region['preserveAsImage']:continue
                rect=pymupdf.Rect(region['bbox'])
                page.add_redact_annot(rect,fill=local_fill(rect),cross_out=False)
            page.apply_redactions(images=0,graphics=0,text=0)
            background=backgrounds/f'page-{number:03d}.png'
            page.get_pixmap(matrix=pymupdf.Matrix(2,2),alpha=False).save(background)
            rows=[]
            for index,region in enumerate(regions,1):
                if region['preserveAsImage']:continue
                rect=pymupdf.Rect(region['bbox'])
                text=region['text']
                if not text:continue
                color=region.get('color',(0,0,0))
                rows.append({'index':index,'text':text,'fontSizePt':max(7,float(region['fontSize'])),
                    'color':'#'+''.join(f'{max(0,min(255,round(value*255))):02X}' for value in color),
                    'position':{'left':rect.x0*4/3,'top':rect.y0*4/3,
                                'width':max(8,(rect.width+max(4,rect.width*.04))*4/3),
                                'height':max(8,(rect.height+2)*4/3)}})
            pages.append({'page':number,'background':str(background),'regions':rows})
    manifest=workspace/'conversion.json'
    manifest.write_text(json.dumps({'width':width,'height':height,'fontFamily':'Noto Sans CJK SC','pages':pages},ensure_ascii=False))
    deck=Presentation()
    deck.slide_width=Pt(width*.75)
    deck.slide_height=Pt(height*.75)
    for page in pages:
        slide=deck.slides.add_slide(deck.slide_layouts[6])
        slide.shapes.add_picture(page['background'],Pt(0),Pt(0),width=deck.slide_width,height=deck.slide_height)
        for region in page['regions']:
            position=region['position']
            shape=slide.shapes.add_textbox(
                Pt(position['left']*.75),Pt(position['top']*.75),
                Pt(position['width']*.75),Pt(position['height']*.75),
            )
            shape.name=f"pdf-text-{page['page']}-{region['index']}"
            frame=shape.text_frame
            frame.text=region['text']
            frame.margin_left=frame.margin_right=0
            frame.margin_top=frame.margin_bottom=0
            frame.word_wrap=False
            frame.auto_size=MSO_AUTO_SIZE.NONE
            color=region['color'].lstrip('#')
            rgb=RGBColor(*(int(color[offset:offset+2],16) for offset in (0,2,4)))
            for paragraph in frame.paragraphs:
                for run in paragraph.runs:
                    run.font.name='Noto Sans CJK SC'
                    run.font.size=Pt(region['fontSizePt'])
                    run.font.color.rgb=rgb
    pending=destination.with_suffix('.pending.pptx')
    deck.save(pending)
    pending.replace(destination)
    deck=Presentation(destination)
    if len(deck.slides)!=course['pageCount']:raise ValueError('PDF 转换后的 PPTX 页数不一致')
    return {'pptx':str(destination),'pageCount':len(deck.slides),
            'editableTextRegions':sum(len(page['regions']) for page in pages),
            'formulaRegionsPreservedAsImage':formula_count,
            'pagesRequiringImageTextInspection':[p['page_id'] for p in course['pages'] if p['imagesRequireTextInspection']]}


def restore_source_pdf_slide_images(source_pdf, image_paths):
    """Use original PDF pixels for source-language video while retaining editable PPTX text."""
    import os
    import tempfile
    import pymupdf
    from PIL import Image

    paths = [Path(path) for path in image_paths]
    with pymupdf.open(source_pdf) as document:
        if len(document) != len(paths):
            raise ValueError('PDF 页数与视频课件页数不一致')
        for page, path in zip(document, paths):
            if not path.is_file():
                raise FileNotFoundError(path)
            with Image.open(path) as current:
                width, height = current.size
            scale = pymupdf.Matrix(width / page.rect.width, height / page.rect.height)
            pixmap = page.get_pixmap(matrix=scale, alpha=False)
            descriptor, temporary_name = tempfile.mkstemp(prefix=path.stem + '.', suffix='.png', dir=path.parent)
            os.close(descriptor)
            temporary = Path(temporary_name)
            try:
                if pixmap.width == width and pixmap.height == height:
                    pixmap.save(temporary)
                else:
                    with Image.frombytes('RGB', (pixmap.width, pixmap.height), pixmap.samples) as image:
                        image.resize((width, height), Image.Resampling.LANCZOS).save(temporary)
                temporary.replace(path)
            finally:
                temporary.unlink(missing_ok=True)
    return paths

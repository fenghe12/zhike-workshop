"""Create versioned course tasks for approved page-level revisions."""

import copy
import json
import re
import shutil
import uuid
from pathlib import Path
from pptx import Presentation

from .narration import export_lecture_text, validate_spoken_text
from .pptx_notes import write_speaker_notes
from .task_model import utc_now
from .teacher_features import _transaction


def _page_number(value):
    text = str(value)
    if text.startswith("src_page_"):
        text = text[9:]
    elif text.startswith("page_"):
        text = text[5:]
    if not text.isdigit() or int(text) < 1:
        raise ValueError("局部修订的页面编号无效")
    return int(text)


def _replace_task_paths(value, base_id, release_id, base_root, release_root):
    if isinstance(value, str):
        return value.replace(str(base_root), str(release_root)).replace(
            f"/artifacts/tasks/{base_id}/", f"/artifacts/tasks/{release_id}/"
        )
    if isinstance(value, list):
        return [_replace_task_paths(item, base_id, release_id, base_root, release_root) for item in value]
    if isinstance(value, dict):
        return {key: _replace_task_paths(item, base_id, release_id, base_root, release_root)
                for key, item in value.items()}
    return value


def _reserve(settings, plan_id):
    reservation = uuid.uuid4().hex
    with _transaction(settings) as data:
        plan = data.get("patchPlans", {}).get(plan_id)
        if not plan:
            raise ValueError("局部修订方案不存在")
        if plan.get("status") != "APPROVED":
            raise ValueError("局部修订方案尚未批准")
        if plan.get("releaseTaskId"):
            return copy.deepcopy(plan), None
        if plan.get("releaseStatus") == "PREPARING":
            raise ValueError("修订版任务正在准备中")
        plan.update(releaseStatus="PREPARING", releaseRequestId=reservation,
                    releaseStartedAt=utc_now(), updatedAt=utc_now())
        return copy.deepcopy(plan), reservation


def _finish(settings, plan_id, reservation, release_id=None, error=None):
    with _transaction(settings) as data:
        plan = data["patchPlans"][plan_id]
        if plan.get("releaseRequestId") != reservation:
            raise RuntimeError("修订版任务状态已被其他请求修改")
        plan.pop("releaseRequestId", None)
        plan.update(releaseStatus="QUEUED" if release_id else "FAILED", updatedAt=utc_now())
        if release_id:
            plan["releaseTaskId"] = release_id
            plan.pop("releaseError", None)
        else:
            plan["releaseError"] = str(error or "任务准备失败")


def start_script_release(application, plan_id, user_id):
    settings = application.settings
    plan, reservation = _reserve(settings, plan_id)
    if reservation is None:
        task = application.get_task(plan["releaseTaskId"])
        return {"patchPlanId": plan_id, "releaseTaskId": task["id"], "status": task["status"],
                "created": False}
    release_id = None
    try:
        base = application.get_task(plan["course_id"])
        if (user_id not in {base.get("createdBy"), *base.get('teacherDemoUsers',[])}
                or base.get("status") != "completed"):
            raise ValueError("原课程尚未完成或不属于当前用户")
        if base.get("mode") not in {"upload", "topic"}:
            raise ValueError("当前仅支持上传和主题课程的讲稿修订")
        changes = plan.get("changes") or []
        if not changes or any(item.get("kind") not in {"script", "slide_text"} for item in changes):
            raise ValueError("当前可生成讲稿和课件文字修订；其他类型尚未开放生成")
        replacements = {}
        slide_edits = {}
        for item in changes:
            number = _page_number(item["page_id"])
            if item['kind']=='script':
                if number in replacements:
                    raise ValueError("同一页面重复设置了讲稿")
                replacements[number] = validate_spoken_text(item.get("text"))
            else:
                old_text=str(item.get('oldText') or '').strip()
                new_text=str(item.get('text') or '').strip()
                if not old_text or not new_text or old_text==new_text:
                    raise ValueError('课件文字修订内容无效')
                slide_edits.setdefault(number,[]).append((old_text,new_text))
        payload = {"name": base["name"] + " · 修订版", "mode": base["mode"],
                   "inputs": copy.deepcopy(base["inputs"]),
                   "sourceLanguage": base["sourceLanguage"],
                   "targetLanguages": list(base["targetLanguages"]),
                   "voiceId": base.get("voiceId", "default"), "autoExecute": False}
        release = application.add_task(payload, user_id)
        release_id = release["id"]
        artifact_root = Path(settings.artifact_root)
        base_root = artifact_root / "tasks" / base["id"]
        release_root = artifact_root / "tasks" / release_id
        shutil.copytree(base_root / "s1", release_root / "s1")
        if slide_edits and (release_root/'s1/source.pdf').exists():
            raise ValueError('PDF 来源课程的课件文字需要在原 PDF 中修订后重新上传')
        structure_path = release_root / "s1/course_structure.json"
        structure = _replace_task_paths(json.loads(structure_path.read_text(encoding="utf-8")),
                                        base["id"], release_id, base_root, release_root)
        pages = {int(page["page_id"]): page for page in structure.get("pages", [])}
        if not (set(replacements)|set(slide_edits)) <= set(pages):
            raise ValueError("修订方案引用了不存在的课程页面")
        for number, text in replacements.items():
            pages[number].update(script_text=text, script_source="teacher_patch")
        if slide_edits:
            deck_path=release_root/'s1/generated.pptx'
            deck=Presentation(deck_path)
            for number, edits in slide_edits.items():
                slide=deck.slides[number-1]
                for old_text,new_text in edits:
                    found=0
                    for shape in slide.shapes:
                        if not shape.has_text_frame:continue
                        for paragraph in shape.text_frame.paragraphs:
                            for run in paragraph.runs:
                                if old_text in run.text:
                                    found+=run.text.count(old_text)
                                    run.text=run.text.replace(old_text,new_text)
                    if found!=1:
                        raise ValueError(f'第 {number} 页待修订文字须在一个文本片段中恰好出现一次；实际找到 {found} 次')
                    page=pages[number]
                    for field in ('title','key_points'):
                        value=page.get(field)
                        if isinstance(value,str):page[field]=value.replace(old_text,new_text)
                        elif isinstance(value,list):page[field]=[x.replace(old_text,new_text) if isinstance(x,str) else x for x in value]
            temporary=deck_path.with_suffix('.patch-pending.pptx')
            deck.save(temporary)
            temporary.replace(deck_path)
        structure_path.write_text(json.dumps(structure, ensure_ascii=False, indent=2), encoding="utf-8")
        narration_path = release_root / "s1/source_narration.json"
        if narration_path.is_file():
            narration_path.write_text(json.dumps(structure, ensure_ascii=False, indent=2), encoding="utf-8")
        export_lecture_text(structure, release_root / "s1/lecture.txt")
        write_speaker_notes(release_root / "s1/generated.pptx", structure["pages"])
        assets = {}
        references = release_root / "references"
        references.mkdir(parents=True, exist_ok=True)
        for role in ("referenceAudio", "referenceImage"):
            existing = base.get("assets", {}).get(role)
            if not existing:
                raise ValueError("原课程缺少参考声音或照片")
            source = application.pipeline_runner._resolve_asset(existing)
            destination = references / (role + source.suffix.lower())
            shutil.copy2(source, destination)
            path = "/artifacts/" + str(destination.relative_to(artifact_root))
            assets[role] = {**existing, "path": path, "url": path}
        s1_output = _replace_task_paths(copy.deepcopy(base["outputs"]["s1"]),
                                        base["id"], release_id, base_root, release_root)
        def bind(task, _data):
            task["assets"] = assets
            task["voiceId"] = base.get("voiceId", "default")
            task["translationPolicy"] = copy.deepcopy(base.get("translationPolicy", {}))
            task["patchPlanId"] = plan_id
            task["baseTaskId"] = base["id"]
            task["outputs"]["s1"] = s1_output
            task["stages"][0].update(status="completed", progress=100, completedAt=utc_now())
            task.update(currentStage=2, overallProgress=20, updatedAt=utc_now())
            return task
        application.mutate_task(release_id, bind)
        _finish(settings, plan_id, reservation, release_id=release_id)
        application.execute_task(release_id, start_stage=2, asynchronous=True)
        return {"patchPlanId": plan_id, "releaseTaskId": release_id,
                "status": "processing", "created": True}
    except Exception as error:
        if release_id:
            def fail(task, _data):
                task.update(status="failed", error={"message": str(error), "stage": 1,
                                                    "time": utc_now()}, updatedAt=utc_now())
                return task
            application.mutate_task(release_id, fail)
        with _transaction(settings) as data:
            saved = data["patchPlans"][plan_id]
            if saved.get("releaseRequestId") == reservation:
                saved.pop("releaseRequestId", None)
            saved.update(releaseStatus="FAILED", releaseError=str(error), updatedAt=utc_now())
        raise


def _subtitle_seconds(value):
    match=re.fullmatch(r'(\d{2}):(\d{2}):(\d{2}),(\d{3})',value.strip())
    if not match:raise ValueError('字幕时间格式无效')
    hours,minutes,seconds,milliseconds=map(int,match.groups())
    return hours*3600+minutes*60+seconds+milliseconds/1000


def _edit_page_subtitle(path, timeline_path, page_number, old_text, new_text):
    timeline=json.loads(Path(timeline_path).read_text(encoding='utf-8'))
    page=next((item for item in timeline.get('pages',[]) if int(item['page_id'])==page_number),None)
    if page is None:raise ValueError('字幕修订引用了不存在的时间轴页面')
    start=float(page.get('start_time',page.get('start')))
    end=float(page.get('end_time',page.get('end')))
    blocks=re.split(r'\n\s*\n',Path(path).read_text(encoding='utf-8').strip())
    changes=0;updated=[]
    for block in blocks:
        lines=block.splitlines()
        if len(lines)<3 or ' --> ' not in lines[1]:
            updated.append(block);continue
        left,right=lines[1].split(' --> ',1)
        midpoint=(_subtitle_seconds(left)+_subtitle_seconds(right))/2
        if start-.25<=midpoint<=end+.25:
            body='\n'.join(lines[2:])
            changes+=body.count(old_text)
            lines[2:]=body.replace(old_text,new_text).split('\n')
        updated.append('\n'.join(lines))
    if changes<1:raise ValueError(f'第 {page_number} 页字幕中没有找到待修订原文')
    Path(path).write_text('\n\n'.join(updated)+'\n',encoding='utf-8')
    return changes


def start_subtitle_release(application,plan_id,user_id):
    settings=application.settings
    plan,reservation=_reserve(settings,plan_id)
    if reservation is None:
        task=application.get_task(plan['releaseTaskId'])
        return {'patchPlanId':plan_id,'releaseTaskId':task['id'],
                'status':task['status'],'created':False}
    release_id=None
    try:
        base=application.get_task(plan['course_id'])
        if (user_id not in {base.get('createdBy'),*base.get('teacherDemoUsers',[])}
                or base.get('status')!='completed'):
            raise ValueError('原课程尚未完成或当前用户无权修订')
        changes=plan.get('changes') or []
        if not changes or any(item.get('kind')!='subtitle' for item in changes):
            raise ValueError('字幕修订不能与其他修订类型混合执行')
        codes=set(base.get('outputs',{}).get('s5',{}).get('languages',{}))
        if any(item.get('locale') not in codes for item in changes):
            raise ValueError('字幕语种不属于原课程')
        payload={'name':base['name']+' · 字幕修订版','mode':base['mode'],
                 'inputs':copy.deepcopy(base['inputs']),
                 'sourceLanguage':base['sourceLanguage'],
                 'targetLanguages':list(base['targetLanguages']),
                 'voiceId':base.get('voiceId','default'),'autoExecute':False}
        release=application.add_task(payload,user_id)
        release_id=release['id']
        artifact_root=Path(settings.artifact_root)
        base_root=artifact_root/'tasks'/base['id']
        release_root=artifact_root/'tasks'/release_id
        for directory in ('s1','s2','s3'):
            source=base_root/directory
            if not source.is_dir():raise ValueError('原课程缺少阶段文件：'+directory)
            shutil.copytree(source,release_root/directory)
        for code in codes:
            for directory in ('s3_output','s4_output'):
                source=base_root/'runtime'/code/directory
                if not source.is_dir():raise ValueError('原课程缺少阶段文件：'+code+'/'+directory)
                shutil.copytree(source,release_root/'runtime'/code/directory)
        for path in release_root.rglob('*.json'):
            value=json.loads(path.read_text(encoding='utf-8'))
            updated=_replace_task_paths(value,base['id'],release_id,base_root,release_root)
            if updated!=value:path.write_text(json.dumps(updated,ensure_ascii=False,indent=2),encoding='utf-8')
        for item in changes:
            number=_page_number(item['page_id'])
            code=item['locale']
            workspace=release_root/'runtime'/code/'s3_output'
            _edit_page_subtitle(workspace/'subtitles.srt',workspace/'timeline.json',number,
                                str(item.get('oldText') or '').strip(),
                                str(item.get('text') or '').strip())
        assets={};references=release_root/'references';references.mkdir(parents=True,exist_ok=True)
        for role in ('referenceAudio','referenceImage'):
            existing=base.get('assets',{}).get(role)
            if not existing:raise ValueError('原课程缺少参考声音或照片')
            source=application.pipeline_runner._resolve_asset(existing)
            destination=references/(role+source.suffix.lower())
            shutil.copy2(source,destination)
            path='/artifacts/'+str(destination.relative_to(artifact_root))
            assets[role]={**existing,'path':path,'url':path}
        old_outputs=base.get('outputs',{})
        outputs={key:_replace_task_paths(copy.deepcopy(old_outputs[key]),base['id'],release_id,
                                         base_root,release_root)
                 for key in ('s1','s2','s3','s4') if key in old_outputs}
        if len(outputs)!=4:raise ValueError('原课程阶段结果不完整')
        def bind(task,_data):
            task['assets']=assets
            task['voiceId']=base.get('voiceId','default')
            task['translationPolicy']=copy.deepcopy(base.get('translationPolicy',{}))
            task['patchPlanId']=plan_id
            task['baseTaskId']=base['id']
            task['outputs'].update(outputs)
            for stage in task['stages'][:4]:
                stage.update(status='completed',progress=100,completedAt=utc_now())
            task.update(currentStage=5,overallProgress=80,updatedAt=utc_now())
            return task
        application.mutate_task(release_id,bind)
        _finish(settings,plan_id,reservation,release_id=release_id)
        application.execute_task(release_id,start_stage=5,asynchronous=True)
        return {'patchPlanId':plan_id,'releaseTaskId':release_id,
                'status':'processing','created':True}
    except Exception as error:
        if release_id:
            def fail(task,_data):
                task.update(status='failed',error={'message':str(error),'stage':5,
                                                   'time':utc_now()},updatedAt=utc_now())
                return task
            application.mutate_task(release_id,fail)
        with _transaction(settings) as data:
            saved=data['patchPlans'][plan_id]
            if saved.get('releaseRequestId')==reservation:saved.pop('releaseRequestId',None)
            saved.update(releaseStatus='FAILED',releaseError=str(error),updatedAt=utc_now())
        raise

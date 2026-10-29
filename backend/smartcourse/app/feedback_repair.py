"""Persist feedback repair jobs and create reviewable, isolated course revisions."""
import json
import os
import socket
import threading
from pathlib import Path

from .course_access import require_course
from .course_studio import LANGUAGE_NAMES, identifier
from .course_revisions import create_revision
from .course_review import load_course
from .narration import validate_spoken_text
from .task_model import utc_now

_workers = set()
_guard = threading.Lock()


def submit_repair(app, task, owner, payload):
    require_course(task, owner, 'edit')
    feedback = app.store.snapshot().get('userFeedback', {}).get(payload.get('feedbackId'))
    if not feedback or feedback.get('taskId') != task['id']:
        raise ValueError('反馈不属于当前课程。')
    page = payload.get('page') or feedback.get('page')
    pages = {p['page_id'] for p in load_course(app.studio.root(task['id'])).get('pages', [])}
    if isinstance(page,bool) or page not in pages:
        raise ValueError('请为反馈选择需要修正的实际页码。')
    language = payload.get('language') or feedback.get('language')
    if language not in task.get('outputs',{}).get('s2',{}).get('languages',{}):
        raise ValueError('请为反馈选择课程中已生成的语种。')
    categories = feedback.get('categories', [])
    kind = payload.get('kind') or next((x for x in ('translation','narration','speech','subtitles','avatar','layout') if x in categories), None)
    if kind not in {'translation','narration','speech','subtitles','avatar','layout'}:
        raise ValueError('此类反馈需要在后台填写处理说明。')
    def reserve(data):
        repairs = data.setdefault('courseStudio',{}).setdefault('repairs',{})
        previous = next((x for x in repairs.values() if x['feedbackId'] == feedback['id'] and x['courseId'] == task['id']), None)
        if previous:
            return previous
        record = {'id':identifier('repair'), 'feedbackId':feedback['id'], 'courseId':task['id'], 'createdBy':owner,
                  'kind':kind,'page':page,'language':language,'status':'queued','attempts':0,'createdAt':utc_now(),
                  'correction':payload.get('correction',''), 'oldText':payload.get('oldText','')}
        repairs[record['id']] = record
        data['userFeedback'][feedback['id']].update(status='reviewing',repairId=record['id'])
        return record
    record = app.store.transact(reserve)
    if not record.get('releaseTaskId') and record.get('status') in {'queued','working'}:
        start_worker(app,record['id'])
    return app.studio.repair_status(record)


def start_worker(app, job_id):
    with _guard:
        if job_id in _workers:
            return
        _workers.add(job_id)
    def run():
        try:
            _work(app, job_id)
        except Exception as error:
            record = app.studio.state().get('repairs',{}).get(job_id)
            if record:
                app.studio.save('repairs',job_id,{**record,'status':'failed','message':str(error)[:1000],'updatedAt':utc_now()})
        finally:
            with _guard:
                _workers.discard(job_id)
    threading.Thread(target=run,name='course-feedback-'+job_id,daemon=True).start()


def recover_repairs(app):
    from .task_recovery import process_alive
    for record in app.studio.state().get('repairs',{}).values():
        if record.get('status') not in {'queued','working'} or record.get('releaseTaskId'):
            continue
        if record.get('workerHost') == socket.gethostname() and process_alive(record.get('workerPid')) and record.get('workerPid') != os.getpid():
            continue
        start_worker(app,record['id'])


def _work(app, job_id):
    import fcntl
    folder = app.settings.data_root/'feedback-repairs'
    folder.mkdir(parents=True,exist_ok=True)
    with (folder/(job_id+'.lock')).open('a') as lock:
        try:
            fcntl.flock(lock,fcntl.LOCK_EX|fcntl.LOCK_NB)
        except BlockingIOError:
            return
        record = app.studio.state()['repairs'][job_id]
        if record.get('releaseTaskId'):
            return
        task = app.studio.task(record['courseId'],record['createdBy'],'edit')
        feedback = app.store.snapshot()['userFeedback'][record['feedbackId']]
        record.update(status='working',workerPid=os.getpid(),workerHost=socket.gethostname(),updatedAt=utc_now())
        app.studio.save('repairs',job_id,record)
        kind, page, code = record['kind'], record['page'], record['language']
        if kind in {'speech','avatar'}:
            changes = [{'page':page,'language':code,'kind':kind}]
        elif kind == 'layout':
            changes = [{'page':page,'language':code,'kind':'layout'}]
        elif kind == 'subtitles':
            if not record.get('oldText') or not record.get('correction'):
                raise ValueError('字幕反馈需要提供原文和正确文字。')
            changes = [{'page':page,'language':code,'kind':'subtitle','oldText':record['oldText'],'text':record['correction']}]
        else:
            original = next(p for p in load_course(app.studio.root(task['id']))['pages'] if p['page_id']==page)
            translated = next(p for p in load_course(app.studio.root(task['id']),code)['pages'] if p['page_id']==page)
            slide_texts = []
            deck_path = app.studio.root(task['id'])/'s2'/code/'translated.pptx'
            if deck_path.is_file():
                from pptx import Presentation
                slide_texts = [shape.text for shape in Presentation(deck_path).slides[page-1].shapes if shape.has_text_frame]
            prompt = ('Correct this page using the original source and teacher feedback. '
                      'Return JSON {"script_text":"complete corrected narration",'
                      '"replacements":[{"oldText":"exact existing translated slide text","text":"corrected translation"}]}. '
                      'Use the requested target language. Preserve every substantive source claim, number and condition. '
                      'Explain directly to students, without reporting what the slide shows. '
                      'Do not change unrelated content or invent facts. Return an empty replacements list if slide text is already correct.\n'
                      +json.dumps({'source':original,'current':translated,'slideTexts':slide_texts,
                                   'language':LANGUAGE_NAMES[code],'feedback':feedback['comment'],
                                   'requestedCorrection':record.get('correction','')},ensure_ascii=False))
            last_error = '模型未返回有效修订。'
            for attempt in range(record.get('attempts',0),4):
                record['attempts']=attempt+1
                app.studio.save('repairs',job_id,record)
                try:
                    resource = app.pipeline_runner.resources.STAGES['S2'][0]
                    with app.pipeline_runner.resources.acquire(resource):
                        folder_images=app.studio.root(task['id'])/'s1/source-slide-images'
                        images=sorted(folder_images.glob('*.png')) if folder_images.is_dir() else []
                        answer = app.pipeline_runner.gateway.respond_json(prompt+'\nPrevious validation error: '+last_error,
                            json_retries=0,json_record_dir=folder/job_id/f'attempt-{attempt+1}',max_output_tokens=5000,
                            **({'image_paths':[images[page-1]]} if len(images)>=page else {}))
                    script = validate_spoken_text(answer.get('script_text'))
                    from .language_validation import check_target_language
                    if check_target_language(script,LANGUAGE_NAMES[code])['status']=='mismatch':
                        raise ValueError('修订讲稿的语种不符合目标语种。')
                    changes = [{'page':page,'language':code,'kind':'script','text':script}]
                    replacements = answer.get('replacements',[])
                    if not isinstance(replacements,list) or len(replacements)>40:
                        raise ValueError('课件文字修订格式无效。')
                    for replacement in replacements:
                        old, new = replacement.get('oldText'), replacement.get('text')
                        if not isinstance(old,str) or not old or not isinstance(new,str) or not new or not any(old in x for x in slide_texts):
                            raise ValueError('课件文字修订未引用实际译文。')
                        changes.append({'page':page,'language':code,'kind':'slide_text','oldText':old,'text':new})
                    break
                except Exception as error:
                    last_error = str(error)
            else:
                raise ValueError('局部修订达到 3 次重试上限：'+last_error)
        release = create_revision(app,task,record['createdBy'],changes,execute=False,feedback_id=feedback['id'])
        record.update(releaseTaskId=release['task']['id'],status='generating',updatedAt=utc_now())
        app.studio.save('repairs',job_id,record)
        app.execute_task(release['task']['id'],start_stage=release['startStage'],asynchronous=True)

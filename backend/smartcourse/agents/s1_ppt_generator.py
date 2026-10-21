"""
S1 PPT生成器 (PPTGenerator) v5
- ERNIE-Image本地权重生图
- 百度/Bing搜图 + CLIP排序 + 生图兜底
- 文字重叠修复（动态高度+强制间距）
- 四角预留数字人空间
- 布局变化/多语言/slide_body 150字
"""

import os
# ==================== 统一配置加载 ====================
import yaml as _yaml
_AGENT_DIR = os.path.dirname(os.path.abspath(__file__))
with open(os.path.join(_AGENT_DIR, "config.yaml"), "r") as _cf:
    _CFG = _yaml.safe_load(_cf)

import os
import re
import io
import json
import uuid
import httpx
import asyncio
import hashlib
import random
import logging

try:
    from .device import select_device
except ImportError:
    from device import select_device
from datetime import datetime
from typing import List, Dict, Any, Optional, Tuple

try:
    from .s1_courseware_skill import S1CoursewareSkill
except ImportError:
    from s1_courseware_skill import S1CoursewareSkill

try:
    from ..rag.service import CourseRAGService
except ImportError:
    from rag.service import CourseRAGService

from PIL import Image as PILImage
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

# -----------------------------
# Config
# -----------------------------
MODEL = "kivy-qwen3_6-plus"
API_URL = f"http://api-hub.inner.chj.cloud/llm-gateway/v1/chat/completions/{MODEL}"
MAX_TOKENS = 8192
ENABLE_THINKING = False

GW_TOKEN = os.environ.get(_CFG['api'].get('gw_token_env', 'SMARTCOURSE_GW_TOKEN'), '')

# ERNIE-Image local weights
IMAGE_GEN_MODEL = _CFG['paths'].get('image_gen_model', '')
IMAGE_GEN_STEPS = 50
IMAGE_GEN_GUIDANCE = 4.0

CLIP_THRESHOLD = 0.75
BAIDU_CANDIDATES = 10

REQUEST_TIMEOUT_SEC = 90.0
IMAGE_DOWNLOAD_TIMEOUT = 15.0
MAX_RETRIES = 5
BACKOFF_BASE_SEC = 1.5
MAX_BACKOFF_SEC = 12.0

WORKSPACE_ROOT = _CFG['paths']['workspace']

# PPT colors
BG_COLOR = RGBColor(0x1e, 0x3a, 0x2f)
TITLE_COLOR = RGBColor(0xFF, 0xFF, 0xFF)
BULLET_COLOR = RGBColor(0xF5, 0xF0, 0xD0)
BODY_COLOR = RGBColor(0xE0, 0xE0, 0xE0)
SUBTITLE_COLOR = RGBColor(0xAA, 0xAA, 0xAA)
ACCENT_COLOR = RGBColor(0x4E, 0xC9, 0xB0)
DIVIDER_COLOR = RGBColor(0x3a, 0x5a, 0x4f)

# PPT dimensions (16:9)
SLIDE_WIDTH = Inches(13.333)
SLIDE_HEIGHT = Inches(7.5)

# Corner reservation for digital human avatar
SAFE_X_MIN = 0.5
SAFE_X_MAX = 12.8
SAFE_Y_MAX = 5.0

logging.basicConfig(
    level=logging.INFO,
    format="[%(asctime)s] [%(levelname)s] %(message)s",
    datefmt="%Y-%m-%d %H:%M:%S",
)
logger = logging.getLogger("PPTGenerator")


# =====================================================
# Gateway API
# =====================================================
def normalize_message_content(content: Any) -> str:
    if isinstance(content, str):
        return content.strip()
    if isinstance(content, list):
        parts = []
        for item in content:
            if isinstance(item, str):
                parts.append(item)
            elif isinstance(item, dict):
                t = item.get("text", "")
                if t:
                    parts.append(str(t))
        return "\n".join(parts).strip()
    if isinstance(content, dict):
        if "text" in content:
            return str(content["text"]).strip()
        return json.dumps(content, ensure_ascii=False)
    return str(content).strip()


def extract_first_json_object(text: str) -> str:
    text = text.strip()
    if text.startswith("```"):
        text = re.sub(r"^```(?:json)?\s*", "", text)
        text = re.sub(r"\s*```$", "", text)
        text = text.strip()
    try:
        json.loads(text)
        return text
    except Exception:
        pass
    for start_ch, end_ch in [("{", "}"), ("[", "]")]:
        start = text.find(start_ch)
        if start == -1:
            continue
        depth = 0
        in_str = False
        esc = False
        for i in range(start, len(text)):
            ch = text[i]
            if in_str:
                if esc:
                    esc = False
                elif ch == "\\":
                    esc = True
                elif ch == '"':
                    in_str = False
            else:
                if ch == '"':
                    in_str = True
                elif ch == start_ch:
                    depth += 1
                elif ch == end_ch:
                    depth -= 1
                    if depth == 0:
                        return text[start: i + 1]
    return text


def should_retry_exception(exc: Exception) -> bool:
    if isinstance(exc, (httpx.TimeoutException, httpx.TransportError)):
        return True
    if isinstance(exc, httpx.HTTPStatusError):
        code = exc.response.status_code
        return code == 429 or 500 <= code < 600
    return False


async def call_gateway(
    messages: List[Dict[str, str]],
    max_tokens: int = MAX_TOKENS,
    client: Optional[httpx.AsyncClient] = None,
    temperature: float = 0.3,
) -> str:
    if not GW_TOKEN:
        raise ValueError("GW token not found.")
    headers = {
        "Content-Type": "application/json",
        "BCS-APIHub-RequestId": str(uuid.uuid4()),
        "X-CHJ-GWToken": GW_TOKEN,
    }
    payload = {
        "model": MODEL,
        "max_tokens": max_tokens,
        "enable_thinking": ENABLE_THINKING,
        "temperature": temperature,
        "messages": messages,
    }
    close = False
    if client is None:
        client = httpx.AsyncClient(timeout=REQUEST_TIMEOUT_SEC)
        close = True
    try:
        resp = await client.post(API_URL, headers=headers, json=payload)
        resp.raise_for_status()
        data = resp.json()
    finally:
        if close:
            await client.aclose()
    content = data["choices"][0]["message"]["content"]
    return normalize_message_content(content)


async def call_llm(
    messages: List[Dict[str, str]],
    client: httpx.AsyncClient,
    max_retries: int = MAX_RETRIES,
    temperature: float = 0.3,
) -> str:
    for attempt in range(max_retries):
        try:
            return await call_gateway(messages, client=client, temperature=temperature)
        except Exception as e:
            logger.warning(f"LLM call failed (attempt {attempt + 1}/{max_retries}): {e}")
            if attempt == max_retries - 1:
                raise
            if should_retry_exception(e):
                delay = min(BACKOFF_BASE_SEC * (2 ** attempt), MAX_BACKOFF_SEC) + random.uniform(0, 0.5)
                await asyncio.sleep(delay)
    raise RuntimeError("Unreachable")


async def call_llm_json(
    messages: List[Dict[str, str]],
    client: httpx.AsyncClient,
) -> Any:
    raw = await call_llm(messages, client)
    json_str = extract_first_json_object(raw)
    return json.loads(json_str)


# =====================================================
# Step 1: LLM Content Generation
# =====================================================
def query_course_rag(
    service: Optional[CourseRAGService], query: str, *, top_k: int = 10,
    current_chapter: Optional[str] = None,
) -> Dict[str, Any]:
    if service is None:
        return {"query": query, "knowledgeUnits": [], "allowedTerms": [], "allowedFormulas": [], "sources": []}
    result = service.query(
        {
            "query": query,
            "topK": top_k,
            "courseId": "ai-literacy-general",
            "courseVersion": "2026.1",
            "currentChapter": current_chapter,
        }
    )
    return result.get("context") or {}


async def generate_course_brief(
    topic: str,
    total_pages: int,
    language: str,
    client: httpx.AsyncClient,
    skill: S1CoursewareSkill,
    rag_context: Dict[str, Any],
    *,
    target_learners: Optional[str] = None,
    duration_minutes: Optional[int] = None,
) -> Dict[str, Any]:
    messages = skill.brief_messages(
        topic,
        total_pages,
        language,
        rag_context,
        target_learners=target_learners,
        duration_minutes=duration_minutes,
    )
    result = await call_llm_json(messages, client)
    if isinstance(result, dict) and not result.get("course_thesis"):
        nested = result.get("course_brief") or result.get("brief")
        if isinstance(nested, dict):
            result = nested
    if not isinstance(result, dict):
        raise ValueError(f"Expected course brief object, got {type(result)}")
    if not result.get("course_thesis"):
        raise ValueError("Course brief is missing course_thesis")
    if not result.get("sections"):
        raise ValueError("Course brief is missing sections")
    if not result.get("terminology"):
        raise ValueError("Course brief is missing terminology")
    brief_audit = skill.audit_brief(result)
    result = skill.repair_brief_deterministically(result, brief_audit)
    brief_audit = skill.audit_brief(result)
    result["style_audit"] = brief_audit
    if not brief_audit["passed"]:
        raise ValueError(
            "Course brief failed style and structure audit: "
            + json.dumps(brief_audit, ensure_ascii=False)
        )
    return result


async def generate_outline(
    topic: str,
    total_pages: int,
    language: str,
    client: httpx.AsyncClient,
    *,
    course_brief: Optional[Dict[str, Any]] = None,
    section_contexts: Optional[Dict[str, Dict[str, Any]]] = None,
    skill: Optional[S1CoursewareSkill] = None,
) -> List[Dict]:
    skill = skill or S1CoursewareSkill()
    if course_brief:
        messages = skill.outline_messages(course_brief, total_pages, language, section_contexts)
    else:
        messages = [
            {
                "role": "system",
                "content": (
                    f"You are a professional course designer. Generate a structured course outline in {language}.\n"
                    "Return JSON array of chapters. Each chapter has title and a list of page topics.\n"
                    f"Total content pages must equal exactly {total_pages} (excluding cover and table of contents).\n"
                    f'Format: [{{"chapter": "chapter title in {language}", "pages": ["page topic 1"]}}]\n'
                    "Return only JSON."
                ),
            },
            {"role": "user", "content": f"Course topic: {topic}\nLanguage: {language}"},
        ]
    result = await call_llm_json(messages, client)
    if isinstance(result, dict):
        result = (
            result.get("outline")
            or result.get("chapters")
            or result.get("sections")
            or result.get("course_outline")
            or result
        )
    if not isinstance(result, list):
        raise ValueError(f"Expected list, got {type(result)}")
    actual = sum(len(ch.get("pages", [])) for ch in result)
    if actual != total_pages:
        raise ValueError(f"Outline has {actual} pages, expected {total_pages}.")
    return result


async def generate_page_content(
    topic: str, chapter: str, page_topic: str, page_id: int,
    language: str, client: httpx.AsyncClient,
    *,
    course_brief: Optional[Dict[str, Any]] = None,
    page_spec: Optional[Dict[str, Any]] = None,
    rag_context: Optional[Dict[str, Any]] = None,
    skill: Optional[S1CoursewareSkill] = None,
) -> Dict:
    skill = skill or S1CoursewareSkill()
    if course_brief and page_spec:
        messages = skill.page_messages(topic, course_brief, page_spec, language, rag_context)
    else:
        messages = [
            {
                "role": "system",
                "content": (
                    f"You are a professional course content creator. Write ALL content in {language}.\n"
                    "Generate content for one presentation slide and return JSON with title, key_points, slide_body, script_text and image_keyword."
                ),
            },
            {
                "role": "user",
                "content": f"Course: {topic}\nChapter: {chapter}\nPage topic: {page_topic}\nPage ID: {page_id}",
            },
        ]
    result = await call_llm_json(messages, client)
    if isinstance(result, dict) and not result.get("title"):
        nested = result.get("page") or result.get("slide") or result.get("content")
        if isinstance(nested, dict):
            result = nested
    result["page_id"] = page_id
    result["chapter"] = chapter
    if page_spec:
        result["section_id"] = page_spec.get("section_id")
        result["page_spec"] = page_spec
        result["retrieval_query"] = page_spec.get("retrieval_query", "")
    result["rag_context"] = skill.compact_rag_context(rag_context, limit=8)
    return result


async def generate_all_content(
    topic: str,
    total_pages: int,
    language: str,
    client: httpx.AsyncClient,
    *,
    rag_service: Optional[CourseRAGService] = None,
    skill: Optional[S1CoursewareSkill] = None,
    target_learners: Optional[str] = None,
    duration_minutes: Optional[int] = None,
    require_rag: bool = True,
) -> Tuple[Dict[str, Any], List[Dict], Dict[str, Any]]:
    skill = skill or S1CoursewareSkill()
    if require_rag and rag_service is None:
        raise RuntimeError("S1 topic-mode generation requires a ready course RAG service")
    initial_context = query_course_rag(
        rag_service,
        f"{topic} 课程主旨 教学大纲 基本概念 典型案例 专业术语",
        top_k=14,
    )
    if require_rag and not initial_context.get("knowledgeUnits"):
        raise RuntimeError("Course RAG returned no knowledge units for course planning")
    logger.info(f"Generating course brief for '{topic}'...")
    course_brief = await generate_course_brief(
        topic,
        total_pages,
        language,
        client,
        skill,
        initial_context,
        target_learners=target_learners,
        duration_minutes=duration_minutes,
    )

    section_contexts: Dict[str, Dict[str, Any]] = {}
    for section in course_brief.get("sections", []):
        section_id = str(section.get("section_id", ""))
        section_query = str(section.get("retrieval_query") or section.get("section_question") or section.get("title") or topic)
        section_contexts[section_id] = query_course_rag(
            rag_service,
            section_query,
            top_k=10,
            current_chapter=section.get("title"),
        )

    logger.info(f"Generating outline for '{topic}' with {total_pages} pages...")
    outline = await generate_outline(
        topic,
        total_pages,
        language,
        client,
        course_brief=course_brief,
        section_contexts=section_contexts,
        skill=skill,
    )

    logger.info(f"Outline: {len(outline)} chapters")
    for ch in outline:
        logger.info(f"  {ch['chapter']}: {len(ch.get('pages', []))} pages")

    all_pages = []
    page_id = 1
    for chapter_info in outline:
        chapter_title = chapter_info["chapter"]
        section_id = str(chapter_info.get("section_id", ""))
        for raw_page_spec in chapter_info.get("pages", []):
            if isinstance(raw_page_spec, str):
                page_spec = {"page_topic": raw_page_spec, "retrieval_query": raw_page_spec, "section_id": section_id}
            else:
                page_spec = dict(raw_page_spec)
                page_spec.setdefault("section_id", section_id)
            page_topic = str(page_spec.get("page_topic") or page_spec.get("title") or "")
            retrieval_query = str(page_spec.get("retrieval_query") or page_topic)
            page_context = query_course_rag(
                rag_service,
                retrieval_query,
                top_k=8,
                current_chapter=chapter_title,
            )
            if require_rag and not page_context.get("knowledgeUnits"):
                raise RuntimeError(f"Course RAG returned no knowledge units for page {page_id}: {page_topic}")
            logger.info(f"  Generating page {page_id}: {page_topic}")
            page = await generate_page_content(
                topic,
                chapter_title,
                page_topic,
                page_id,
                language,
                client,
                course_brief=course_brief,
                page_spec=page_spec,
                rag_context=page_context,
                skill=skill,
            )
            audit = skill.audit_page(page, course_brief.get("terminology", []))
            page = skill.repair_page_deterministically(page, audit)
            audit = skill.audit_page(page, course_brief.get("terminology", []))
            rewrite_round = 0
            while not audit["passed"] and rewrite_round < skill.config.page_rewrite_limit:
                page = await call_llm_json(
                    skill.rewrite_messages(page, audit, course_brief, page_context),
                    client,
                )
                if isinstance(page, dict) and not page.get("title"):
                    nested_page = page.get("page") or page.get("slide") or page.get("content")
                    if isinstance(nested_page, dict):
                        page = nested_page
                page["page_id"] = page_id
                page["chapter"] = chapter_title
                page["section_id"] = section_id
                page["page_spec"] = page_spec
                page["retrieval_query"] = retrieval_query
                page["rag_context"] = skill.compact_rag_context(page_context, limit=8)
                audit = skill.audit_page(page, course_brief.get("terminology", []))
                rewrite_round += 1
            page["quality_audit"] = audit
            page["rewrite_rounds"] = rewrite_round
            all_pages.append(page)
            page_id += 1

    quality_report = skill.audit_deck(course_brief, all_pages)
    return course_brief, all_pages, quality_report


# =====================================================
# Step 2a: Image Search (Baidu + Bing)
# =====================================================
async def baidu_image_search(keyword: str, num: int = BAIDU_CANDIDATES) -> List[str]:
    urls = []
    headers = {
        "User-Agent": (
            "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
            "AppleWebKit/537.36 (KHTML, like Gecko) "
            "Chrome/120.0.0.0 Safari/537.36"
        ),
        "Referer": "https://image.baidu.com",
        "Accept": "application/json, text/html, */*",
    }

    try:
        async with httpx.AsyncClient(timeout=IMAGE_DOWNLOAD_TIMEOUT, follow_redirects=True) as client:
            # Method 1: acjson API
            params = {
                "tn": "resultjson_com",
                "logid": str(uuid.uuid4()).replace("-", "")[:20],
                "ipn": "rj", "ct": "201326592", "fp": "result",
                "word": keyword, "pn": 0, "rn": num, "gsm": "1e", "ie": "utf-8",
            }
            resp = await client.get("https://image.baidu.com/search/acjson",
                                    params=params, headers=headers)
            if resp.status_code == 200:
                try:
                    data = resp.json()
                    for item in data.get("data", []):
                        if not isinstance(item, dict):
                            continue
                        img_url = item.get("thumbURL") or item.get("middleURL") or item.get("hoverURL")
                        if img_url and img_url.startswith("http"):
                            urls.append(img_url)
                            if len(urls) >= num:
                                break
                except Exception:
                    pass

            # Method 2: HTML scraping
            if not urls:
                resp2 = await client.get("https://image.baidu.com/search/index",
                                         params={"tn": "baiduimage", "word": keyword, "ie": "utf-8"},
                                         headers=headers)
                if resp2.status_code == 200:
                    found = re.findall(r'"thumbURL"\s*:\s*"(https?://[^"]+)"', resp2.text)
                    for u in found:
                        if u not in urls:
                            urls.append(u)
                            if len(urls) >= num:
                                break

            # Method 3: Bing China
            if not urls:
                try:
                    bing_headers = {**headers, "Referer": "https://cn.bing.com"}
                    resp3 = await client.get("https://cn.bing.com/images/search",
                                             params={"q": keyword, "form": "HDRSC2", "first": 1},
                                             headers=bing_headers)
                    if resp3.status_code == 200:
                        found = re.findall(r'murl&quot;:&quot;(https?://[^&]+)&quot;', resp3.text)
                        for u in found:
                            urls.append(u)
                            if len(urls) >= num:
                                break
                except Exception:
                    pass

    except Exception as e:
        logger.warning(f"Image search failed for '{keyword}': {e}")

    logger.info(f"Image search '{keyword}': found {len(urls)} images")
    return urls


async def download_image(url: str) -> Optional[bytes]:
    try:
        headers = {
            "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) Chrome/120.0.0.0",
            "Referer": "https://image.baidu.com",
        }
        async with httpx.AsyncClient(timeout=IMAGE_DOWNLOAD_TIMEOUT, follow_redirects=True) as client:
            resp = await client.get(url, headers=headers)
            if resp.status_code == 200 and len(resp.content) > 1000:
                PILImage.open(io.BytesIO(resp.content)).verify()
                return resp.content
    except Exception:
        pass
    return None


# =====================================================
# Step 2b: CLIP Scoring
# =====================================================
class CLIPScorer:
    def __init__(self):
        self._model = None

    def _load(self):
        if self._model is not None:
            return
        logger.info("Loading CLIP model...")
        try:
            import cn_clip.clip as clip
            from cn_clip.clip import load_from_name
            import torch
            self._device = select_device()
            self._model, self._preprocess = load_from_name(
                "ViT-B-16", device=self._device, download_root=os.path.expanduser("~/.cache/clip")
            )
            self._model.eval()
            self._clip_module = clip
            self._torch = torch
            self._backend = "cn_clip"
        except ImportError:
            import open_clip
            import torch
            self._device = select_device()
            model, _, preprocess = open_clip.create_model_and_transforms(
                "ViT-B-32", pretrained="laion2b_s34b_b79k")
            self._model = model.to(self._device).eval()
            self._preprocess = preprocess
            self._tokenizer = open_clip.get_tokenizer("ViT-B-32")
            self._torch = torch
            self._backend = "open_clip"
        logger.info(f"CLIP loaded ({self._backend}) on {self._device}")

    def score(self, text: str, image_bytes: bytes) -> float:
        self._load()
        torch = self._torch
        try:
            img = PILImage.open(io.BytesIO(image_bytes)).convert("RGB")
            img_input = self._preprocess(img).unsqueeze(0).to(self._device)
            if self._backend == "cn_clip":
                text_input = self._clip_module.tokenize([text]).to(self._device)
            else:
                text_input = self._tokenizer([text]).to(self._device)
            with torch.no_grad():
                img_feat = self._model.encode_image(img_input)
                txt_feat = self._model.encode_text(text_input)
                img_feat = img_feat / img_feat.norm(dim=-1, keepdim=True)
                txt_feat = txt_feat / txt_feat.norm(dim=-1, keepdim=True)
                sim = (img_feat @ txt_feat.T).item()
            return float(sim)
        except Exception as e:
            logger.warning(f"CLIP scoring failed: {e}")
            return 0.0


# =====================================================
# Step 2c: ERNIE-Image Generation (local weights)
# =====================================================
class ImageGenerator:
    """ERNIE-Image based generator, loads from local path."""

    def __init__(self, model_path: str = IMAGE_GEN_MODEL):
        self._pipe = None
        self._model_path = model_path

    def _load(self):
        if self._pipe is not None:
            return
        logger.info(f"Loading ERNIE-Image from {self._model_path}...")
        import torch
        from diffusers import ErnieImagePipeline

        self._pipe = ErnieImagePipeline.from_pretrained(
            self._model_path,
            torch_dtype=torch.bfloat16,
        ).to("cuda")
        self._torch = torch
        logger.info("ERNIE-Image loaded.")

    def generate(self, prompt: str, width: int = 848, height: int = 512, seed: int = -1) -> bytes:
        self._load()
        torch = self._torch
        if seed < 0:
            seed = random.randint(0, 2**31)

        logger.info(f"ERNIE generating: '{prompt[:60]}...' ({width}x{height})")

        image = self._pipe(
            prompt=prompt,
            height=height,
            width=width,
            num_inference_steps=IMAGE_GEN_STEPS,
            guidance_scale=IMAGE_GEN_GUIDANCE,
            use_pe=True,
            generator=torch.Generator("cuda").manual_seed(seed),
        ).images[0]

        buf = io.BytesIO()
        image.save(buf, format="PNG")
        return buf.getvalue()


# =====================================================
# Step 2d: Image Pipeline (Search → CLIP → Generate)
# =====================================================
def _save_image(img_bytes: bytes, image_dir: str, page_id: int, source: str, score: float) -> Optional[str]:
    try:
        os.makedirs(image_dir, exist_ok=True)
        img_path = os.path.join(image_dir, f"page_{page_id:02d}.png")
        img = PILImage.open(io.BytesIO(img_bytes)).convert("RGB")
        img.thumbnail((500, 400), PILImage.LANCZOS)
        img.save(img_path, "PNG")
        logger.info(f"Page {page_id}: Saved ({source}, score={score:.3f})")
        return img_path
    except Exception as e:
        logger.error(f"Image save failed: {e}")
        return None


async def get_page_image(
    keyword: str, text_description: str,
    clip_scorer: CLIPScorer, img_gen: ImageGenerator,
    image_dir: str, page_id: int,
) -> Optional[str]:
    best_bytes = None
    best_score = 0.0
    source = "none"
    query = keyword if keyword else text_description

    # Step 1: Search
    urls = await baidu_image_search(keyword, num=BAIDU_CANDIDATES)
    if urls:
        tasks = [download_image(url) for url in urls]
        results = await asyncio.gather(*tasks, return_exceptions=True)
        downloaded = [b for b in results if isinstance(b, bytes) and b]

        if downloaded:
            for img_bytes in downloaded:
                score = clip_scorer.score(query, img_bytes)
                if score > best_score:
                    best_score = score
                    best_bytes = img_bytes
            logger.info(f"Page {page_id}: Search best CLIP = {best_score:.3f}")
        else:
            logger.info(f"Page {page_id}: All downloads failed")

    # Step 2: Generate fallback if no image or CLIP too low
    need_gen = (best_bytes is None) or (best_score < CLIP_THRESHOLD)

    if need_gen:
        if best_bytes is None:
            reason = "no search results or all downloads failed"
        else:
            reason = f"CLIP {best_score:.3f} < {CLIP_THRESHOLD}"
        logger.info(f"Page {page_id}: {reason}, generating with ERNIE-Image...")

        try:
            gen_prompt = (
                f"Professional educational illustration for a presentation slide about: {keyword}. "
                "Clean, minimalist, flat design, white or light background, "
                "suitable for dark green presentation slide, no text in image. "
                "High quality, clear details, educational diagram style."
            )
            gen_bytes = await asyncio.get_event_loop().run_in_executor(
                None, img_gen.generate, gen_prompt, 848, 512
            )

            if gen_bytes:
                gen_score = clip_scorer.score(query, gen_bytes)
                logger.info(f"Page {page_id}: Generated CLIP = {gen_score:.3f}")

                if gen_score > best_score or best_bytes is None:
                    best_bytes = gen_bytes
                    best_score = gen_score
                    source = "generated"
                else:
                    source = "search"
            else:
                logger.warning(f"Page {page_id}: Generation returned empty")
                source = "search" if best_bytes else "none"

        except Exception as e:
            logger.error(f"Page {page_id}: ERNIE-Image failed: {e}")
            if best_bytes is not None:
                source = "search"
                logger.info(f"Page {page_id}: Falling back to search (score={best_score:.3f})")
    else:
        source = "search"

    if best_bytes:
        return _save_image(best_bytes, image_dir, page_id, source, best_score)

    logger.warning(f"Page {page_id}: No image from any source")
    return None


# =====================================================
# Step 3: LLM Layout Decision
# =====================================================
async def decide_layout(
    page: Dict, has_image: bool, client: httpx.AsyncClient
) -> Dict:
    n_points = len(page.get("key_points", []))
    title = page.get("title", "")
    body_len = len(page.get("slide_body", ""))
    page_id = page.get("page_id", 0)

    messages = [
        {
            "role": "system",
            "content": (
                "You are a professional presentation designer.\n"
                "The slide is 13.333 x 7.5 inches, dark green chalkboard background.\n\n"
                "CORNER RESERVATION for digital human teacher avatar:\n"
                "- Top-left: 2.5 x 1.5 inches reserved (chapter label only)\n"
                "- Top-right: 2.5 x 1.5 inches reserved (keep clear)\n"
                "- Bottom-left: 3.0 x 2.5 inches reserved (keep clear)\n"
                "- Bottom-right: 3.0 x 2.5 inches reserved (primary avatar)\n\n"
                "SAFE ZONE for content:\n"
                "- X: 0.5 to 9.8 inches (leave right side for image or avatar)\n"
                "- Y: 1.0 to 5.0 inches (leave top for title, bottom for avatar)\n"
                "- Title: Y 0.3 to 0.6, full width\n\n"
                "Content elements: Title, Key Points (bullets), Body Text (paragraphs), Image (optional)\n\n"
                "IMPORTANT - LAYOUT VARIATION:\n"
                f"This is slide {page_id}.\n"
                "Each slide MUST have a different layout. Rules:\n"
                "- Odd slides: key points ABOVE body text (points_first=true)\n"
                "- Even slides: body text ABOVE key points (points_first=false)\n"
                "- Slide 1,4,7: points start at Y=1.3\n"
                "- Slide 2,5,8: points start at Y=1.6\n"
                "- Slide 3,6,9: points start at Y=1.8\n"
                "- Vary content_x between 0.6 and 1.2\n"
                "- Vary widths between 5.5 and 8.5 (narrower if has image)\n"
                "- If has image: text left 55%, image right 35%\n"
                "- Image Y varies: 1.3, 1.8, 2.2 depending on page\n\n"
                "Return JSON:\n"
                "{\n"
                '  "title_y": float,\n'
                '  "points_x": float, "points_y": float, "points_w": float, "points_h": float,\n'
                '  "body_x": float, "body_y": float, "body_w": float, "body_h": float,\n'
                '  "points_first": boolean,\n'
                '  "image_x": float, "image_y": float, "image_w": float, "image_h": float\n'
                "}\n"
                "image_* fields only if has_image=true.\n"
                "Return only JSON."
            ),
        },
        {
            "role": "user",
            "content": json.dumps({
                "page_id": page_id,
                "title": title,
                "num_key_points": n_points,
                "body_text_length": body_len,
                "has_image": has_image,
            }, ensure_ascii=False),
        },
    ]

    try:
        layout = await call_llm_json(messages, client)
        return layout
    except Exception as e:
        logger.warning(f"Layout decision failed: {e}, using default with variation")
        return _default_layout(page_id, has_image)


def _default_layout(page_id: int, has_image: bool) -> Dict:
    offset = (page_id % 3) * 0.15
    swap = (page_id % 2 == 0)
    base_px = 0.7 + offset
    upper_y = 1.3 + offset
    lower_y = 3.5 + offset

    if has_image:
        tw = 6.0 + offset
        return {
            "title_y": 0.4,
            "points_x": base_px, "points_w": tw,
            "points_y": upper_y if not swap else lower_y, "points_h": 1.8 + offset,
            "body_x": base_px, "body_w": tw,
            "body_y": lower_y if not swap else upper_y, "body_h": 1.5 + offset * 0.5,
            "points_first": not swap,
            "image_x": 7.8 + offset, "image_y": 1.3 + offset * 2,
            "image_w": 3.5 - offset * 0.5, "image_h": 2.8,
        }
    else:
        tw = 8.5 + offset
        return {
            "title_y": 0.4,
            "points_x": base_px, "points_w": tw,
            "points_y": upper_y if not swap else lower_y, "points_h": 2.0 + offset,
            "body_x": base_px, "body_w": tw,
            "body_y": lower_y if not swap else upper_y, "body_h": 1.8,
            "points_first": not swap,
        }


# =====================================================
# Step 4: PPT Rendering
# =====================================================
def clamp_layout(layout: Dict, key_points: List[str] = None, slide_body: str = "") -> Dict:
    """Dynamic height + forced gap between blocks."""
    s = layout.copy()

    n_points = len(key_points) if key_points else 0
    body_chars = len(slide_body) if slide_body else 0
    est_points_h = max(n_points * 0.45, 1.0) if n_points > 0 else 0
    est_body_h = max((body_chars / 80) * 0.35, 0.8) if body_chars > 0 else 0

    s["title_y"] = max(0.3, min(s.get("title_y", 0.4), 0.6))
    content_start = s["title_y"] + 0.9
    min_gap = 0.3

    for kx in ["points_x", "body_x"]:
        if kx in s:
            s[kx] = max(SAFE_X_MIN, min(s[kx], 2.0))

    points_first = s.get("points_first", True)

    if points_first:
        first_h_est, second_h_est = est_points_h, est_body_h
    else:
        first_h_est, second_h_est = est_body_h, est_points_h

    first_key = "points_y" if points_first else "body_y"
    first_y = max(s.get(first_key, content_start), content_start)

    first_h_key = "points_h" if points_first else "body_h"
    first_h = max(min(s.get(first_h_key, first_h_est), first_h_est + 0.5), 0.5)

    second_y = first_y + first_h + min_gap

    second_h_key = "body_h" if points_first else "points_h"
    second_h = max(min(s.get(second_h_key, second_h_est), second_h_est + 0.5), 0.5)

    if second_y + second_h > SAFE_Y_MAX:
        available = SAFE_Y_MAX - content_start - min_gap
        total_est = first_h_est + second_h_est
        ratio = (first_h_est / total_est) if total_est > 0 else 0.5
        first_h = max(available * ratio, 0.5)
        second_h = max(available * (1 - ratio), 0.5)
        first_y = content_start
        second_y = first_y + first_h + min_gap

    if points_first:
        s["points_y"], s["points_h"] = first_y, first_h
        s["body_y"], s["body_h"] = second_y, second_h
    else:
        s["body_y"], s["body_h"] = first_y, first_h
        s["points_y"], s["points_h"] = second_y, second_h

    has_image = "image_x" in s
    for kw in ["points_w", "body_w"]:
        if kw in s:
            kx_key = kw.replace("_w", "_x")
            start_x = s.get(kx_key, 0.8)
            max_w = (s.get("image_x", 8.0) - start_x - 0.4) if has_image else (SAFE_X_MAX - start_x)
            s[kw] = max(min(s[kw], max_w), 4.0)

    if "image_x" in s:
        s["image_x"] = max(min(s["image_x"], 9.5), 5.0)
    if "image_y" in s:
        s["image_y"] = max(s.get("image_y", 1.3), content_start)
        img_bottom = s["image_y"] + s.get("image_h", 3.0)
        if img_bottom > SAFE_Y_MAX:
            s["image_h"] = max(SAFE_Y_MAX - s["image_y"], 1.5)

    return s


def set_slide_bg(slide, color: RGBColor):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color


def add_text_box(
    slide, text: str, x: float, y: float, w: float, h: float,
    font_size: int = 24, color: RGBColor = TITLE_COLOR,
    bold: bool = False, alignment=PP_ALIGN.LEFT, font_name: str = "Microsoft YaHei",
):
    txBox = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.color.rgb = color
    p.font.bold = bold
    p.font.name = font_name
    p.alignment = alignment
    return txBox


def add_bullet_list(
    slide, points: List[str], x: float, y: float, w: float, h: float,
    font_size: int = 20, color: RGBColor = BULLET_COLOR,
    line_spacing: float = 1.5, font_name: str = "Microsoft YaHei",
):
    txBox = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = txBox.text_frame
    tf.word_wrap = True
    for i, point in enumerate(points):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = f"▸ {point}"
        p.font.size = Pt(font_size)
        p.font.color.rgb = color
        p.font.name = font_name
        p.space_after = Pt(font_size * (line_spacing - 1))
    return txBox


def add_body_text(
    slide, text: str, x: float, y: float, w: float, h: float,
    font_size: int = 15, color: RGBColor = BODY_COLOR,
    font_name: str = "Microsoft YaHei",
):
    txBox = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = txBox.text_frame
    tf.word_wrap = True
    sentences = [s.strip() for s in re.split(r'(?<=[。！？.!?])\s*', text) if s.strip()]
    for i, sentence in enumerate(sentences):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = sentence
        p.font.size = Pt(font_size)
        p.font.color.rgb = color
        p.font.name = font_name
        p.space_after = Pt(6)
        p.line_spacing = Pt(font_size * 1.6)
    return txBox


def add_divider_line(slide, x: float, y: float, w: float):
    shape = slide.shapes.add_shape(1, Inches(x), Inches(y), Inches(w), Pt(1))
    shape.fill.solid()
    shape.fill.fore_color.rgb = DIVIDER_COLOR
    shape.line.fill.background()


def render_cover_slide(prs: Presentation, topic: str, page_count: int):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, BG_COLOR)
    add_text_box(slide, topic, 1.5, 2.2, 10.0, 2.0, font_size=48, color=TITLE_COLOR, bold=True, alignment=PP_ALIGN.CENTER)
    add_divider_line(slide, 4.0, 4.3, 5.0)
    add_text_box(slide, f"共 {page_count} 页", 1.5, 4.6, 10.0, 0.8, font_size=22, color=SUBTITLE_COLOR, alignment=PP_ALIGN.CENTER)
    add_text_box(slide, datetime.now().strftime("%Y年%m月%d日"), 1.5, 5.5, 10.0, 0.6, font_size=16, color=SUBTITLE_COLOR, alignment=PP_ALIGN.CENTER)


def render_toc_slide(prs: Presentation, chapters: List[str]):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, BG_COLOR)
    add_text_box(slide, "目 录", 0.8, 0.4, 11.0, 0.8, font_size=36, color=TITLE_COLOR, bold=True)
    add_divider_line(slide, 0.8, 1.2, 11.5)
    add_bullet_list(slide, [f"{i+1}. {ch}" for i, ch in enumerate(chapters)],
                    1.2, 1.6, 9.0, 4.5, font_size=24, color=BULLET_COLOR, line_spacing=2.0)


def render_content_slide(
    prs: Presentation, page: Dict, layout: Dict, image_path: Optional[str],
):
    key_points = page.get("key_points", [])
    slide_body = page.get("slide_body", "")
    layout = clamp_layout(layout, key_points=key_points, slide_body=slide_body)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, BG_COLOR)

    title = page.get("title", "")
    chapter = page.get("chapter", "")
    page_id = page.get("page_id", "")

    if chapter:
        add_text_box(slide, chapter, 0.5, 0.1, 5.0, 0.3, font_size=11, color=SUBTITLE_COLOR)
    add_text_box(slide, str(page_id), 11.5, 0.1, 1.0, 0.3, font_size=11, color=SUBTITLE_COLOR, alignment=PP_ALIGN.RIGHT)

    title_y = layout.get("title_y", 0.4)
    add_text_box(slide, title, 0.5, title_y, 11.0, 0.7, font_size=32, color=TITLE_COLOR, bold=True)
    add_divider_line(slide, 0.5, title_y + 0.7, 11.0)

    points_first = layout.get("points_first", True)

    if points_first:
        if key_points:
            add_bullet_list(slide, key_points, layout.get("points_x", 0.8), layout["points_y"],
                            layout.get("points_w", 6.5), layout["points_h"], font_size=20, color=BULLET_COLOR)
        if slide_body:
            add_body_text(slide, slide_body, layout.get("body_x", 0.8), layout["body_y"],
                          layout.get("body_w", 6.5), layout["body_h"], font_size=15, color=BODY_COLOR)
    else:
        if slide_body:
            add_body_text(slide, slide_body, layout.get("body_x", 0.8), layout["body_y"],
                          layout.get("body_w", 6.5), layout["body_h"], font_size=15, color=BODY_COLOR)
        if key_points:
            add_bullet_list(slide, key_points, layout.get("points_x", 0.8), layout["points_y"],
                            layout.get("points_w", 6.5), layout["points_h"], font_size=20, color=BULLET_COLOR)

    if image_path and os.path.exists(image_path):
        try:
            slide.shapes.add_picture(image_path,
                                     Inches(layout.get("image_x", 8.0)), Inches(layout.get("image_y", 1.5)),
                                     Inches(layout.get("image_w", 3.5)), Inches(layout.get("image_h", 2.8)))
        except Exception as e:
            logger.warning(f"Failed to add image: {e}")


def render_end_slide(prs: Presentation, topic: str):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, BG_COLOR)
    add_text_box(slide, "谢谢观看", 1.5, 2.5, 10.0, 1.5, font_size=48, color=TITLE_COLOR, bold=True, alignment=PP_ALIGN.CENTER)
    add_divider_line(slide, 4.0, 4.2, 5.0)
    add_text_box(slide, topic, 1.5, 4.6, 10.0, 0.8, font_size=20, color=SUBTITLE_COLOR, alignment=PP_ALIGN.CENTER)


# =====================================================
# Main Pipeline
# =====================================================
class PPTGenerator:
    def __init__(
        self,
        workspace_root: str = WORKSPACE_ROOT,
        *,
        rag_root: Optional[str] = None,
        require_rag: bool = True,
    ):
        self.workspace_root = workspace_root
        self.clip_scorer = CLIPScorer()
        self.img_gen = ImageGenerator()
        self.courseware_skill = S1CoursewareSkill()
        self.require_rag = require_rag
        self.rag_root = rag_root or os.environ.get("SMARTCOURSE_RAG_ROOT", "")
        self.rag_service: Optional[CourseRAGService] = None
        if self.rag_root:
            self.rag_service = CourseRAGService(self.rag_root)
            if not self.rag_service.status().get("ready"):
                raise RuntimeError(f"Course RAG index is not ready: {self.rag_root}")

    def run(self, task_id: str, course_topic: str, total_pages: int = 10,
            target_language: str = "zh", image_source: str = "both",
            target_learners: Optional[str] = None,
            duration_minutes: Optional[int] = None) -> Dict:
        return asyncio.run(
            self.arun(
                task_id,
                course_topic,
                total_pages,
                target_language,
                image_source,
                target_learners=target_learners,
                duration_minutes=duration_minutes,
            )
        )

    async def arun(self, task_id: str, course_topic: str, total_pages: int = 10,
                   target_language: str = "zh", image_source: str = "both",
                   *, target_learners: Optional[str] = None,
                   duration_minutes: Optional[int] = None) -> Dict:
        logger.info(f"=== PPTGenerator START: '{course_topic}', {total_pages} pages ===")

        output_dir = os.path.join(self.workspace_root, task_id, "s1_output")
        image_dir = os.path.join(output_dir, "images")
        os.makedirs(output_dir, exist_ok=True)
        os.makedirs(image_dir, exist_ok=True)

        async with httpx.AsyncClient(timeout=REQUEST_TIMEOUT_SEC) as client:
            # Step 1
            logger.info("Step 1: Generating teacher preparation plan and grounded content...")
            course_brief, pages, quality_report = await generate_all_content(
                course_topic,
                total_pages,
                target_language,
                client,
                rag_service=self.rag_service,
                skill=self.courseware_skill,
                target_learners=target_learners,
                duration_minutes=duration_minutes,
                require_rag=self.require_rag,
            )

            # Step 2
            if image_source != "none":
                logger.info("Step 2: Getting images...")
                for page in pages:
                    visual_plan = page.get("visual_plan") or {}
                    visual_objects = " ".join(str(item) for item in visual_plan.get("objects", []))
                    keyword = page.get("image_keyword") or f"{page.get('title', '')} {visual_objects}".strip()
                    title = page.get("title", "")
                    pid = page.get("page_id", 0)

                    if image_source == "search":
                        img_path = await self._get_search_only(keyword, title, image_dir, pid)
                    elif image_source == "generate":
                        img_path = await self._get_gen_only(keyword, image_dir, pid)
                    else:  # both
                        img_path = await get_page_image(
                            keyword, title, self.clip_scorer, self.img_gen, image_dir, pid)
                    page["image_path"] = img_path
            else:
                logger.info("Step 2: Skipping images")
                for page in pages:
                    page["image_path"] = None

            # Step 3
            logger.info("Step 3: Deciding layouts...")
            layouts = []
            for page in pages:
                has_image = page.get("image_path") is not None
                layout = await decide_layout(page, has_image, client)
                layouts.append(layout)

            # Step 4
            logger.info("Step 4: Rendering PPT...")
            prs = Presentation()
            prs.slide_width = SLIDE_WIDTH
            prs.slide_height = SLIDE_HEIGHT

            render_cover_slide(prs, course_topic, total_pages)

            chapters = []
            seen = set()
            for page in pages:
                ch = page.get("chapter", "")
                if ch and ch not in seen:
                    chapters.append(ch)
                    seen.add(ch)
            if chapters:
                render_toc_slide(prs, chapters)

            for page, layout in zip(pages, layouts):
                render_content_slide(prs, page, layout, page.get("image_path"))

            render_end_slide(prs, course_topic)

            pptx_path = os.path.join(output_dir, "generated.pptx")
            prs.save(pptx_path)
            logger.info(f"PPT saved: {pptx_path}")

            # Step 5: Save outputs
            brief_path = os.path.join(output_dir, "course_brief.json")
            with open(brief_path, "w", encoding="utf-8") as f:
                json.dump(course_brief, f, indent=2, ensure_ascii=False)

            quality_path = os.path.join(output_dir, "course_quality_report.json")
            with open(quality_path, "w", encoding="utf-8") as f:
                json.dump(quality_report, f, indent=2, ensure_ascii=False)

            structure = {
                "task_id": task_id,
                "input_mode": "mode_b",
                "course_topic": course_topic,
                "course_brief": course_brief,
                "quality_report": quality_report,
                "pages": [{
                    "page_id": p["page_id"], "chapter": p.get("chapter", ""),
                    "section_id": p.get("section_id", ""),
                    "title": p.get("title", ""), "key_points": p.get("key_points", []),
                    "slide_body": p.get("slide_body", ""), "script_text": p.get("script_text", ""),
                    "image_keyword": p.get("image_keyword", ""), "ppt_image": p.get("image_path", ""),
                    "opening_question": p.get("opening_question", ""),
                    "worked_example": p.get("worked_example", ""),
                    "visual_plan": p.get("visual_plan", {}),
                    "terminology_used": p.get("terminology_used", []),
                    "retrieval_query": p.get("retrieval_query", ""),
                    "source_unit_ids": p.get("source_unit_ids", []),
                    "rag_context": p.get("rag_context", {}),
                    "quality_audit": p.get("quality_audit", {}),
                    "rewrite_rounds": p.get("rewrite_rounds", 0),
                } for p in pages],
            }
            structure_path = os.path.join(output_dir, "course_structure.json")
            with open(structure_path, "w", encoding="utf-8") as f:
                json.dump(structure, f, indent=2, ensure_ascii=False)

            script_path = os.path.join(output_dir, "script.txt")
            with open(script_path, "w", encoding="utf-8") as f:
                for p in pages:
                    f.write(f"=== Page {p['page_id']}: {p.get('title', '')} ===\n")
                    f.write(f"[章节] {p.get('chapter', '')}\n")
                    f.write(f"[要点] {' | '.join(p.get('key_points', []))}\n")
                    f.write(f"[正文] {p.get('slide_body', '')}\n")
                    f.write(f"[讲稿]\n{p.get('script_text', '')}\n\n")

            logger.info(f"=== PPTGenerator DONE ===")
            logger.info(f"  PPT: {pptx_path}")
            logger.info(f"  Structure: {structure_path}")
            logger.info(f"  Script: {script_path}")
            logger.info(f"  Course brief: {brief_path}")
            logger.info(f"  Quality report: {quality_path}")

            return {
                "task_id": task_id,
                "pptx_path": pptx_path,
                "structure_path": structure_path,
                "script_path": script_path,
                "course_brief_path": brief_path,
                "quality_report_path": quality_path,
                "pages": len(pages),
            }

    async def _get_search_only(self, keyword, title, image_dir, page_id):
        urls = await baidu_image_search(keyword)
        if not urls:
            return None
        best_bytes, best_score = None, 0.0
        tasks = [download_image(url) for url in urls]
        results = await asyncio.gather(*tasks, return_exceptions=True)
        for img_bytes in results:
            if isinstance(img_bytes, bytes) and img_bytes:
                score = self.clip_scorer.score(keyword, img_bytes)
                if score > best_score:
                    best_score = score
                    best_bytes = img_bytes
        if best_bytes and best_score > 0.3:
            return _save_image(best_bytes, image_dir, page_id, "search", best_score)
        return None

    async def _get_gen_only(self, keyword, image_dir, page_id):
        try:
            prompt = (
                f"Professional educational illustration: {keyword}. "
                "Clean, minimalist, flat design, suitable for presentation. No text in image."
            )
            gen_bytes = await asyncio.get_event_loop().run_in_executor(
                None, self.img_gen.generate, prompt, 848, 512)
            return _save_image(gen_bytes, image_dir, page_id, "generated", 1.0)
        except Exception as e:
            logger.error(f"Generation failed for page {page_id}: {e}")
            return None


# =====================================================
# CLI
# =====================================================
def main():
    import argparse
    parser = argparse.ArgumentParser(description="S1 PPT Generator v5")
    parser.add_argument("--task-id", type=str, default=str(uuid.uuid4())[:8])
    parser.add_argument("--topic", type=str, required=True, help="Course topic")
    parser.add_argument("--pages", type=int, default=10, help="Total content pages")
    parser.add_argument("--language", type=str, default="zh", help="Content language")
    parser.add_argument("--image-source", type=str, default="both",
                        choices=["search", "generate", "both", "none"],
                        help="search=Baidu/Bing only, generate=ERNIE-Image only, both=search+fallback, none=no images")
    parser.add_argument("--workspace", type=str, default=WORKSPACE_ROOT)
    parser.add_argument("--rag-root", type=str, default=os.environ.get("SMARTCOURSE_RAG_ROOT", ""))
    parser.add_argument("--allow-no-rag", action="store_true")
    parser.add_argument("--target-learners", type=str, default="高校本科生")
    parser.add_argument("--duration-minutes", type=int, default=45)
    args = parser.parse_args()

    gen = PPTGenerator(
        workspace_root=args.workspace,
        rag_root=args.rag_root,
        require_rag=not args.allow_no_rag,
    )
    result = gen.run(
        task_id=args.task_id, course_topic=args.topic,
        total_pages=args.pages, target_language=args.language,
        image_source=args.image_source,
        target_learners=args.target_learners,
        duration_minutes=args.duration_minutes,
    )

    print("\n" + "=" * 60)
    print("PPT Generation Complete")
    print("=" * 60)
    print(f"  Task ID:    {result['task_id']}")
    print(f"  Pages:      {result['pages']}")
    print(f"  PPT:        {result['pptx_path']}")
    print(f"  Structure:  {result['structure_path']}")
    print(f"  Script:     {result['script_path']}")
    print(f"  Brief:      {result['course_brief_path']}")
    print(f"  Quality:    {result['quality_report_path']}")
    print("=" * 60)


if __name__ == "__main__":
    main()

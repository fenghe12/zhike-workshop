
"""S1 课程生成智能体：PPT内容提取 + 讲稿结构化"""
import os
import json
from pathlib import Path
from pptx import Presentation


class CourseAgent:
    def __init__(self, workspace="/mnt/pfs-gv8sxa/tts/dhg/smartcourse/tasks"):
        self.workspace = Path(workspace)

    def run(self, task_id, ppt_file, script_file=None):
        task_dir = self.workspace / task_id
        output_dir = task_dir / "s1_output"
        output_dir.mkdir(parents=True, exist_ok=True)

        pages = self._extract_ppt(ppt_file)
        pages = self._attach_script(pages, script_file)

        result = {
            "task_id": task_id,
            "page_count": len(pages),
            "pages": pages,
            "structured_json": str(output_dir / "course_structure.json")
        }

        with open(result["structured_json"], "w", encoding="utf-8") as f:
            json.dump(result, f, ensure_ascii=False, indent=2)

        return result

    def _extract_ppt(self, ppt_file):
        prs = Presentation(ppt_file)
        pages = []
        for slide_idx, slide in enumerate(prs.slides, start=1):
            title = ""
            key_points = []
            other_texts = []
            for shape in slide.shapes:
                if not shape.has_text_frame:
                    continue
                text = shape.text_frame.text.strip()
                if not text:
                    continue
                if shape == slide.shapes.title:
                    title = text
                elif len(text) < 80:
                    key_points.append(text)
                else:
                    other_texts.append(text)
            pages.append({
                "page_id": slide_idx,
                "title": title,
                "key_points": key_points,
                "script_text": "",
                "other_texts": other_texts
            })
        return pages

    def _attach_script(self, pages, script_file=None):
        paragraphs = []
        if script_file and os.path.isfile(script_file):
            with open(script_file, "r", encoding="utf-8") as f:
                script = f.read()
            paragraphs = [p.strip() for p in script.split("\n\n") if p.strip()]
        for i, page in enumerate(pages):
            if i < len(paragraphs):
                page["script_text"] = paragraphs[i]
            else:
                values = [*page.get("key_points", []), *page.get("other_texts", [])]
                page["script_text"] = "。".join(value for value in values if value)
        return pages

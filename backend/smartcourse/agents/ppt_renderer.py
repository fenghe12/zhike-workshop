"""Render PPTX slides to readable PNG files without LibreOffice.

The renderer covers the elements used by SmartCourse generated presentations
and common teaching slides: backgrounds, text boxes, placeholders, tables,
pictures and basic shapes. Geometry is taken from the PPTX document, so the
same images can also be used for avatar placement analysis.
"""

from __future__ import annotations

import os
from pathlib import Path
from typing import Any, Iterable

from PIL import Image, ImageDraw, ImageFont


EMU_PER_POINT = 12_700


def _font_path() -> str:
    configured = os.environ.get("SMARTCOURSE_CJK_FONT", "").strip()
    candidates = [
        configured,
        "/home/ma-user/work/smartcourse/shared/smartcourse-cn-font.ttc",
        "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc",
        "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
        "/System/Library/Fonts/STHeiti Medium.ttc",
    ]
    for candidate in candidates:
        if candidate and Path(candidate).is_file():
            return candidate
    raise FileNotFoundError("未找到可用于 PPT 渲染的字体")


def _font(size: int) -> ImageFont.FreeTypeFont:
    return ImageFont.truetype(_font_path(), max(12, int(size)))


def _rgb(value: Any, default: tuple[int, int, int]) -> tuple[int, int, int]:
    try:
        rgb = value.rgb
        if rgb is not None:
            return tuple(int(component) for component in bytes(rgb))
    except Exception:
        pass
    return default


def _fill_color(fill: Any, default: tuple[int, int, int] | None = None):
    try:
        if fill.type is not None:
            return _rgb(fill.fore_color, default or (255, 255, 255))
    except Exception:
        pass
    return default


def _line_color(shape: Any, default: tuple[int, int, int] = (210, 218, 226)):
    try:
        return _rgb(shape.line.color, default)
    except Exception:
        return default


def _wrap_text(draw: ImageDraw.ImageDraw, text: str, font, max_width: int) -> list[str]:
    lines: list[str] = []
    for source_line in str(text).splitlines() or [""]:
        if not source_line:
            lines.append("")
            continue
        current = ""
        for char in source_line:
            candidate = current + char
            if current and draw.textlength(candidate, font=font) > max_width:
                lines.append(current)
                current = char
            else:
                current = candidate
        if current:
            lines.append(current)
    return lines


def _paragraph_font(paragraph: Any, default_size: int, scale: float):
    runs = list(paragraph.runs)
    run = runs[0] if runs else None
    size_points = None
    if run is not None and run.font.size is not None:
        size_points = float(run.font.size.pt)
    size = int((size_points or default_size) * scale)
    return _font(size), run


def _draw_text_frame(
    image: Image.Image,
    draw: ImageDraw.ImageDraw,
    shape: Any,
    box: tuple[int, int, int, int],
    scale: float,
    is_title: bool,
) -> None:
    x1, y1, x2, y2 = box
    margin_x = max(12, int(12 * scale))
    margin_y = max(8, int(8 * scale))
    cursor_y = y1 + margin_y
    max_width = max(20, x2 - x1 - margin_x * 2)
    default_points = 31 if is_title else 20
    for paragraph in shape.text_frame.paragraphs:
        value = paragraph.text.strip()
        if not value:
            cursor_y += max(8, int(6 * scale))
            continue
        font, run = _paragraph_font(paragraph, default_points, scale)
        color = (24, 42, 59) if not is_title else (8, 78, 97)
        bold = is_title
        if run is not None:
            color = _rgb(run.font.color, color)
            bold = bool(run.font.bold) or bold
        prefix = ""
        try:
            if int(paragraph.level) >= 0 and not is_title:
                prefix = "• "
        except Exception:
            pass
        lines = _wrap_text(draw, prefix + value, font, max_width)
        line_height = int(font.size * 1.30)
        for line in lines:
            if cursor_y + line_height > y2 - margin_y:
                return
            draw.text(
                (x1 + margin_x, cursor_y),
                line,
                font=font,
                fill=color,
                stroke_width=1 if bold and font.size >= 28 else 0,
                stroke_fill=color,
            )
            cursor_y += line_height
        cursor_y += max(5, int(5 * scale))


def _draw_table(image: Image.Image, draw: ImageDraw.ImageDraw, shape: Any, box, scale: float):
    x1, y1, x2, y2 = box
    rows = len(shape.table.rows)
    cols = len(shape.table.columns)
    if not rows or not cols:
        return
    cell_w = (x2 - x1) / cols
    cell_h = (y2 - y1) / rows
    font = _font(int(14 * scale))
    for row_index, row in enumerate(shape.table.rows):
        for col_index, cell in enumerate(row.cells):
            left = int(x1 + col_index * cell_w)
            top = int(y1 + row_index * cell_h)
            right = int(x1 + (col_index + 1) * cell_w)
            bottom = int(y1 + (row_index + 1) * cell_h)
            fill = (231, 246, 242) if row_index == 0 else (255, 255, 255)
            draw.rectangle((left, top, right, bottom), fill=fill, outline=(170, 185, 198), width=2)
            text = cell.text.strip()
            for line_index, line in enumerate(_wrap_text(draw, text, font, max(10, right-left-16))[:3]):
                draw.text((left + 8, top + 6 + line_index * int(font.size * 1.2)), line, font=font, fill=(25, 45, 61))


def _shape_box(shape: Any, x_scale: float, y_scale: float) -> tuple[int, int, int, int]:
    x1 = int(int(shape.left) * x_scale)
    y1 = int(int(shape.top) * y_scale)
    x2 = x1 + max(1, int(int(shape.width) * x_scale))
    y2 = y1 + max(1, int(int(shape.height) * y_scale))
    return x1, y1, x2, y2


def render_pptx_to_images(
    pptx_path: str | Path,
    output_dir: str | Path,
    width: int = 2500,
) -> list[str]:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    presentation = Presentation(str(pptx_path))
    output = Path(output_dir)
    output.mkdir(parents=True, exist_ok=True)
    height = max(1, round(width * int(presentation.slide_height) / int(presentation.slide_width)))
    x_scale = width / int(presentation.slide_width)
    y_scale = height / int(presentation.slide_height)
    point_scale = width / (int(presentation.slide_width) / EMU_PER_POINT)
    results: list[str] = []

    for slide_index, slide in enumerate(presentation.slides, start=1):
        background = (242, 246, 249)
        try:
            background = _fill_color(slide.background.fill, background) or background
        except Exception:
            pass
        image = Image.new("RGB", (width, height), background)
        draw = ImageDraw.Draw(image)
        title_shape = slide.shapes.title
        for shape in slide.shapes:
            box = _shape_box(shape, x_scale, y_scale)
            shape_type = getattr(shape, "shape_type", None)
            if shape_type == MSO_SHAPE_TYPE.PICTURE:
                try:
                    picture = Image.open(__import__("io").BytesIO(shape.image.blob)).convert("RGB")
                    picture = picture.resize((box[2]-box[0], box[3]-box[1]), Image.Resampling.LANCZOS)
                    image.paste(picture, (box[0], box[1]))
                except Exception:
                    draw.rectangle(box, fill=(226, 233, 238), outline=(173, 187, 198), width=2)
                continue
            fill = _fill_color(getattr(shape, "fill", None), None)
            if fill is not None:
                draw.rounded_rectangle(box, radius=max(0, int(5 * point_scale)), fill=fill, outline=_line_color(shape), width=max(1, int(point_scale)))
            if getattr(shape, "has_table", False):
                _draw_table(image, draw, shape, box, point_scale)
            if getattr(shape, "has_text_frame", False) and shape.text.strip():
                _draw_text_frame(image, draw, shape, box, point_scale, shape is title_shape)
        path = output / f"slide_{slide_index:03d}.png"
        image.save(path, "PNG", optimize=True)
        results.append(str(path))
    return results


def occupied_regions(pptx_path: str | Path, image_width: int = 2500) -> list[list[tuple[int, int, int, int]]]:
    """Include active formula branches, grouped objects, layouts and master text."""
    import zipfile
    from lxml import etree as E
    from app.pptx_math_images import inherited, related
    from app.pptx_text_fit import slide_parts
    a='http://schemas.openxmlformats.org/drawingml/2006/main'
    p='http://schemas.openxmlformats.org/presentationml/2006/main'
    m='http://schemas.openxmlformats.org/officeDocument/2006/math'
    mc='http://schemas.openxmlformats.org/markup-compatibility/2006'
    ns={'a':a,'p':p}
    with zipfile.ZipFile(pptx_path) as archive:
        data={name:archive.read(name) for name in archive.namelist()}
    size=E.fromstring(data['ppt/presentation.xml']).find('{'+p+'}sldSz')
    scale=image_width/int(size.get('cx'))
    slides=[]
    for slide_part in slide_parts(data):
        parts=[slide_part]
        layout=related(data,slide_part,'slideLayout');parts.append(layout)
        if E.fromstring(data[slide_part]).get('showMasterSp','1')!='0':
            parts.append(related(data,layout,'slideMaster'))
        regions=[]
        for part in parts:
            root=E.fromstring(data[part])
            for shape in root.iter():
                if shape.tag not in {'{'+p+'}sp','{'+p+'}pic','{'+p+'}graphicFrame','{'+p+'}grpSp'}:
                    continue
                ancestors=list(shape.iterancestors())
                if any(node.tag=='{'+mc+'}Fallback' or node.tag=='{'+p+'}grpSp' for node in ancestors):
                    continue
                contains_text=any((node.text or '').strip() for node in shape.iter() if node.tag in {'{'+a+'}t','{'+m+'}t'})
                contains_media=shape.tag in {'{'+p+'}pic','{'+p+'}graphicFrame'} or any(node.tag=='{'+p+'}pic' for node in shape.iter())
                if not contains_text and not contains_media:
                    continue
                transform=shape.find('p:spPr/a:xfrm',ns)
                if transform is None:transform=shape.find('p:grpSpPr/a:xfrm',ns)
                if transform is None:transform=shape.find('p:xfrm',ns)
                if transform is None and shape.tag=='{'+p+'}sp':
                    try:transform=inherited(data,part,shape,include_transform=True)[3]
                    except (ValueError,KeyError,AttributeError):continue
                if transform is None:continue
                offset=transform.find('a:off',ns);extent=transform.find('a:ext',ns)
                if offset is None or extent is None:continue
                x,y=int(offset.get('x')),int(offset.get('y'))
                w,h=int(extent.get('cx')),int(extent.get('cy'))
                if w>0 and h>0:regions.append((round(x*scale),round(y*scale),round((x+w)*scale),round((y+h)*scale)))
        slides.append(list(dict.fromkeys(regions)))
    return slides
